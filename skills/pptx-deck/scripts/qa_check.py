#!/usr/bin/env python3
"""Quality checks for a finished .pptx. Returns JSON {status, issues}.

Usage:
    qa_check.py DECK.pptx [--spec SPEC.json] [--template TEMPLATE.pptx] [--json]
                [--min-body 18] [--min-caption 12] [--min-title 28] [--allow-footers]

Status
------
``blocked``  the package itself is invalid or unreadable (ooxml.py validation);
``revise``   at least one issue with severity ``error``;
``pass``     no errors (warnings may remain).
Exit status: 0 pass, 1 revise, 2 blocked.

Each issue is ``{"slide", "shape", "severity", "code", "message"}`` (slide is
1-based, 0 for deck-level issues; shape is a name or id path, "" if none).

Checks (code - severity)
------------------------
off-slide - error for text/pictures/charts/tables, warning for decorative shapes
    shape box extends beyond the slide (``off-slide-hidden``: wholly outside).
text-overflow - error when the text clearly needs more height than its box
    (heuristic, from _common.py's estimator and inherited font sizes), warning
    when marginal; boxes with shrink-on-overflow autofit are skipped.
font-too-small - error: a run is below the caption minimum, a title below the
    title minimum, or a paragraph of body text below the body minimum (only
    paragraphs over 90 characters count as body text; short text may be a
    caption). The title is a title placeholder, else the largest text in the
    top quarter of the slide. Sizes include autofit font scaling.
leftover-placeholder - error: empty placeholder, "Click to add ...", "Lorem
    ipsum", TODO / TBD / XXX, "[bracketed]" prompts, template sample phrases
    (and, with --template, any sentence copied verbatim from the template).
slide-number-field / date-field / footer-placeholder - error (warning with
    --allow-footers): slide number, date or footer fields and placeholders.
too-many-fonts - warning: more than 3 distinct font families.
font-missing - warning: a font family is not installed on this machine.
image-reused - error for large pictures, warning for small marks: the same image
    blob appears more than once across slides.
no-notes - warning: slide without speaker notes.
low-contrast - warning: text colour against a solid fill or slide background
    below 4.5:1 (3:1 for large text), only where both are determinable.
layout-repeat - warning: the same layout structure used on more than 2
    consecutive slides.
With --spec: slide-count (error), title-missing (error).

Minimum font sizes default to x-typography.minPt in the deck-spec schema.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

sys.path.insert(0, str(Path(__file__).resolve().parent))
import _common as C  # noqa: E402

C.require_pptx()

from pptx import Presentation  # noqa: E402

import ooxml  # noqa: E402
import template_profile as TP  # noqa: E402
from inventory import abs_box, iter_shapes_abs, localname  # noqa: E402

A = TP.A
P = TP.P
EMU = C.EMU_PER_INCH

EXIT_PASS, EXIT_REVISE, EXIT_BLOCKED = 0, 1, 2

PLACEHOLDER_PATTERNS = [
    (re.compile(r"click to (add|edit)|click here to", re.I), "prompt text 'Click to ...'"),
    (re.compile(r"lorem ipsum|dolor sit amet|consectetur adipiscing", re.I), "Lorem ipsum filler"),
    (re.compile(r"\b(TODO|TBD|FIXME)\b|\bX{3,}\b"), "TODO / TBD / XXX marker"),
    (re.compile(r"\[[^\]\n]{0,60}\]"), "[bracketed] placeholder"),
    (
        re.compile(
            r"\b(your (title|text|name|company|content|logo|headline|subtitle)( goes)? here|text goes here|title goes here|"
            r"add (your )?(text|title|content|subtitle|description)( here)?|insert (your )?(text|title|content|image|picture)( here)?|"
            r"(sample|dummy|placeholder) (text|title|content)|type (your )?(text|title) here|(headline|subtitle|description|caption) here|"
            r"presentation title|company name here|name surname|john doe|jane doe)\b",
            re.I,
        ),
        "template sample phrase",
    ),
]

STRAY_BRACKET = re.compile(r"\[")


class Issues:
    def __init__(self) -> None:
        self.items: List[Dict[str, Any]] = []
        self._seen: Set[Tuple[Any, ...]] = set()

    def add(self, slide: int, shape: str, severity: str, code: str, message: str) -> None:
        key = (slide, shape, code, message)
        if key in self._seen:
            return
        self._seen.add(key)
        self.items.append({"slide": slide, "shape": shape, "severity": severity, "code": code, "message": message})


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------


def shape_label(shape: Any, path: str) -> str:
    return "%s (%s)" % (shape.name, path) if shape.name else path


def is_title_shape(shape: Any, slot_role: Optional[str]) -> bool:
    try:
        if shape.is_placeholder and shape.placeholder_format.type is not None and shape.placeholder_format.type.name in TP.TITLE_PH:
            return True
    except Exception:
        pass
    return slot_role == "title"


def slide_text_blocks(prs: Any, slide: Any, index: int, resolver: Any) -> List[Dict[str, Any]]:
    """Text-bearing frames of a slide with abs box (EMU), paragraphs, runs and z-order position."""
    blocks: List[Dict[str, Any]] = []
    for z, (shape, path, _depth, tf) in enumerate(iter_shapes_abs(slide.shapes)):
        el = shape._element
        tag = localname(el)
        box = abs_box(shape, tf)
        if tag == "graphicFrame" and getattr(shape, "has_table", False) and shape.has_table:
            tbl = shape.table
            cw = [c.width for c in tbl.columns]
            rh = [r.height for r in tbl.rows]
            x0, y0 = (box[0], box[1]) if box else (0.0, 0.0)
            for ri, row in enumerate(tbl.rows):
                for ci, cell in enumerate(row.cells):
                    if cell.is_spanned:
                        continue
                    cx = x0 + sum(cw[:ci])
                    cy = y0 + sum(rh[:ri])
                    blocks.append(
                        {
                            "shape": shape,
                            "path": "%s@%d,%d" % (path, ri + 1, ci + 1),
                            "txbody": cell.text_frame._txBody,
                            "box": (cx, cy, cx + sum(cw[ci : ci + max(1, cell.span_width)]), cy + sum(rh[ri : ri + max(1, cell.span_height)])),
                            "z": z,
                            "cell": True,
                        }
                    )
            continue
        if tag != "sp":
            continue
        txb = el.find(P + "txBody")
        if txb is None:
            continue
        if box is None:
            try:
                base = shape._base_placeholder
                while base is not None and C.shape_bounds(base) is None:
                    base = base._base_placeholder
                bb = C.shape_bounds(base) if base is not None else None
                if bb is not None:
                    box = tuple(float(v) for v in bb)  # type: ignore[assignment]
            except Exception:
                box = None
        blocks.append({"shape": shape, "path": path, "txbody": txb, "box": box, "z": z, "cell": False})
    return blocks


def flat_text(txb: Any) -> str:
    return TP.txbody_text(txb).replace("\v", " ")


def is_text_shape(shape: Any) -> bool:
    try:
        return bool(shape.has_text_frame and shape.text_frame.text.strip())
    except Exception:
        return False


# --------------------------------------------------------------------------
# Individual checks
# --------------------------------------------------------------------------


def check_geometry(prs: Any, slide: Any, index: int, issues: Issues) -> None:
    sw, sh = prs.slide_width, prs.slide_height
    tol = int(0.02 * EMU)
    for shape, path, _d, tf in iter_shapes_abs(slide.shapes):
        if localname(shape._element) == "grpSp":
            continue
        box = abs_box(shape, tf)
        if box is None:
            continue
        l, t, r, b = box
        if not (l < -tol or t < -tol or r > sw + tol or b > sh + tol):
            continue
        content = is_text_shape(shape) or localname(shape._element) in ("pic", "graphicFrame")
        wholly = r <= 0 or b <= 0 or l >= sw or t >= sh
        over = max(-l, -t, r - sw, b - sh) / EMU
        if wholly:
            issues.add(index, shape_label(shape, path), "error" if content else "warning", "off-slide-hidden", "shape lies completely outside the slide (never visible)")
        else:
            issues.add(
                index,
                shape_label(shape, path),
                "error" if content else "warning",
                "off-slide",
                "shape extends %.2f in beyond the slide edge (left %.2f, top %.2f, right %.2f, bottom %.2f in; slide %.2f x %.2f in)"
                % (over, l / EMU, t / EMU, r / EMU, b / EMU, sw / EMU, sh / EMU),
            )


def check_overflow(slide: Any, index: int, blocks: List[Dict[str, Any]], issues: Issues, resolver: Any, slide_h_emu: int = 0) -> None:
    for blk in blocks:
        txb, box, shape = blk["txbody"], blk["box"], blk["shape"]
        if box is None or not flat_text(txb).strip():
            continue
        if TP._shape_rotated(shape):
            continue
        bp = txb.find(A + "bodyPr")
        if bp is not None:
            if bp.get("vert") not in (None, "horz"):
                continue
            if bp.find(A + "normAutofit") is not None:
                continue
        grows = bp is not None and bp.find(A + "spAutoFit") is not None
        if not TP.wraps(txb):
            continue
        li, ti, ri, bi = TP.body_insets_in(txb)
        width = (box[2] - box[0]) / EMU - li - ri
        avail = (box[3] - box[1]) / EMU - ti - bi
        if width <= 0.05 or avail <= 0:
            continue
        spacing = TP.line_spacing_factor(txb)
        needed = 0.0
        for p in TP.txbody_paragraphs(txb):
            text = TP.para_text(p).replace("\v", "\n")
            runs = TP.text_runs(p)
            sizes = [resolver.run_size(slide, shape, p, r) for r in runs] or [resolver.run_size(slide, shape, p, None)]
            pt = max(sizes)
            if not text:
                needed += pt * spacing / 72.0
                continue
            _n, h = C.measure_text(text, width, pt, spacing)
            needed += h
            ppr = p.find(A + "pPr")
            if ppr is not None:
                for tag in ("spcBef", "spcAft"):
                    pts = ppr.find("%s%s/%sspcPts" % (A, tag, A))
                    if pts is not None and pts.get("val"):
                        needed += int(pts.get("val")) / 100.0 / 72.0
        if grows:
            # resize-shape autofit: the box grows with its text, so the limit is the slide bottom
            room = (slide_h_emu - box[1]) / EMU - ti - bi if slide_h_emu else avail
            if needed > room + 0.05:
                issues.add(
                    index,
                    shape_label(shape, blk["path"]),
                    "error",
                    "text-overflow",
                    "text needs about %.2f in of height but only %.2f in remain below the box top (box grows with its text and runs off the slide)" % (needed, room),
                )
            continue
        if needed > avail * 1.25 + 0.1:
            sev = "error"
        elif needed > avail + 0.05:
            sev = "warning"
        else:
            continue
        issues.add(
            index,
            shape_label(shape, blk["path"]),
            sev,
            "text-overflow",
            "text needs about %.2f in of height but the box offers %.2f in (estimate)" % (needed, avail),
        )


def check_fonts_and_sizes(
    slide: Any,
    index: int,
    blocks: List[Dict[str, Any]],
    issues: Issues,
    resolver: Any,
    mins: Dict[str, float],
    fonts_used: Dict[str, Set[int]],
    title_ids: Set[str],
) -> None:
    sh_in = None
    for blk in blocks:
        txb, shape = blk["txbody"], blk["shape"]
        if not flat_text(txb).strip():
            continue
        is_title = blk["path"] in title_ids
        for p in TP.txbody_paragraphs(txb):
            ptext = TP.para_text(p).replace("\v", " ")
            if not ptext.strip():
                continue
            for r in TP.text_runs(p):
                t = r.find(A + "t")
                if t is None or not (t.text or "").strip():
                    continue
                size = resolver.run_size(slide, shape, p, r)
                fonts_used.setdefault(resolver.run_font(slide, shape, p, r), set()).add(index)
                label = shape_label(shape, blk["path"])
                if size < mins["caption"] - 0.01:
                    issues.add(index, label, "error", "font-too-small", "text at %.4g pt is below the caption minimum of %g pt" % (size, mins["caption"]))
                elif is_title and size < mins["title"] - 0.01:
                    issues.add(index, label, "error", "font-too-small", "title text at %.4g pt is below the title minimum of %g pt" % (size, mins["title"]))
                elif not is_title and size < mins["body"] - 0.01 and len(ptext.strip()) > 90:
                    issues.add(index, label, "error", "font-too-small", "body paragraph at %.4g pt is below the body minimum of %g pt (short captions may use %g pt)" % (size, mins["body"], mins["caption"]))


def check_leftovers(
    slide: Any,
    index: int,
    blocks: List[Dict[str, Any]],
    issues: Issues,
    template_sentences: Set[str],
    allow_footers: bool,
) -> None:
    foot_sev = "warning" if allow_footers else "error"
    for blk in blocks:
        shape, txb = blk["shape"], blk["txbody"]
        label = shape_label(shape, blk["path"])
        text = flat_text(txb)
        ph_type = None
        try:
            if shape.is_placeholder:
                ph_type = shape.placeholder_format.type.name if shape.placeholder_format.type is not None else "OBJECT"
        except Exception:
            ph_type = None
        if ph_type == "SLIDE_NUMBER":
            issues.add(index, label, foot_sev, "slide-number-field", "slide number placeholder on the slide")
            continue
        if ph_type == "DATE":
            issues.add(index, label, foot_sev, "date-field", "date placeholder on the slide")
            continue
        if ph_type == "FOOTER":
            issues.add(index, label, foot_sev, "footer-placeholder", "footer placeholder on the slide")
            continue
        if TP.has_field(txb, ("slidenum",)):
            issues.add(index, label, foot_sev, "slide-number-field", "text contains a slide number field")
        if TP.has_field(txb, ("datetime",)):
            issues.add(index, label, foot_sev, "date-field", "text contains a date field")
        if not text.strip():
            if ph_type is not None:
                issues.add(index, label, "error", "leftover-placeholder", "empty %s placeholder (shows prompt text when edited)" % ph_type.lower())
            continue
        for rx, what in PLACEHOLDER_PATTERNS:
            m = rx.search(text)
            if m:
                issues.add(index, label, "error", "leftover-placeholder", "%s found: %r" % (what, m.group(0)[:60]))
        if template_sentences:
            for line in text.split("\n"):
                if line.strip().lower() in template_sentences:
                    issues.add(index, label, "error", "leftover-template-text", "sentence copied from the template was left unchanged: %r" % line.strip()[:60])
    # empty picture placeholders
    for shape, path, _d, _t in iter_shapes_abs(slide.shapes):
        try:
            if localname(shape._element) == "sp" and shape.is_placeholder and shape.placeholder_format.type is not None and shape.placeholder_format.type.name == "PICTURE":
                issues.add(index, shape_label(shape, path), "error", "leftover-placeholder", "empty picture placeholder")
        except Exception:
            pass


def check_images(prs: Any, issues: Issues) -> None:
    sw, sh = prs.slide_width, prs.slide_height
    seen: Dict[str, List[Tuple[int, str, float]]] = {}
    for i, slide in enumerate(prs.slides, start=1):
        for shape, path, _d, tf in iter_shapes_abs(slide.shapes):
            if localname(shape._element) != "pic":
                continue
            try:
                blob = shape.image.blob
            except Exception:
                continue
            h = hashlib.sha1(blob).hexdigest()
            box = abs_box(shape, tf)
            frac = ((box[2] - box[0]) * (box[3] - box[1])) / float(sw * sh) if box and sw and sh else 0.0
            seen.setdefault(h, []).append((i, shape_label(shape, path), frac))
    for h, uses in seen.items():
        slides = sorted({u[0] for u in uses})
        if len(slides) < 2:
            continue
        big = max(u[2] for u in uses) > 0.03
        for slide_no, label, _f in uses[1:]:
            if slide_no == uses[0][0]:
                continue
            issues.add(
                slide_no,
                label,
                "error" if big else "warning",
                "image-reused",
                "same image (sha1 %s) is used on slides %s" % (h[:8], ", ".join(str(s) for s in slides)),
            )


def check_notes(prs: Any, issues: Issues) -> None:
    for i, slide in enumerate(prs.slides, start=1):
        text = ""
        if slide.has_notes_slide:
            try:
                text = slide.notes_slide.notes_text_frame.text
            except Exception:
                text = ""
        if not text.strip():
            issues.add(i, "", "warning", "no-notes", "slide has no speaker notes")


def _slide_background(resolver: Any, slide: Any) -> Optional[str]:
    """Solid slide background colour (slide, then layout, then master), or None."""
    for holder in (slide, slide.slide_layout, slide.slide_layout.slide_master):
        bg = holder._element.find(".//" + P + "bg")
        if bg is None:
            continue
        bgpr = bg.find(P + "bgPr")
        if bgpr is not None:
            fill = bgpr.find(A + "solidFill")
            if fill is not None:
                return resolver.color_of(slide, fill)
            return None  # gradient / picture background
        ref = bg.find(P + "bgRef")
        if ref is not None:
            clr = next((c for c in ref if localname(c) in ("srgbClr", "schemeClr", "sysClr", "prstClr")), None)
            return resolver.color_of(slide, clr) if clr is not None else None
    return None


def _shape_fill(resolver: Any, slide: Any, shape: Any) -> Tuple[str, Optional[str]]:
    """("solid", hex) / ("none", None) / ("unknown", None) for a shape's own fill."""
    el = shape._element
    sppr = el.find(P + "spPr")
    if sppr is None:
        return "none", None
    if sppr.find(A + "noFill") is not None:
        return "none", None
    fill = sppr.find(A + "solidFill")
    if fill is not None:
        col = resolver.color_of(slide, fill)
        return ("solid", col) if col else ("unknown", None)
    if sppr.find(A + "gradFill") is not None or sppr.find(A + "blipFill") is not None or sppr.find(A + "pattFill") is not None:
        return "unknown", None
    style = el.find(P + "style")
    if style is not None:
        ref = style.find(A + "fillRef")
        if ref is not None and ref.get("idx") not in (None, "0"):
            clr = next((c for c in ref if localname(c) in ("srgbClr", "schemeClr")), None)
            col = resolver.color_of(slide, clr) if clr is not None else None
            return ("solid", col) if col else ("unknown", None)
    return "none", None


def check_contrast(slide: Any, index: int, blocks: List[Dict[str, Any]], issues: Issues, resolver: Any) -> None:
    bg = _slide_background(resolver, slide)
    shapes = list(iter_shapes_abs(slide.shapes))
    for blk in blocks:
        box, shape, txb = blk["box"], blk["shape"], blk["txbody"]
        if box is None or not flat_text(txb).strip():
            continue
        cx, cy = (box[0] + box[2]) / 2, (box[1] + box[3]) / 2
        # background: own solid fill, else topmost earlier shape containing the centre, else slide background
        kind, colour = ("none", None)
        if not blk["cell"]:
            kind, colour = _shape_fill(resolver, slide, shape)
        if blk["cell"]:
            kind = "unknown"  # table styles define cell fills
        if kind == "none":
            kind, colour = "bg", bg
            for other, _opath, _od, otf in shapes:
                if other is shape:
                    break
                ob = abs_box(other, otf)
                tag = localname(other._element)
                if ob is None or tag in ("grpSp", "cxnSp"):
                    continue
                if not (ob[0] <= cx <= ob[2] and ob[1] <= cy <= ob[3]):
                    continue
                if tag in ("pic", "graphicFrame"):
                    kind, colour = "unknown", None
                    continue
                k2, c2 = _shape_fill(resolver, slide, other)
                if k2 == "solid":
                    kind, colour = "solid", c2
                elif k2 == "unknown":
                    kind, colour = "unknown", None
        if not colour or kind == "unknown":
            continue
        reported = False
        for p in TP.txbody_paragraphs(txb):
            for r in TP.text_runs(p):
                t = r.find(A + "t")
                if t is None or not (t.text or "").strip() or reported:
                    continue
                fg = resolver.run_color(slide, shape, p, r)
                if not fg:
                    continue
                size = resolver.run_size(slide, shape, p, r)
                rpr = r.find(A + "rPr")
                bold = rpr is not None and rpr.get("b") == "1"
                need = 3.0 if (size >= 24 or (bold and size >= 18.66)) else 4.5
                ratio = C.contrast_ratio(fg, colour)
                if ratio < need:
                    reported = True
                    issues.add(
                        index,
                        shape_label(shape, blk["path"]),
                        "warning",
                        "low-contrast",
                        "text %s on %s has contrast %.2f:1 (needs %.1f:1 at %.4g pt)" % (fg, colour, ratio, need, size),
                    )


def layout_signature(slide: Any) -> str:
    parts = [slide.slide_layout.name]
    items = []
    for shape, _p, depth, tf in iter_shapes_abs(slide.shapes):
        if localname(shape._element) == "grpSp":
            continue
        box = abs_box(shape, tf)
        kind = localname(shape._element)
        if box is None:
            items.append((kind, -1, -1, -1, -1))
        else:
            q = EMU * 0.5
            items.append((kind, round(box[0] / q), round(box[1] / q), round((box[2] - box[0]) / q), round((box[3] - box[1]) / q)))
    items.sort()
    return hashlib.sha1((parts[0] + repr(items)).encode("utf-8")).hexdigest()


def check_layout_runs(prs: Any, issues: Issues) -> None:
    sigs = [layout_signature(s) for s in prs.slides]
    run = 1
    for i in range(1, len(sigs)):
        run = run + 1 if sigs[i] == sigs[i - 1] else 1
        if run == 3:
            issues.add(i + 1, "", "warning", "layout-repeat", "slides %d-%d use the same layout; vary the layout after two in a row" % (i - 1, i + 1))
        elif run > 3:
            issues.add(i + 1, "", "warning", "layout-repeat", "slide %d continues a run of %d slides with the same layout" % (i + 1, run))


def check_fonts_global(fonts_used: Dict[str, Set[int]], issues: Issues) -> None:
    if len(fonts_used) > 3:
        issues.add(0, "", "warning", "too-many-fonts", "%d distinct fonts are used (%s); keep to 3 or fewer" % (len(fonts_used), ", ".join(sorted(fonts_used))))
    for name, slides in sorted(fonts_used.items()):
        if not C.font_available(name, assume_office=True):
            issues.add(0, "", "warning", "font-missing", "font %r is not installed on this machine (used on slide(s) %s); it will be substituted when rendered" % (name, ", ".join(str(s) for s in sorted(slides)[:8])))


def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^\w\s]", "", text.lower())).strip()


def spec_titles(slide_spec: Dict[str, Any]) -> List[str]:
    t = slide_spec.get("title")
    if isinstance(t, str):
        return [t]
    if isinstance(t, dict):
        return [v for v in (t.get("full"), t.get("short")) if isinstance(v, str) and v.strip()]
    return []


def check_spec(prs: Any, spec: Dict[str, Any], issues: Issues) -> None:
    sspecs = spec.get("slides") or []
    n = len(prs.slides)
    if len(sspecs) != n:
        issues.add(0, "", "error", "slide-count", "deck has %d slide(s) but the spec lists %d" % (n, len(sspecs)))
    for i, (slide, ss) in enumerate(zip(prs.slides, sspecs), start=1):
        titles = spec_titles(ss)
        if not titles:
            continue
        blob = _norm(" ".join(shape.text_frame.text for shape, _p, _d, _t in iter_shapes_abs(slide.shapes) if getattr(shape, "has_text_frame", False) and shape.has_text_frame))
        if not any(_norm(t) and _norm(t) in blob for t in titles):
            issues.add(i, "", "error", "title-missing", "none of the spec titles is on the slide: %s" % " | ".join(repr(t) for t in titles))


def template_sentences_of(path: str) -> Set[str]:
    """Lower-cased text lines of 15+ characters found anywhere in a template deck."""
    out: Set[str] = set()
    prs = Presentation(path)
    for slide in prs.slides:
        for shape, _p, _d, _t in iter_shapes_abs(slide.shapes):
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(r.text for r in para.runs).strip()
                    if len(line) >= 15:
                        out.add(line.lower())
    return out


# --------------------------------------------------------------------------
# Driver
# --------------------------------------------------------------------------


def run_checks(
    path: str,
    spec: Optional[Dict[str, Any]] = None,
    mins: Optional[Dict[str, float]] = None,
    template: Optional[str] = None,
    allow_footers: bool = False,
) -> Dict[str, Any]:
    issues = Issues()
    # 1. package validity
    try:
        report = ooxml.validate(path)
    except (OSError, zipfile.BadZipFile, ValueError, KeyError) as exc:
        issues.add(0, "", "error", "package-unreadable", "cannot read the package: %s" % exc)
        return {"status": "blocked", "file": path, "issues": issues.items}
    except Exception as exc:  # corrupt XML etc.
        issues.add(0, "", "error", "package-unreadable", "cannot read the package: %s" % exc)
        return {"status": "blocked", "file": path, "issues": issues.items}
    for pr in report.problems:
        issues.add(0, pr["part"], "error", "package-" + pr["code"], pr["message"])
    if report.problems:
        return {"status": "blocked", "file": path, "issues": issues.items}
    try:
        prs = Presentation(path)
    except Exception as exc:
        issues.add(0, "", "error", "package-unreadable", "python-pptx cannot open the deck: %s" % exc)
        return {"status": "blocked", "file": path, "issues": issues.items}

    limits = C.typography_limits().get("minPt", {})
    m = {"body": float(limits.get("body", 18)), "caption": float(limits.get("caption", 12)), "title": float(limits.get("title", 28))}
    for k, v in (mins or {}).items():
        if v is not None:
            m[k] = float(v)

    resolver = TP.StyleResolver(prs)
    tpl_sentences: Set[str] = template_sentences_of(template) if template else set()
    fonts_used: Dict[str, Set[int]] = {}
    for i, slide in enumerate(prs.slides, start=1):
        blocks = slide_text_blocks(prs, slide, i, resolver)
        # slide title: a title placeholder, else the largest non-numeric text in the top quarter
        title_paths: Set[str] = set()
        best = None
        for blk in blocks:
            if blk["cell"] or not flat_text(blk["txbody"]).strip():
                continue
            if is_title_shape(blk["shape"], None):
                title_paths.add(blk["path"])
            box = blk["box"]
            if box is None or box[1] > 0.25 * prs.slide_height:
                continue
            txt = flat_text(blk["txbody"]).strip()
            if len(txt) > 140 or TP.NUMERIC_RE.match(txt):
                continue
            sz = max([resolver.run_size(slide, blk["shape"], p, r) for p in TP.txbody_paragraphs(blk["txbody"]) for r in TP.text_runs(p)] or [0])
            if best is None or sz > best[0] + 0.01 or (abs(sz - best[0]) <= 0.01 and box[1] < best[2]):
                best = (sz, blk["path"], box[1])
        # a fallback title must also be the largest text on the slide (otherwise a small label would be taken for it)
        slide_max = max(
            [resolver.run_size(slide, b["shape"], p, r) for b in blocks if not b["cell"] and flat_text(b["txbody"]).strip() for p in TP.txbody_paragraphs(b["txbody"]) for r in TP.text_runs(p)] or [0]
        )
        if not title_paths and best is not None and best[0] >= m["body"] and best[0] >= slide_max - 0.01:
            title_paths.add(best[1])
        check_geometry(prs, slide, i, issues)
        check_overflow(slide, i, blocks, issues, resolver, prs.slide_height)
        check_fonts_and_sizes(slide, i, blocks, issues, resolver, m, fonts_used, title_paths)
        check_leftovers(slide, i, blocks, issues, tpl_sentences, allow_footers)
        check_contrast(slide, i, blocks, issues, resolver)
    check_images(prs, issues)
    check_notes(prs, issues)
    check_layout_runs(prs, issues)
    check_fonts_global(fonts_used, issues)
    if spec is not None:
        check_spec(prs, spec, issues)

    errors = [x for x in issues.items if x["severity"] == "error"]
    order = {"error": 0, "warning": 1}
    issues.items.sort(key=lambda x: (x["slide"], order[x["severity"]], x["code"]))
    return {
        "status": "revise" if errors else "pass",
        "file": path,
        "slideCount": len(prs.slides),
        "counts": {"errors": len(errors), "warnings": len(issues.items) - len(errors)},
        "issues": issues.items,
    }


def main(argv: Optional[List[str]] = None) -> int:
    ap = argparse.ArgumentParser(
        description="QA a .pptx: off-slide shapes, text overflow, small fonts, leftover placeholders, footers, "
        "fonts, reused images, notes, contrast, layout repeats, and (with --spec) slide count and titles. "
        "Prints JSON {status, issues}. Exit 0 pass, 1 revise, 2 blocked.",
    )
    ap.add_argument("input", help="deck .pptx")
    ap.add_argument("--spec", help="deck spec JSON the deck was built from (checks slide count and titles)")
    ap.add_argument("--template", help="template .pptx; sentences copied verbatim from it count as leftover sample text")
    ap.add_argument("--json", action="store_true", help="print JSON only (default prints a readable list, then the status)")
    ap.add_argument("--min-body", type=float, help="minimum body font size in pt (default x-typography.minPt.body)")
    ap.add_argument("--min-caption", type=float, help="minimum caption font size in pt (default x-typography.minPt.caption)")
    ap.add_argument("--min-title", type=float, help="minimum title font size in pt (default x-typography.minPt.title)")
    ap.add_argument("--allow-footers", action="store_true", help="report slide number / date / footer fields as warnings instead of errors")
    args = ap.parse_args(argv)

    spec = None
    if args.spec:
        try:
            spec = C.load_json(args.spec)
        except (OSError, ValueError) as exc:
            print("error: cannot read spec %s: %s" % (args.spec, exc), file=sys.stderr)
            return EXIT_BLOCKED
    if not Path(args.input).is_file():
        result = {"status": "blocked", "file": args.input, "issues": [{"slide": 0, "shape": "", "severity": "error", "code": "package-unreadable", "message": "file not found: %s" % args.input}]}
    else:
        try:
            result = run_checks(
                args.input,
                spec,
                {"body": args.min_body, "caption": args.min_caption, "title": args.min_title},
                args.template,
                args.allow_footers,
            )
        except Exception as exc:  # last resort: report instead of a traceback
            result = {"status": "blocked", "file": args.input, "issues": [{"slide": 0, "shape": "", "severity": "error", "code": "package-unreadable", "message": "QA could not run: %s" % exc}]}
    if args.json:
        print(json.dumps(result, indent=2, ensure_ascii=False))
    else:
        for it in result["issues"]:
            print("%-7s slide %-3s %-22s %s%s" % (it["severity"].upper(), it["slide"] or "-", it["code"], it["message"], "  [%s]" % it["shape"] if it["shape"] else ""))
        print("status: %s (%s)" % (result["status"], "%(errors)d error(s), %(warnings)d warning(s)" % result["counts"] if "counts" in result else "%d problem(s)" % len(result["issues"])))
    return {"pass": EXIT_PASS, "revise": EXIT_REVISE, "blocked": EXIT_BLOCKED}[result["status"]]


if __name__ == "__main__":
    sys.exit(main())
