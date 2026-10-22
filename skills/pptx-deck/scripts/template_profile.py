#!/usr/bin/env python3
"""Profile a user-supplied template deck: page roles, text slots, picture slots, type scale.

Usage:
    template_profile.py TEMPLATE.pptx -o profile.json
    template_profile.py TEMPLATE.pptx            # JSON to stdout

The profile describes every slide of the template so that a deck can be
assembled by *selecting, ordering and duplicating* its slides and then filling
text only (see fill_template.py):

* ``slides[i].role`` - guessed page role: ``cover``, ``agenda``, ``section``
  (transition), ``content`` or ``closing``; ``contentType`` refines content
  slides (text, columns, grid, big-numbers, image-text, chart, table, process,
  quote). The guess uses position, text, type scale and shape mix; it is a
  heuristic, ``roleReason`` says why.
* ``slides[i].slots`` - every text-bearing shape, including shapes inside groups
  and table cells, each with a stable slot id ``s<slide>:<shape id path>`` (for
  example ``s4:6`` or ``s6:18/2``; a table cell is ``s7:12@2,3`` with 1-based
  row,column), its name, placeholder type, sample text, bounds in inches,
  font sizes in points, an estimated capacity in characters at the current
  font size (soft: about 20 percent headroom over the box estimate and never
  below the sample text) and an ``editable`` flag. Page numbers, dates, tiny
  text and locked shapes are not editable. ``persistent`` marks running
  header/footer labels that repeat on most slides and are kept when unfilled.
* ``slides[i].pictures`` - picture slots (replaceable image frames) with
  pixel size, role (photo, logo, background) and a ``replaceable`` flag.
* ``summary`` - deck-level type scale, theme palette, fonts and the slide
  indices recommended per role.

Slide numbers are 1-based. The module also provides ``StyleResolver`` (inherited
font size, font and colour lookup through layout, master and theme) which
qa_check.py and fill_template.py import.

Exit status: 0 ok, 2 unreadable input.
"""

from __future__ import annotations

import argparse
import colorsys
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Dict, Iterator, List, Optional, Tuple

sys.path.insert(0, str(Path(__file__).resolve().parent))
import _common as C  # noqa: E402

C.require_pptx()

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.opc.constants import RELATIONSHIP_TYPE as RT  # noqa: E402

from inventory import abs_box, iter_shapes_abs, localname  # noqa: E402

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A = "{%s}" % NS_A
P = "{%s}" % NS_P
EMU = C.EMU_PER_INCH

#: Headroom over the geometric estimate when stating a slot's capacity.
CAPACITY_HEADROOM = 1.2

TITLE_PH = {"TITLE", "CENTER_TITLE", "VERTICAL_TITLE"}
BODY_PH = {"BODY", "OBJECT", "SUBTITLE", "VERTICAL_BODY", "VERTICAL_OBJECT"}
FOOTER_PH = {"SLIDE_NUMBER": "page-number", "DATE": "date", "FOOTER": "footer"}


def _xp(el: Any, expr: str) -> List[Any]:
    return etree._Element.xpath(el, expr, namespaces={"a": NS_A, "p": NS_P, "r": NS_R})


# --------------------------------------------------------------------------
# Theme, colours, inherited text properties
# --------------------------------------------------------------------------

_SCHEME_TO_THEME = {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}


def _rgb_to_hex(r: float, g: float, b: float) -> str:
    return "#%02X%02X%02X" % (max(0, min(255, round(r))), max(0, min(255, round(g))), max(0, min(255, round(b))))


def apply_color_mods(hex_color: str, clr_el: Any) -> Optional[str]:
    """Apply lumMod/lumOff/tint/shade children of a colour element; None when alpha or other mods make it unknowable."""
    r, g, b = C.hex_to_tuple(hex_color)
    for mod in clr_el:
        name = localname(mod)
        try:
            val = int(mod.get("val", "0")) / 100000.0
        except ValueError:
            return None
        if name == "alpha":
            if val < 0.999:
                return None
        elif name in ("lumMod", "lumOff"):
            h, l, s = colorsys.rgb_to_hls(r / 255.0, g / 255.0, b / 255.0)
            l = l * val if name == "lumMod" else l + val
            l = max(0.0, min(1.0, l))
            rr, gg, bb = colorsys.hls_to_rgb(h, l, s)
            r, g, b = rr * 255, gg * 255, bb * 255
        elif name == "tint":
            r, g, b = r * val + 255 * (1 - val), g * val + 255 * (1 - val), b * val + 255 * (1 - val)
        elif name == "shade":
            r, g, b = r * val, g * val, b * val
        elif name in ("satMod", "satOff", "hueMod", "hueOff", "hue", "sat", "lum", "red", "green", "blue", "gamma", "invGamma", "comp", "inv", "gray"):
            return None
    return _rgb_to_hex(r, g, b)


class Theme:
    """Colours and fonts of one theme part."""

    def __init__(self, blob: Optional[bytes]) -> None:
        self.colors: Dict[str, str] = {}
        self.major = ""
        self.minor = ""
        if not blob:
            return
        try:
            root = etree.fromstring(blob)
        except etree.XMLSyntaxError:
            return
        scheme = root.find(".//" + A + "clrScheme")
        if scheme is not None:
            for ch in scheme:
                name = localname(ch)
                inner = next(iter(ch), None)
                if inner is None:
                    continue
                val = inner.get("val") if localname(inner) == "srgbClr" else inner.get("lastClr")
                if val and re.fullmatch(r"[0-9A-Fa-f]{6}", val):
                    self.colors[name] = "#" + val.upper()
        for tag, attr in (("majorFont", "major"), ("minorFont", "minor")):
            el = root.find(".//%s%s/%slatin" % (A, tag, A))
            if el is not None and el.get("typeface"):
                setattr(self, attr, el.get("typeface"))


class StyleResolver:
    """Look up inherited text size, font, colour and fills for shapes of an open presentation."""

    def __init__(self, prs: Any) -> None:
        self.prs = prs
        self._themes: Dict[Any, Theme] = {}
        pres_el = prs.part._element
        self._default_text_style = pres_el.find(P + "defaultTextStyle")

    # -- theme / master ----------------------------------------------------
    def master_of(self, slide: Any) -> Any:
        return slide.slide_layout.slide_master

    def theme(self, slide: Any) -> Theme:
        master = self.master_of(slide)
        key = master.part
        if key not in self._themes:
            blob = None
            try:
                blob = master.part.part_related_by(RT.THEME).blob
            except Exception:
                pass
            self._themes[key] = Theme(blob)
        return self._themes[key]

    def scheme_color(self, slide: Any, name: str) -> Optional[str]:
        """Resolve a scheme colour name (bg1, tx2, accent3 ...) through the master colour map."""
        th = self.theme(slide)
        master = self.master_of(slide)
        clr_map = master._element.find(P + "clrMap")
        target = name
        if clr_map is not None and clr_map.get(name):
            target = clr_map.get(name)
        target = _SCHEME_TO_THEME.get(target, target)
        return th.colors.get(target)

    def color_of(self, slide: Any, fill_el: Any) -> Optional[str]:
        """Hex colour of an element holding a colour child (``a:solidFill``, ``a:srgbClr`` ...); None if unknowable."""
        if fill_el is None:
            return None
        if localname(fill_el) in ("srgbClr", "schemeClr", "sysClr", "prstClr"):
            clr = fill_el
        else:
            clr = next((c for c in fill_el if localname(c) in ("srgbClr", "schemeClr", "sysClr", "prstClr")), None)
        if clr is None:
            return None
        kind = localname(clr)
        base: Optional[str] = None
        if kind == "srgbClr":
            base = "#" + clr.get("val", "").upper()
        elif kind == "sysClr":
            base = "#" + clr.get("lastClr", "").upper() if clr.get("lastClr") else None
            if base is None:
                base = {"windowText": "#000000", "window": "#FFFFFF"}.get(clr.get("val", ""))
        elif kind == "prstClr":
            base = {"black": "#000000", "white": "#FFFFFF"}.get(clr.get("val", ""))
        elif kind == "schemeClr":
            base = self.scheme_color(slide, clr.get("val", ""))
        if not base or not re.fullmatch(r"#[0-9A-F]{6}", base):
            return None
        return apply_color_mods(base, clr)

    # -- inheritance chain -------------------------------------------------
    def _placeholder_type(self, shape: Any) -> Optional[str]:
        try:
            if shape.is_placeholder:
                t = shape.placeholder_format.type
                return t.name if t is not None else "OBJECT"
        except Exception:
            return None
        return None

    def defrprs(self, slide: Any, shape: Any, level: int = 0) -> List[Any]:
        """``a:defRPr`` / ``a:lvlNpPr`` candidates for paragraph *level* in priority order."""
        out: List[Any] = []
        tag = "%slvl%dpPr" % (A, level + 1)

        def add(container: Any) -> None:
            if container is None:
                return
            pPr = container.find(tag)
            if pPr is not None:
                out.append(pPr)
                d = pPr.find(A + "defRPr")
                if d is not None:
                    out.append(d)

        el = getattr(shape, "_element", None)
        if el is not None and localname(el) == "sp":
            tb = el.find(P + "txBody")
            if tb is not None:
                add(tb.find(A + "lstStyle"))
        ph_type = self._placeholder_type(shape)
        if ph_type is not None:
            base = shape
            for _ in range(3):
                try:
                    base = base._base_placeholder
                except Exception:
                    base = None
                if base is None:
                    break
                tb = base._element.find(P + "txBody")
                if tb is not None:
                    add(tb.find(A + "lstStyle"))
            master = self.master_of(slide)
            tx = master._element.find(P + "txStyles")
            if tx is not None:
                if ph_type in TITLE_PH:
                    add(tx.find(P + "titleStyle"))
                elif ph_type in BODY_PH:
                    add(tx.find(P + "bodyStyle"))
                else:
                    add(tx.find(P + "otherStyle"))
        add(self._default_text_style)
        return out

    @staticmethod
    def para_level(p_el: Any) -> int:
        ppr = p_el.find(A + "pPr")
        try:
            return max(0, min(8, int(ppr.get("lvl", "0")))) if ppr is not None else 0
        except ValueError:
            return 0

    def font_scale(self, shape: Any) -> float:
        el = getattr(shape, "_element", None)
        if el is None:
            return 1.0
        na = el.find(".//" + A + "bodyPr/" + A + "normAutofit")
        if na is not None and na.get("fontScale"):
            try:
                return int(na.get("fontScale")) / 100000.0
            except ValueError:
                return 1.0
        return 1.0

    def run_size(self, slide: Any, shape: Any, p_el: Any, r_el: Any, scaled: bool = True) -> float:
        """Effective font size in points of a run (autofit fontScale applied when *scaled*)."""
        sz = None
        rpr = r_el.find(A + "rPr") if r_el is not None else None
        if rpr is not None and rpr.get("sz"):
            sz = rpr.get("sz")
        if sz is None:
            for d in self.defrprs(slide, shape, self.para_level(p_el)):
                v = d.get("sz") if localname(d) == "defRPr" else None
                if v:
                    sz = v
                    break
        try:
            size = int(sz) / 100.0 if sz else 18.0
        except ValueError:
            size = 18.0
        return size * self.font_scale(shape) if scaled else size

    def run_font(self, slide: Any, shape: Any, p_el: Any, r_el: Any) -> str:
        """Effective Latin font family of a run, theme fonts resolved."""
        th = self.theme(slide)
        face = None
        rpr = r_el.find(A + "rPr") if r_el is not None else None
        if rpr is not None:
            lat = rpr.find(A + "latin")
            if lat is not None:
                face = lat.get("typeface")
        if not face:
            for d in self.defrprs(slide, shape, self.para_level(p_el)):
                if localname(d) != "defRPr":
                    continue
                lat = d.find(A + "latin")
                if lat is not None and lat.get("typeface"):
                    face = lat.get("typeface")
                    break
        if not face or face == "+mn-lt":
            return th.minor or "Calibri"
        if face == "+mj-lt":
            return th.major or th.minor or "Calibri"
        return face

    def run_color(self, slide: Any, shape: Any, p_el: Any, r_el: Any) -> Optional[str]:
        """Explicit or inherited solid text colour of a run, or None when not determinable."""
        rpr = r_el.find(A + "rPr") if r_el is not None else None
        if rpr is not None:
            fill = rpr.find(A + "solidFill")
            if fill is not None:
                return self.color_of(slide, fill)
            if rpr.find(A + "gradFill") is not None:
                return None
        for d in self.defrprs(slide, shape, self.para_level(p_el)):
            if localname(d) != "defRPr":
                continue
            fill = d.find(A + "solidFill")
            if fill is not None:
                return self.color_of(slide, fill)
        return None


# --------------------------------------------------------------------------
# Text frames
# --------------------------------------------------------------------------


def text_runs(p_el: Any) -> List[Any]:
    """a:r and a:fld children (text-bearing) of a paragraph element."""
    return [c for c in p_el if localname(c) in ("r", "fld")]


def para_text(p_el: Any) -> str:
    parts = []
    for c in p_el:
        n = localname(c)
        if n in ("r", "fld"):
            t = c.find(A + "t")
            parts.append(t.text or "" if t is not None else "")
        elif n == "br":
            parts.append("\v")
    return "".join(parts)


def txbody_paragraphs(txbody: Any) -> List[Any]:
    return [c for c in txbody if localname(c) == "p"]


def txbody_text(txbody: Any) -> str:
    return "\n".join(para_text(p) for p in txbody_paragraphs(txbody))


def has_field(txbody: Any, kinds: Tuple[str, ...]) -> bool:
    for fld in txbody.iter(A + "fld"):
        t = (fld.get("type") or "").lower()
        if any(t.startswith(k) for k in kinds):
            return True
    return False


def body_insets_in(txbody: Any, is_cell: bool = False) -> Tuple[float, float, float, float]:
    bp = txbody.find(A + "bodyPr")

    def val(name: str, default: int) -> float:
        raw = bp.get(name) if bp is not None else None
        return (int(raw) if raw is not None else default) / EMU

    return val("lIns", 91440), val("tIns", 45720), val("rIns", 91440), val("bIns", 45720)


def wraps(txbody: Any) -> bool:
    bp = txbody.find(A + "bodyPr")
    return not (bp is not None and bp.get("wrap") == "none")


def dominant_size(resolver: StyleResolver, slide: Any, shape: Any, txbody: Any) -> Tuple[float, List[float]]:
    """(character-weighted dominant font size, sorted distinct sizes) of a text body."""
    weights: Counter = Counter()
    sizes = set()
    first = None
    for p in txbody_paragraphs(txbody):
        runs = text_runs(p)
        for r in runs:
            t = r.find(A + "t")
            n = len(t.text or "") if t is not None else 0
            sz = round(resolver.run_size(slide, shape, p, r), 2)
            if first is None:
                first = sz
            if n:
                weights[sz] += n
                sizes.add(sz)
        if not runs and first is None:
            first = round(resolver.run_size(slide, shape, p, None), 2)
    if weights:
        dom = max(weights.items(), key=lambda kv: (kv[1], kv[0]))[0]
    else:
        dom = first if first is not None else 18.0
        sizes.add(dom)
    return dom, sorted(sizes)


def line_spacing_factor(txbody: Any) -> float:
    for p in txbody_paragraphs(txbody):
        ppr = p.find(A + "pPr")
        if ppr is not None:
            pct = ppr.find("%slnSpc/%sspcPct" % (A, A))
            if pct is not None and pct.get("val"):
                try:
                    return max(0.6, min(3.0, int(pct.get("val")) / 100000.0)) * C.DEFAULT_LINE_SPACING
                except ValueError:
                    pass
    return C.DEFAULT_LINE_SPACING


def capacity_for(width_in: float, height_in: float, font_pt: float, txbody: Any, sample: str) -> Dict[str, int]:
    li, ti, ri, bi = body_insets_in(txbody)
    w = max(0.0, width_in - li - ri)
    h = max(0.0, height_in - ti - bi)
    cap = C.estimate_text_capacity(w, h, font_pt, line_spacing=line_spacing_factor(txbody))
    base = max(cap.max_chars, len(sample.replace("\n", "")))
    return {
        "charsPerLine": cap.chars_per_line,
        "lines": cap.max_lines,
        "chars": cap.max_chars,
        "softChars": int(base * CAPACITY_HEADROOM),
    }


# --------------------------------------------------------------------------
# Slot discovery
# --------------------------------------------------------------------------

NUMERIC_RE = re.compile(r"^[\s$€£¥+\-–~<>]*[\d][\d.,\s]*[%xX×kKmMbB]?\+?$")
INDEX_RE = re.compile(r"^\s*0?\d{1,2}\s*[.)]?\s*$")


def slot_id(slide_index: int, path: str, cell: Optional[Tuple[int, int]] = None) -> str:
    sid = "s%d:%s" % (slide_index, path)
    if cell is not None:
        sid += "@%d,%d" % (cell[0] + 1, cell[1] + 1)
    return sid


def parse_slot_id(sid: str) -> Tuple[Optional[int], str, Optional[Tuple[int, int]]]:
    """Split ``s4:6/2@1,3`` into (slide index or None, shape path, (row, col) 0-based or None).

    A bare path (``6``, ``6/2``) is accepted and yields a None slide index.
    """
    m = re.fullmatch(r"(?:s?(\d+):)?([0-9]+(?:/[0-9]+)*)(?:@(\d+),(\d+))?", sid.strip())
    if not m:
        raise ValueError("cannot read slot id %r (expected like s4:6, s4:6/2 or s4:12@2,3)" % sid)
    idx = int(m.group(1)) if m.group(1) else None
    cell = (int(m.group(3)) - 1, int(m.group(4)) - 1) if m.group(3) else None
    if cell and (cell[0] < 0 or cell[1] < 0):
        raise ValueError("slot id %r: table cells are 1-based" % sid)
    return idx, m.group(2), cell


def _bounds_dict(box: Tuple[float, float, float, float]) -> Dict[str, float]:
    l, t, r, b = box
    return {"left": round(l / EMU, 3), "top": round(t / EMU, 3), "width": round((r - l) / EMU, 3), "height": round((b - t) / EMU, 3)}


def _is_locked(shape: Any) -> bool:
    el = shape._element
    for lock in el.iter(A + "spLocks"):
        if lock.get("noTextEdit") == "1":
            return True
    return False


def _shape_rotated(shape: Any) -> bool:
    try:
        return abs(float(shape.rotation)) > 1.0
    except Exception:
        return False


def _txbody_of_shape(shape: Any) -> Optional[Any]:
    el = shape._element
    if localname(el) == "sp":
        return el.find(P + "txBody")
    return None


def collect_raw_slots(prs: Any, slide: Any, index: int, resolver: StyleResolver) -> List[Dict[str, Any]]:
    """Text slots of one slide, before deck-level role assignment."""
    sw, sh = prs.slide_width, prs.slide_height
    out: List[Dict[str, Any]] = []
    for shape, path, depth, tf in iter_shapes_abs(slide.shapes):
        el = shape._element
        tag = localname(el)
        box = abs_box(shape, tf)
        if box is None:
            # inherited geometry: take the layout / master placeholder box
            try:
                base = shape._base_placeholder
                while base is not None and C.shape_bounds(base) is None:
                    base = base._base_placeholder
                bb = C.shape_bounds(base) if base is not None else None
                if bb is not None:
                    box = tuple(float(v) for v in bb)  # type: ignore[assignment]
            except Exception:
                box = None
        if tag == "graphicFrame" and getattr(shape, "has_table", False) and shape.has_table:
            tbl = shape.table
            col_w = [c.width for c in tbl.columns]
            row_h = [r.height for r in tbl.rows]
            x0 = box[0] if box else 0.0
            y0 = box[1] if box else 0.0
            for ri, row in enumerate(tbl.rows):
                for ci, cell in enumerate(row.cells):
                    if cell.is_spanned:
                        continue
                    txb = cell.text_frame._txBody
                    cx = x0 + sum(col_w[:ci])
                    cy = y0 + sum(row_h[:ri])
                    wid = sum(col_w[ci : ci + max(1, cell.span_width)])
                    hei = sum(row_h[ri : ri + max(1, cell.span_height)])
                    out.append(
                        {
                            "shape": shape,
                            "txbody": txb,
                            "path": path,
                            "cell": (ri, ci),
                            "kind": "table-cell",
                            "box": (cx, cy, cx + wid, cy + hei),
                            "placeholder": None,
                            "name": shape.name,
                            "tableHeader": ri == 0,
                        }
                    )
            continue
        txb = _txbody_of_shape(shape)
        if txb is None:
            continue
        ph_type = resolver._placeholder_type(shape)
        text = txbody_text(txb)
        if not text.strip() and ph_type is None and not has_field(txb, ("slidenum", "datetime")):
            continue  # decorative shape without text
        out.append(
            {
                "shape": shape,
                "txbody": txb,
                "path": path,
                "cell": None,
                "kind": "text",
                "box": box,
                "placeholder": ph_type,
                "name": shape.name,
                "tableHeader": False,
            }
        )
    return out


def _slide_title_guess(slots: List[Dict[str, Any]]) -> str:
    for s in slots:
        if s["role"] == "title" and s["sample"].strip():
            return s["sample"].strip().split("\n")[0]
    return ""


def build_slot(raw: Dict[str, Any], index: int, slide: Any, resolver: StyleResolver, sw: int, sh: int) -> Dict[str, Any]:
    shape, txb = raw["shape"], raw["txbody"]
    sample = txbody_text(txb)
    dom, sizes = dominant_size(resolver, slide, shape, txb)
    first_font = ""
    for p in txbody_paragraphs(txb):
        runs = text_runs(p)
        if runs:
            first_font = resolver.run_font(slide, shape, p, runs[0])
            break
    box = raw["box"]
    bounds = _bounds_dict(box) if box else None
    slot: Dict[str, Any] = {
        "id": slot_id(index, raw["path"], raw["cell"]),
        "name": raw["name"],
        "kind": raw["kind"],
        "placeholder": raw["placeholder"],
        "sample": sample,
        "paragraphs": len(txbody_paragraphs(txb)),
        "bounds": bounds,
        "fontSizes": sizes,
        "fontPt": dom,
        "font": first_font,
        "wrap": wraps(txb),
        "role": "body",
        "editable": True,
        "editableReason": None,
        "persistent": False,
    }
    if bounds:
        slot["capacity"] = capacity_for(bounds["width"], bounds["height"], dom, txb, sample)
    else:
        slot["capacity"] = None
    # classification that needs only the slot itself
    stripped = sample.strip()
    ph = raw["placeholder"]
    if ph in FOOTER_PH or has_field(txb, ("slidenum",)):
        slot["role"] = FOOTER_PH.get(ph or "", "page-number")
        slot["editable"] = False
        slot["editableReason"] = "page number / date / footer field"
    elif has_field(txb, ("datetime",)):
        slot["role"] = "date"
        slot["editable"] = False
        slot["editableReason"] = "date field"
    elif re.fullmatch(r"\d{1,3}", stripped) and bounds and (
        bounds["top"] + bounds["height"] / 2 > 0.82 * sh / EMU or bounds["top"] + bounds["height"] / 2 < 0.08 * sh / EMU
    ):
        slot["role"] = "page-number"
        slot["editable"] = False
        slot["editableReason"] = "page number"
    elif ph in TITLE_PH:
        slot["role"] = "title"
    elif raw["kind"] == "table-cell":
        slot["role"] = "table-header" if raw.get("tableHeader") else "table-cell"
    elif stripped and NUMERIC_RE.match(stripped) and len(stripped) <= 10:
        slot["role"] = "figure"
    elif ph == "SUBTITLE":
        slot["role"] = "subtitle"
    if slot["editable"]:
        if _is_locked(shape):
            slot["editable"], slot["editableReason"] = False, "locked (noTextEdit)"
        elif dom < 8 and stripped:
            slot["editable"], slot["editableReason"] = False, "very small text (%.1f pt)" % dom
        elif bounds and (bounds["width"] < 0.25 or bounds["height"] < 0.12):
            slot["editable"], slot["editableReason"] = False, "very small shape"
    return slot


def _assign_text_roles(slots: List[Dict[str, Any]], slide_h_in: float) -> None:
    """title / subtitle / body / label roles for slots left as 'body' by build_slot."""
    cands = [s for s in slots if s["role"] == "body" and s["editable"] and s["sample"].strip() and s["bounds"]]
    if not cands:
        return
    # title: largest font in the upper 45 percent, ties broken by position
    upper = [s for s in cands if s["bounds"]["top"] < 0.45 * slide_h_in and s["kind"] == "text"]
    pool = upper or cands
    top_pt = max(s["fontPt"] for s in pool)
    big = sorted((s for s in pool if s["fontPt"] >= top_pt - 0.01), key=lambda s: (s["bounds"]["top"], s["bounds"]["left"]))
    title = big[0] if big and not any(x["role"] == "title" for x in slots) else None
    if title is not None and len(title["sample"]) <= 140 and title["fontPt"] >= 24:
        title["role"] = "title"
    for s in cands:
        if s["role"] != "body":
            continue
        txt = s["sample"].strip()
        long_text = len(txt) >= 60 or s["paragraphs"] >= 2
        if long_text:
            s["role"] = "body"
        elif title is not None and s["bounds"]["top"] >= title["bounds"]["top"] and s["fontPt"] < title["fontPt"] and len(txt) < 60 and s["fontPt"] >= 0.55 * title["fontPt"] and abs(s["bounds"]["top"] - title["bounds"]["top"]) < 1.6:
            s["role"] = "subtitle"
        else:
            s["role"] = "label"


# --------------------------------------------------------------------------
# Pictures
# --------------------------------------------------------------------------


def collect_pictures(prs: Any, slide: Any, index: int) -> List[Dict[str, Any]]:
    out: List[Dict[str, Any]] = []
    sw, sh = prs.slide_width, prs.slide_height
    for shape, path, depth, tf in iter_shapes_abs(slide.shapes):
        el = shape._element
        tag = localname(el)
        is_pic = tag == "pic"
        is_fill = tag == "sp" and el.find("%sspPr/%sblipFill" % (P, A)) is not None
        is_empty_ph = False
        try:
            if tag == "sp" and shape.is_placeholder and shape.placeholder_format.type is not None and shape.placeholder_format.type.name == "PICTURE":
                is_empty_ph = True
        except Exception:
            pass
        if not (is_pic or is_fill or is_empty_ph):
            continue
        box = abs_box(shape, tf)
        if box is None:
            try:
                base = shape._base_placeholder
                bb = C.shape_bounds(base) if base is not None else None
                if bb is not None:
                    box = tuple(float(v) for v in bb)  # type: ignore[assignment]
            except Exception:
                box = None
        if box is None:
            continue
        entry: Dict[str, Any] = {
            "id": slot_id(index, path),
            "name": shape.name,
            "bounds": _bounds_dict(box),
            "placeholder": resolve_ph(shape),
        }
        try:
            img = shape.image
            entry["imagePx"] = list(img.size)
            entry["sha1"] = img.sha1
            entry["contentType"] = img.content_type
        except Exception:
            entry["imagePx"] = None
            entry["empty"] = True
        descr = el.find(".//" + P + "cNvPr")
        if descr is not None and descr.get("descr"):
            entry["alt"] = descr.get("descr")
        frac = ((box[2] - box[0]) * (box[3] - box[1])) / float(sw * sh) if sw and sh else 0.0
        entry["areaFraction"] = round(frac, 4)
        out.append(entry)
    return out


def resolve_ph(shape: Any) -> Optional[str]:
    try:
        if shape.is_placeholder:
            t = shape.placeholder_format.type
            return t.name if t is not None else None
    except Exception:
        pass
    return None


# --------------------------------------------------------------------------
# Page-role guessing
# --------------------------------------------------------------------------

CLOSING_KW = ("thank you", "thanks", "the end", "any questions", "questions?", "q&a", "q & a", "get in touch", "contact us", "stay in touch", "see you")
AGENDA_KW = ("agenda", "table of contents", "contents", "outline", "what we will cover", "what we'll cover", "topics", "roadmap", "today's plan", "overview of")
SECTION_KW = ("section", "chapter", "part one", "part two", "part 1", "part 2", "up next", "next up", "intermission", "transition")


def _count_columns(slots: List[Dict[str, Any]]) -> int:
    """Largest group of body/label blocks that share a row (similar top) and are side by side."""
    blocks = [s for s in slots if s["role"] in ("body", "label", "figure") and s["bounds"] and s["editable"] and s["sample"].strip()]
    best = 0
    for s in blocks:
        row = [t for t in blocks if abs(t["bounds"]["top"] - s["bounds"]["top"]) <= 0.35]
        row.sort(key=lambda t: t["bounds"]["left"])
        # side by side: no horizontal overlap
        distinct = []
        for t in row:
            if all(t["bounds"]["left"] >= d["bounds"]["left"] + d["bounds"]["width"] - 0.05 for d in distinct):
                distinct.append(t)
        best = max(best, len(distinct))
    return best


def guess_role(index: int, count: int, slots: List[Dict[str, Any]], pictures: List[Dict[str, Any]], info: Dict[str, Any], body_median: float) -> Tuple[str, Optional[str], str]:
    live = [s for s in slots if s["editable"] and s["sample"].strip() and s["role"] not in ("page-number", "date", "footer") and not s.get("persistent")]
    total_chars = sum(len(s["sample"]) for s in live)
    text_blob = " ".join(s["sample"].lower() for s in live)
    title = next((s for s in live if s["role"] == "title"), None)
    title_pt = title["fontPt"] if title else 0.0
    title_text = title["sample"].lower() if title else ""

    if index == 1:
        return "cover", None, "first slide"
    short_blob = " ".join(s["sample"].lower() for s in live if len(s["sample"]) <= 60)
    if index > 1 and index >= count - 1 and any(k in short_blob for k in CLOSING_KW):
        return "closing", None, "closing phrase in a short text block"
    if index == count and total_chars <= 300 and len(live) <= 8:
        return "closing", None, "last slide with little text"
    idx_slots = [s for s in live if INDEX_RE.match(s["sample"])]
    if any(k in title_text for k in AGENDA_KW) or (len(idx_slots) >= 4 and len(live) >= 8 and total_chars < 700):
        return "agenda", None, "agenda keyword in title" if any(k in title_text for k in AGENDA_KW) else "numbered list of short items"
    big_title = title_pt >= 1.4 * body_median if body_median else False
    if any(k in title_text for k in SECTION_KW) or (len(live) <= 3 and total_chars <= 140 and (big_title or len(live) <= 2) and not info["hasChart"] and not info["hasTable"]):
        return "section", None, "few text blocks, large title" if big_title else "very little text"
    # content
    big_pics = [p for p in pictures if p["areaFraction"] >= 0.2]
    figures = [s for s in live if s["role"] == "figure"]
    columns = _count_columns(live)
    if info["hasChart"]:
        ctype, why = "chart", "native chart"
    elif info["hasTable"]:
        ctype, why = "table", "table"
    elif big_pics:
        ctype, why = "image-text", "large picture next to text"
    elif len(figures) >= 2:
        ctype, why = "big-numbers", "%d numeric figures" % len(figures)
    elif info["connectors"] >= 3 and columns >= 3:
        ctype, why = "process", "connectors between %d steps" % columns
    elif columns >= 4:
        ctype, why = "grid", "%d blocks in a row" % columns
    elif columns == 3:
        ctype, why = "columns", "three text columns"
    elif columns == 2:
        ctype, why = "columns", "two text columns"
    elif total_chars and text_blob.lstrip().startswith(("“", '"')):
        ctype, why = "quote", "starts with a quotation mark"
    else:
        ctype, why = "text", "single text block"
    return "content", ctype, why


# --------------------------------------------------------------------------
# Whole-deck profile
# --------------------------------------------------------------------------


def _pos_key(bounds: Dict[str, float], pt: float) -> Tuple[float, float, float, float]:
    return (round(bounds["left"] * 5) / 5, round(bounds["top"] * 5) / 5, round(bounds["width"] * 2) / 2, round(pt))


def build_profile(prs: Any, source: str = "") -> Dict[str, Any]:
    resolver = StyleResolver(prs)
    sw, sh = prs.slide_width, prs.slide_height
    sw_in, sh_in = sw / EMU, sh / EMU
    n = len(prs.slides)
    per_slide: List[Dict[str, Any]] = []

    for i, slide in enumerate(prs.slides, start=1):
        raws = collect_raw_slots(prs, slide, i, resolver)
        slots = [build_slot(r, i, slide, resolver, sw, sh) for r in raws]
        _assign_text_roles(slots, sh_in)
        pics = collect_pictures(prs, slide, i)
        info = {
            "hasChart": any(getattr(s, "has_chart", False) and s.has_chart for s in slide.shapes),
            "hasTable": any(getattr(s, "has_table", False) and s.has_table for s in slide.shapes),
            "connectors": sum(1 for s, _p, _d, _t in iter_shapes_abs(slide.shapes) if localname(s._element) == "cxnSp" or (localname(s._element) == "sp" and s.name.lower().startswith(("autoshape", "straight", "arrow", "connector")) and not (s.has_text_frame and s.text_frame.text.strip()))),
            "groups": sum(1 for s in slide.shapes if localname(s._element) == "grpSp"),
        }
        notes = ""
        if slide.has_notes_slide:
            try:
                notes = slide.notes_slide.notes_text_frame.text
            except Exception:
                notes = ""
        per_slide.append({"index": i, "layout": slide.slide_layout.name, "slots": slots, "pictures": pics, "info": info, "notes": notes, "hidden": slide._element.get("show") == "0"})

    # running header / footer labels: same position and size on many slides
    keys: Dict[Tuple[float, float, float, float], List[Tuple[int, Dict[str, Any]]]] = defaultdict(list)
    for ps in per_slide:
        for s in ps["slots"]:
            mid = (s["bounds"]["top"] + s["bounds"]["height"] / 2) if s["bounds"] else 0.5 * sh_in
            in_band = mid < 0.15 * sh_in or mid > 0.85 * sh_in
            if s["editable"] and s["bounds"] and in_band and s["role"] not in ("title", "table-cell", "table-header") and 0 < len(s["sample"].strip()) <= 40 and s["kind"] == "text":
                keys[_pos_key(s["bounds"], s["fontPt"])].append((ps["index"], s))
    threshold = max(3, int(0.4 * max(1, n - 2)))
    for key, members in keys.items():
        slides_with = {i for i, _ in members}
        if len(slides_with) >= threshold:
            texts = Counter(s["sample"].strip() for _, s in members)
            same_text = texts.most_common(1)[0][1] >= 0.6 * len(members)
            for _i, s in members:
                s["role"] = "running-label"
                s["persistent"] = bool(same_text)

    # repeated small pictures are logos
    pic_keys: Dict[Tuple[float, float, str], List[Dict[str, Any]]] = defaultdict(list)
    for ps in per_slide:
        for p in ps["pictures"]:
            pic_keys[(round(p["bounds"]["left"] * 4) / 4, round(p["bounds"]["top"] * 4) / 4, p.get("sha1") or "")].append(p)
    for p_list in pic_keys.values():
        for p in p_list:
            small = p["areaFraction"] <= 0.03
            repeated = len(p_list) >= 3
            if small and repeated:
                p["role"] = "logo"
            elif small:
                p["role"] = "icon"
            elif p["areaFraction"] >= 0.85:
                p["role"] = "background"
            else:
                p["role"] = "photo"
    for ps in per_slide:
        for p in ps["pictures"]:
            p.setdefault("role", "photo")
            p["replaceable"] = p["role"] not in ("logo", "background") or bool(p.get("empty"))
            if p["role"] == "background":
                p["replaceableReason"] = "full-bleed background (decorative)"
            elif p["role"] == "logo":
                p["replaceableReason"] = "logo / repeated mark"

    # deck-level type scale
    weights: Dict[float, Dict[str, int]] = defaultdict(lambda: {"slots": 0, "chars": 0})
    by_role: Dict[str, Counter] = defaultdict(Counter)
    for ps in per_slide:
        for s in ps["slots"]:
            if s["editable"] and s["sample"].strip():
                n_chars = len(s["sample"])
                weights[s["fontPt"]]["slots"] += 1
                weights[s["fontPt"]]["chars"] += n_chars
                by_role[s["role"]][s["fontPt"]] += n_chars
    body_sizes = [(pt, c) for pt, c in by_role.get("body", Counter()).items()]
    body_median = 0.0
    if body_sizes:
        total = sum(c for _pt, c in body_sizes)
        acc = 0
        for pt, c in sorted(body_sizes):
            acc += c
            if acc >= total / 2:
                body_median = pt
                break

    # roles
    roles_to_slides: Dict[str, List[int]] = defaultdict(list)
    ctype_to_slides: Dict[str, List[int]] = defaultdict(list)
    for ps in per_slide:
        role, ctype, why = guess_role(ps["index"], n, ps["slots"], ps["pictures"], ps["info"], body_median)
        ps["role"], ps["contentType"], ps["roleReason"] = role, ctype, why
        roles_to_slides[role].append(ps["index"])
        if ctype:
            ctype_to_slides[ctype].append(ps["index"])

    # palette / fonts
    theme = resolver.theme(prs.slides[0]) if n else Theme(None)
    palette = {k: v for k, v in theme.colors.items()}
    used_fonts = Counter()
    for ps in per_slide:
        for s in ps["slots"]:
            if s["font"] and s["sample"].strip():
                used_fonts[s["font"]] += len(s["sample"])

    def sizes_for(role: str) -> List[float]:
        return sorted(by_role.get(role, Counter()).keys(), reverse=True)

    summary = {
        "slideCount": n,
        "slideSizeIn": [round(sw_in, 3), round(sh_in, 3)],
        "typeScale": {
            "title": sizes_for("title"),
            "subtitle": sizes_for("subtitle"),
            "figure": sizes_for("figure"),
            "body": sizes_for("body"),
            "label": sizes_for("label") + [x for x in sizes_for("running-label") if x not in sizes_for("label")],
            "bodyMedian": body_median,
            "all": [{"pt": pt, "slots": v["slots"], "chars": v["chars"]} for pt, v in sorted(weights.items(), reverse=True)],
        },
        "palette": {
            "dark1": palette.get("dk1"),
            "light1": palette.get("lt1"),
            "dark2": palette.get("dk2"),
            "light2": palette.get("lt2"),
            "accents": [palette[k] for k in ("accent1", "accent2", "accent3", "accent4", "accent5", "accent6") if k in palette],
            "hyperlink": palette.get("hlink"),
        },
        "fonts": {"major": theme.major, "minor": theme.minor, "used": [f for f, _c in used_fonts.most_common()]},
        "rolesToSlides": {k: roles_to_slides.get(k, []) for k in ("cover", "agenda", "section", "content", "closing")},
        "contentTypes": dict(ctype_to_slides),
        "recommended": {k: (roles_to_slides[k][0] if roles_to_slides.get(k) else None) for k in ("cover", "agenda", "section", "content", "closing")},
    }

    slides_out = []
    for ps in per_slide:
        slots_out = ps["slots"]
        slides_out.append(
            {
                "index": ps["index"],
                "layout": ps["layout"],
                "hidden": ps["hidden"],
                "role": ps["role"],
                "contentType": ps["contentType"],
                "roleReason": ps["roleReason"],
                "title": _slide_title_guess(slots_out),
                "hasChart": ps["info"]["hasChart"],
                "hasTable": ps["info"]["hasTable"],
                "notes": ps["notes"],
                "slots": slots_out,
                "pictures": ps["pictures"],
            }
        )
    return {"file": str(source), "slideWidthIn": round(sw_in, 3), "slideHeightIn": round(sh_in, 3), "slideCount": n, "summary": summary, "slides": slides_out}


# --------------------------------------------------------------------------
# Slot lookup (used by fill_template.py)
# --------------------------------------------------------------------------


def find_shape_by_path(slide: Any, path: str) -> Optional[Any]:
    for shape, p, _d, _t in iter_shapes_abs(slide.shapes):
        if p == path:
            return shape
    return None


def find_slot_txbody(slide: Any, sid: str) -> Tuple[Optional[Any], Optional[Any]]:
    """(shape, txBody) for a slot id on *slide*; (None, None) when it does not resolve."""
    _idx, path, cell = parse_slot_id(sid)
    shape = find_shape_by_path(slide, path)
    if shape is None:
        return None, None
    if cell is not None:
        if not (getattr(shape, "has_table", False) and shape.has_table):
            return shape, None
        try:
            return shape, shape.table.cell(cell[0], cell[1]).text_frame._txBody
        except IndexError:
            return shape, None
    return shape, _txbody_of_shape(shape)


def main(argv: Optional[List[str]] = None) -> int:
    ap = argparse.ArgumentParser(
        description="Profile a template .pptx: guessed page roles, text and picture slots with ids, "
        "capacities and editable flags, type scale, palette and fonts, as JSON.",
    )
    ap.add_argument("input", help="template .pptx")
    ap.add_argument("-o", "--output", help="write the profile JSON here (default: stdout)")
    args = ap.parse_args(argv)
    try:
        prs = Presentation(args.input)
        profile = build_profile(prs, args.input)
    except Exception as exc:
        print("error: cannot profile %s: %s" % (args.input, exc), file=sys.stderr)
        return 2
    text = json.dumps(profile, indent=2, ensure_ascii=False)
    if args.output:
        Path(args.output).parent.mkdir(parents=True, exist_ok=True)
        Path(args.output).write_text(text + "\n", encoding="utf-8")
        s = profile["summary"]
        nslots = sum(len(x["slots"]) for x in profile["slides"])
        ned = sum(1 for x in profile["slides"] for sl in x["slots"] if sl["editable"])
        print(
            "wrote %s: %d slide(s), %d text slot(s) (%d editable), %d picture slot(s); roles %s"
            % (args.output, profile["slideCount"], nslots, ned, sum(len(x["pictures"]) for x in profile["slides"]), json.dumps(s["rolesToSlides"])),
            file=sys.stderr,
        )
    else:
        print(text)
    return 0


if __name__ == "__main__":
    sys.exit(main())
