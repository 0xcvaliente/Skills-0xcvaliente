#!/usr/bin/env python3
"""Replace text in an existing .pptx while keeping its formatting.

Usage:
    replace_text.py IN.pptx REPLACEMENTS.json -o OUT.pptx [--dry-run]

REPLACEMENTS.json
-----------------
{
  "slides": {
    "2": {
      "5":        {"paragraphs": ["New first paragraph", "New second paragraph"]},
      "Subtitle": "One paragraph\\nanother paragraph",
      "3/2":      "text inside the group member with path 3/2",
      "7":        {"cells": {"2,3": "new text for row 2, column 3"}},
      "7@3,1":    "same idea: cell row 3, column 1 of table shape 7",
      "9":        {"rows": [["Region", "Revenue"], [null, "13.0"]]},
      "notes":    {"paragraphs": ["New speaker notes"]}
    }
  },
  "find_replace": [
    {"find": "Q3", "replace": "Q4"},
    {"find": "Acme Corp", "replace": "Acme Inc.", "slides": [1, 4]}
  ]
}

Keys inside a slide are matched, in this order, as a shape path from
inventory.py ("3/2"), a numeric shape id, or a shape name. A name that matches
several shapes is an error (use the id or path). Slide numbers are 1-based.

How formatting is kept
----------------------
* Paragraph replacement: the new text goes into the first run of each existing
  paragraph (its run properties stay); the other runs are removed. Extra new
  paragraphs are clones of the last existing paragraph, so they share its
  formatting. Paragraphs beyond the new list are removed. A "\\v" in a string
  becomes a line break.
* find/replace works across run boundaries: the runs a match touches are merged
  into the first of them (first run's formatting wins). Text in groups, tables
  and speaker notes is included. A match never changes position, size, colour
  or font of any shape.

Every key and every find that matches nothing is an error: the summary lists
them and the exit status is 1 (nothing is written). --dry-run does everything
except writing OUT.pptx. Exit status: 0 ok, 1 unmatched/invalid input, 2 unreadable files.
"""

from __future__ import annotations

import argparse
import copy
import json
import sys
from pathlib import Path
from typing import Any, Dict, Iterator, List, Optional, Tuple

sys.path.insert(0, str(Path(__file__).resolve().parent))
import _common as C  # noqa: E402

C.require_pptx()

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402

from inventory import iter_shapes_abs, localname  # noqa: E402

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
A = "{%s}" % NS_A


# --------------------------------------------------------------------------
# Paragraph-level editing (lxml on a:p elements)
# --------------------------------------------------------------------------


def _text_el(run: Any) -> Any:
    t = run.find(A + "t")
    if t is None:
        t = etree.SubElement(run, A + "t")
    return t


def _set_text(run: Any, text: str) -> None:
    t = _text_el(run)
    t.text = text
    if text != text.strip() or "  " in text:
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def _end_props(p: Any) -> Any:
    return p.find(A + "endParaRPr")


def set_paragraph_text(p: Any, text: str) -> None:
    """Put *text* in the first run of paragraph element *p*; drop the other runs."""
    segments = text.split("\v")
    content = [c for c in p if localname(c) in ("r", "br", "fld")]
    first = next((c for c in content if localname(c) in ("r", "fld")), None)
    if first is None:
        first = etree.Element(A + "r")
        end = _end_props(p)
        rpr = etree.SubElement(first, A + "rPr")
        if end is not None:
            rpr.attrib.update(end.attrib)
            for ch in end:
                rpr.append(copy.deepcopy(ch))
            end.addprevious(first)
        else:
            p.append(first)
        etree.SubElement(first, A + "t")
    elif localname(first) == "fld":
        # a field (slide number, date) is replaced by a plain run keeping its properties
        run = etree.Element(A + "r")
        rpr = first.find(A + "rPr")
        if rpr is not None:
            run.append(copy.deepcopy(rpr))
        etree.SubElement(run, A + "t")
        first.addprevious(run)
        p.remove(first)
        content = [run if c is first else c for c in content]
        first = run
    for c in content:
        if c is not first:
            p.remove(c)
    _set_text(first, segments[0])
    anchor = first
    for seg in segments[1:]:
        rpr = first.find(A + "rPr")
        br = etree.Element(A + "br")
        if rpr is not None:
            br.append(copy.deepcopy(rpr))
        anchor.addnext(br)
        run = copy.deepcopy(first)
        _set_text(run, seg)
        br.addnext(run)
        anchor = run


def paragraph_runs(p: Any) -> List[Any]:
    """The text-bearing children (a:r and a:fld) of paragraph element *p*."""
    return [c for c in p if localname(c) in ("r", "fld")]


def run_text(r: Any) -> str:
    t = r.find(A + "t")
    return (t.text or "") if t is not None else ""


def replace_in_paragraph(p: Any, find: str, replace: str) -> int:
    """Replace every occurrence of *find* in paragraph *p*, even across runs. Returns the count."""
    runs = paragraph_runs(p)
    if not runs:
        return 0
    texts = [run_text(r) for r in runs]
    joined = "".join(texts)
    if find not in joined:
        return 0
    # occurrences, non-overlapping, left to right
    spans: List[Tuple[int, int]] = []
    pos = 0
    while True:
        i = joined.find(find, pos)
        if i < 0:
            break
        spans.append((i, i + len(find)))
        pos = i + len(find)
    starts: List[int] = []
    acc = 0
    for t in texts:
        starts.append(acc)
        acc += len(t)

    def run_at(offset: int, end: bool = False) -> int:
        for idx in range(len(texts)):
            lo, hi = starts[idx], starts[idx] + len(texts[idx])
            if (lo <= offset < hi) or (end and lo < offset <= hi):
                return idx
        return len(texts) - 1

    for s, e in reversed(spans):
        a = run_at(s)
        b = run_at(e, end=True) if e > s else a
        if b < a:
            b = a
        head = texts[a][: s - starts[a]]
        tail = texts[b][e - starts[b]:]
        texts[a] = head + replace + tail
        for k in range(a + 1, b + 1):
            texts[k] = None  # type: ignore[assignment]
        # recompute offsets lazily: only runs before a matter for the next (earlier) span
    # write back
    for idx, r in enumerate(runs):
        if texts[idx] is None:
            p.remove(r)
        else:
            if localname(r) == "fld" and texts[idx] != run_text(r):
                # edited field becomes a plain run so PowerPoint does not regenerate its text
                nr = etree.Element(A + "r")
                rpr = r.find(A + "rPr")
                if rpr is not None:
                    nr.append(copy.deepcopy(rpr))
                etree.SubElement(nr, A + "t")
                r.addprevious(nr)
                p.remove(r)
                r = nr
            _set_text(r, texts[idx])
    return len(spans)


# --------------------------------------------------------------------------
# Text-body level
# --------------------------------------------------------------------------


def body_paragraphs(txbody: Any) -> List[Any]:
    return [c for c in txbody if localname(c) == "p"]


def set_body_paragraphs(txbody: Any, texts: List[str]) -> None:
    """Replace the paragraphs of a txBody with *texts*, keeping run and paragraph formatting."""
    paras = body_paragraphs(txbody)
    if not paras:
        paras = [etree.SubElement(txbody, A + "p")]
    if not texts:
        texts = [""]
    last = paras[-1]
    for i, text in enumerate(texts):
        if i < len(paras):
            p = paras[i]
        else:
            p = copy.deepcopy(last)
            body_paragraphs(txbody)[-1].addnext(p)
        set_paragraph_text(p, text)
    for p in paras[len(texts):]:
        txbody.remove(p)


def split_text(value: Any) -> List[str]:
    """Normalise a replacement value ("text" or {"paragraphs": [...]}) to a list of paragraph strings."""
    if isinstance(value, str):
        return value.split("\n")
    if isinstance(value, dict) and isinstance(value.get("paragraphs"), list) and all(isinstance(x, str) for x in value["paragraphs"]):
        return list(value["paragraphs"])
    if isinstance(value, dict) and isinstance(value.get("text"), str):
        return value["text"].split("\n")
    raise ValueError('expected a string or {"paragraphs": ["..."]}')


def shape_txbodies(shape: Any) -> List[Any]:
    """txBody elements of a shape: its own text frame, or every table cell's."""
    out = []
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                out.append(cell.text_frame._txBody)
    elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        out.append(shape.text_frame._txBody)
    return out


def slide_txbodies(slide: Any, include_notes: bool = True) -> Iterator[Tuple[str, Any]]:
    """Yield ``(label, txBody)`` for every text body on a slide (shapes, groups, tables, notes)."""
    for shp, path, _d, _tf in iter_shapes_abs(slide.shapes):
        for tb in shape_txbodies(shp):
            yield "shape %s" % path, tb
    if include_notes and slide.has_notes_slide:
        for shp in slide.notes_slide.shapes:
            if getattr(shp, "has_text_frame", False) and shp.has_text_frame:
                yield "notes", shp.text_frame._txBody


# --------------------------------------------------------------------------
# Key resolution
# --------------------------------------------------------------------------


def resolve_keys(slide: Any, key: str) -> List[Any]:
    """Shapes on *slide* matching a path, a numeric id or a name."""
    shapes = list(iter_shapes_abs(slide.shapes))
    by_path = [s for s, p, _d, _t in shapes if p == key]
    if by_path:
        return by_path
    if key.isdigit():
        by_id = [s for s, _p, _d, _t in shapes if str(s.shape_id) == key]
        if by_id:
            return by_id[:1]
    return [s for s, _p, _d, _t in shapes if s.name == key]


def parse_cell_ref(ref: str) -> Tuple[int, int]:
    r, c = ref.replace(" ", "").split(",")
    r_i, c_i = int(r), int(c)
    if r_i < 1 or c_i < 1:
        raise ValueError("cell references are 1-based")
    return r_i - 1, c_i - 1


# --------------------------------------------------------------------------
# Applying a replacement file
# --------------------------------------------------------------------------


class Result:
    def __init__(self) -> None:
        self.done: List[str] = []
        self.errors: List[str] = []
        self.find_hits: Dict[int, int] = {}


def apply_slide_entry(slide: Any, snum: int, key: str, value: Any, res: Result) -> None:
    label = "slide %d key %r" % (snum, key)
    if key == "notes":
        if not slide.has_notes_slide:
            res.errors.append("%s: slide has no speaker notes to edit" % label)
            return
        try:
            texts = split_text(value)
        except ValueError as exc:
            res.errors.append("%s: %s" % (label, exc))
            return
        set_body_paragraphs(slide.notes_slide.notes_text_frame._txBody, texts)
        res.done.append("%s: notes replaced (%d paragraph(s))" % (label, len(texts)))
        return
    cell_ref: Optional[str] = None
    shape_key = key
    if "@" in key:
        shape_key, cell_ref = key.split("@", 1)
    shapes = resolve_keys(slide, shape_key)
    if not shapes:
        res.errors.append("%s: no shape with that path, id or name" % label)
        return
    if len(shapes) > 1:
        res.errors.append("%s: name matches %d shapes; use a numeric id or a path" % (label, len(shapes)))
        return
    shape = shapes[0]
    is_table = bool(getattr(shape, "has_table", False) and shape.has_table)
    try:
        if is_table or cell_ref:
            if not is_table:
                res.errors.append("%s: shape is not a table" % label)
                return
            cells: Dict[Tuple[int, int], Any] = {}
            if cell_ref:
                cells[parse_cell_ref(cell_ref)] = value
            elif isinstance(value, dict) and isinstance(value.get("cells"), dict):
                for ref, v in value["cells"].items():
                    cells[parse_cell_ref(ref)] = v
            elif isinstance(value, dict) and isinstance(value.get("rows"), list):
                for ri, row in enumerate(value["rows"]):
                    for ci, v in enumerate(row):
                        if v is not None:
                            cells[(ri, ci)] = v
            else:
                res.errors.append('%s: a table needs {"cells": {"r,c": text}}, {"rows": [[...]]} or a key like "7@2,3"' % label)
                return
            tbl = shape.table
            for (ri, ci), v in cells.items():
                if ri >= len(tbl.rows) or ci >= len(tbl.columns):
                    res.errors.append("%s: cell %d,%d is outside the %dx%d table" % (label, ri + 1, ci + 1, len(tbl.rows), len(tbl.columns)))
                    continue
                set_body_paragraphs(tbl.cell(ri, ci).text_frame._txBody, split_text(v))
                res.done.append("%s: table cell %d,%d replaced" % (label, ri + 1, ci + 1))
            return
        if not (getattr(shape, "has_text_frame", False) and shape.has_text_frame):
            res.errors.append("%s: shape %r has no text frame" % (label, shape.name))
            return
        texts = split_text(value)
        set_body_paragraphs(shape.text_frame._txBody, texts)
        res.done.append("%s: %r replaced (%d paragraph(s))" % (label, shape.name, len(texts)))
    except (ValueError, TypeError) as exc:
        res.errors.append("%s: %s" % (label, exc))


def apply_replacements(prs: Any, spec: Dict[str, Any]) -> Result:
    res = Result()
    slides = list(prs.slides)
    n = len(slides)
    unknown = set(spec) - {"slides", "find_replace"}
    for k in sorted(unknown):
        res.errors.append("unknown top-level key %r (expected 'slides' and/or 'find_replace')" % k)
    for skey, entries in (spec.get("slides") or {}).items():
        try:
            snum = int(skey)
        except ValueError:
            res.errors.append("slide key %r is not a slide number" % skey)
            continue
        if not 1 <= snum <= n:
            res.errors.append("slide %d does not exist (deck has %d slides)" % (snum, n))
            continue
        if not isinstance(entries, dict):
            res.errors.append("slide %d: expected an object of shape keys" % snum)
            continue
        for key, value in entries.items():
            apply_slide_entry(slides[snum - 1], snum, key, value, res)

    for i, rule in enumerate(spec.get("find_replace") or [], start=1):
        label = "find_replace[%d]" % i
        if not isinstance(rule, dict) or not isinstance(rule.get("find"), str) or not isinstance(rule.get("replace"), str):
            res.errors.append('%s: needs string "find" and "replace"' % label)
            continue
        if rule["find"] == "":
            res.errors.append("%s: 'find' is empty" % label)
            continue
        only = rule.get("slides")
        if only is not None and not (isinstance(only, list) and all(isinstance(x, int) for x in only)):
            res.errors.append("%s: 'slides' must be a list of 1-based slide numbers" % label)
            continue
        total = 0
        touched = set()
        for idx, slide in enumerate(slides, start=1):
            if only is not None and idx not in only:
                continue
            for _lbl, tb in slide_txbodies(slide):
                for p in body_paragraphs(tb):
                    k = replace_in_paragraph(p, rule["find"], rule["replace"])
                    if k:
                        total += k
                        touched.add(idx)
        if total == 0:
            res.errors.append("%s: %r not found%s" % (label, rule["find"], (" on slides %s" % only) if only else ""))
        else:
            res.done.append("%s: %r -> %r, %d replacement(s) on slide(s) %s" % (label, rule["find"], rule["replace"], total, ",".join(map(str, sorted(touched)))))
    return res


def main(argv: Optional[List[str]] = None) -> int:
    ap = argparse.ArgumentParser(
        description="Replace text in a .pptx without touching position, size, colour or font. "
        "See the module docstring for the REPLACEMENTS.json format.",
    )
    ap.add_argument("input", help="input .pptx")
    ap.add_argument("replacements", help="JSON file with 'slides' and/or 'find_replace'")
    ap.add_argument("-o", "--output", help="output .pptx (required unless --dry-run)")
    ap.add_argument("--dry-run", action="store_true", help="apply in memory and report, but write nothing")
    args = ap.parse_args(argv)
    if not args.output and not args.dry_run:
        ap.error("-o/--output is required unless --dry-run is given")
    try:
        spec = C.load_json(args.replacements)
        prs = Presentation(args.input)
    except (OSError, ValueError, KeyError) as exc:
        print("error: %s" % exc, file=sys.stderr)
        return 2
    if not isinstance(spec, dict):
        print("error: the replacements file must contain a JSON object", file=sys.stderr)
        return 2
    res = apply_replacements(prs, spec)
    for line in res.done:
        print("ok     " + line)
    for line in res.errors:
        print("ERROR  " + line)
    print("%d applied, %d error(s)%s" % (len(res.done), len(res.errors), " (dry run)" if args.dry_run else ""))
    if res.errors:
        print("nothing written because of the errors above", file=sys.stderr)
        return 1
    if not args.dry_run:
        prs.save(args.output)
        print("wrote %s" % args.output)
    return 0


if __name__ == "__main__":
    sys.exit(main())
