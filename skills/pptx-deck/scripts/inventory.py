#!/usr/bin/env python3
"""Dump a structural inventory of an existing .pptx as JSON.

For every slide: 1-based index, layout name, hidden flag, speaker notes, and
every shape, recursively (group members are flattened into the same list and
carry a ``path`` such as ``"3/2"``: shape id 3 is the group, 2 is the member).

Per shape: id, name, path, parent, kind (text, placeholder, picture, table, chart,
group, connector, freeform, shape, graphic), placeholder type/index, bounds in
inches measured on the slide (group transforms applied; ``null`` when the
position is inherited), an out-of-bounds flag, and kind-specific detail:

* text and placeholders: paragraphs with level, alignment, bullet and runs
  (text, font name, size in pt, bold, italic, colour where set explicitly) and
  the text-frame autofit mode;
* tables: the cell text grid;
* charts: chart type, categories and series values;
* pictures: content type, pixel size and alt text.

The ``path`` (or the numeric ``id``) is what replace_text.py accepts as a key.

Usage:
    inventory.py deck.pptx                    # JSON to stdout
    inventory.py deck.pptx -o inventory.json  # JSON to a file
    inventory.py deck.pptx --slides 1,3-5     # only some slides
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any, Dict, Iterator, List, Optional, Tuple

sys.path.insert(0, str(Path(__file__).resolve().parent))
import _common as C  # noqa: E402

C.require_pptx()

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "{%s}" % NS_A
EMU = C.EMU_PER_INCH

Transform = Tuple[float, float, float, float]  # ax, bx, ay, by  (abs = a * v + b)
IDENTITY: Transform = (1.0, 0.0, 1.0, 0.0)


def localname(el: Any) -> str:
    return etree.QName(el).localname if isinstance(el.tag, str) else ""


# --------------------------------------------------------------------------
# Geometry
# --------------------------------------------------------------------------


def own_box(shape: Any) -> Optional[Tuple[int, int, int, int]]:
    """(x, y, cx, cy) in EMU in the shape's own coordinate space, or None when inherited."""
    b = C.shape_bounds(shape)
    if b is None:
        return None
    return b[0], b[1], b[2] - b[0], b[3] - b[1]


def abs_box(shape: Any, tf: Transform) -> Optional[Tuple[float, float, float, float]]:
    """Slide-space (left, top, right, bottom) in EMU after applying the group transform *tf*."""
    box = own_box(shape)
    if box is None:
        return None
    x, y, cx, cy = box
    ax, bx, ay, by = tf
    l, t = ax * x + bx, ay * y + by
    return l, t, l + ax * cx, t + ay * cy


def child_transform(group: Any, tf: Transform) -> Transform:
    """Transform for the members of *group*, composed with the parent transform *tf*."""
    xfrm = group._element.find("{%s}grpSpPr/{%s}xfrm" % (NS_P, NS_A))
    box = own_box(group)
    if xfrm is None or box is None:
        return tf
    ch_off = xfrm.find(A + "chOff")
    ch_ext = xfrm.find(A + "chExt")
    if ch_off is None or ch_ext is None:
        return tf
    cox, coy = int(ch_off.get("x", 0)), int(ch_off.get("y", 0))
    cex, cey = int(ch_ext.get("cx", 0)), int(ch_ext.get("cy", 0))
    x, y, cx, cy = box
    sx = cx / cex if cex else 1.0
    sy = cy / cey if cey else 1.0
    ax, bx, ay, by = tf
    # g(v) = x + (v - cox) * sx  then parent: ax * g + bx
    return (ax * sx, ax * (x - cox * sx) + bx, ay * sy, ay * (y - coy * sy) + by)


def iter_shapes_abs(shapes: Any, tf: Transform = IDENTITY, parent: str = "") -> Iterator[Tuple[Any, str, int, Transform]]:
    """Yield ``(shape, path, depth, transform)`` for every shape, descending into groups.

    ``path`` is the "/"-joined shape ids from the top level, e.g. ``"3/2"``.
    ``transform`` maps the shape's own coordinates to slide coordinates.
    """
    for shp in shapes:
        path = "%s/%d" % (parent, shp.shape_id) if parent else str(shp.shape_id)
        yield shp, path, path.count("/"), tf
        if localname(shp._element) == "grpSp":
            yield from iter_shapes_abs(shp.shapes, child_transform(shp, tf), path)


# --------------------------------------------------------------------------
# Shape description
# --------------------------------------------------------------------------


def shape_kind(shape: Any) -> str:
    tag = localname(shape._element)
    if tag == "grpSp":
        return "group"
    if tag == "cxnSp":
        return "connector"
    if tag == "pic":
        return "placeholder" if shape.is_placeholder else "picture"
    if tag == "graphicFrame":
        if getattr(shape, "has_table", False) and shape.has_table:
            return "table"
        if getattr(shape, "has_chart", False) and shape.has_chart:
            return "chart"
        return "graphic"
    if tag == "sp":
        if shape.is_placeholder:
            return "placeholder"
        if shape._element.find(".//{%s}custGeom" % NS_A) is not None and not _has_text(shape):
            return "freeform"
        if _has_text(shape):
            return "text"
        return "freeform" if shape._element.find(".//{%s}custGeom" % NS_A) is not None else "shape"
    return tag or "unknown"


def _has_text(shape: Any) -> bool:
    try:
        return bool(shape.has_text_frame and shape.text_frame.text.strip())
    except Exception:
        return False


def run_color(font: Any) -> Optional[str]:
    try:
        col = font.color
        if col is None or col.type is None:
            return None
        if col.type == 1:  # MSO_COLOR_TYPE.RGB
            return "#" + str(col.rgb)
        tc = col.theme_color
        return "theme:" + (tc.name.lower() if hasattr(tc, "name") else str(tc))
    except Exception:
        return None


def describe_run(run: Any) -> Dict[str, Any]:
    f = run.font
    out: Dict[str, Any] = {"text": run.text}
    if f.name:
        out["font"] = f.name
    if f.size is not None:
        out["size"] = round(f.size.pt, 2)
    if f.bold is not None:
        out["bold"] = bool(f.bold)
    if f.italic is not None:
        out["italic"] = bool(f.italic)
    col = run_color(f)
    if col:
        out["color"] = col
    return out


def autofit_mode(text_frame: Any) -> str:
    try:
        bp = text_frame._txBody.bodyPr
    except Exception:
        return "inherit"
    for child in bp:
        n = localname(child)
        if n == "normAutofit":
            return "shrink"
        if n == "spAutoFit":
            return "resize-shape"
        if n == "noAutofit":
            return "none"
    return "inherit"


def describe_paragraphs(text_frame: Any) -> List[Dict[str, Any]]:
    paras = []
    for p in text_frame.paragraphs:
        d: Dict[str, Any] = {"level": p.level, "text": p.text, "runs": [describe_run(r) for r in p.runs]}
        if p.alignment is not None:
            d["align"] = p.alignment.name.lower() if hasattr(p.alignment, "name") else str(p.alignment)
        ppr = p._p.find(A + "pPr")
        if ppr is not None:
            if ppr.find(A + "buChar") is not None or ppr.find(A + "buAutoNum") is not None:
                d["bullet"] = True
            elif ppr.find(A + "buNone") is not None:
                d["bullet"] = False
        paras.append(d)
    return paras


def describe_chart(chart: Any) -> Dict[str, Any]:
    out: Dict[str, Any] = {}
    try:
        out["chartType"] = chart.chart_type.name
    except Exception:
        out["chartType"] = "unknown"
    try:
        if chart.has_title and chart.chart_title.has_text_frame:
            out["title"] = chart.chart_title.text_frame.text
    except Exception:
        pass
    cats: List[str] = []
    series: List[Dict[str, Any]] = []
    try:
        for plot in chart.plots:
            if not cats:
                cats = [str(c) for c in plot.categories]
            for s in plot.series:
                series.append({"name": s.name, "values": [None if v is None else v for v in s.values]})
    except Exception:
        pass
    out["categories"] = cats
    out["series"] = series
    return out


def describe_shape(shape: Any, path: str, depth: int, tf: Transform, sw: int, sh: int) -> Dict[str, Any]:
    kind = shape_kind(shape)
    d: Dict[str, Any] = {
        "id": shape.shape_id,
        "name": shape.name,
        "path": path,
        "depth": depth,
        "type": kind,
    }
    if depth:
        d["parent"] = path.rsplit("/", 1)[0]
    if shape.is_placeholder:
        pf = shape.placeholder_format
        d["placeholder"] = {"type": pf.type.name if pf.type is not None else None, "idx": pf.idx}
    box = abs_box(shape, tf)
    if box is None:
        d["bounds"] = None
        d["outOfBounds"] = False
    else:
        l, t, r, b = box
        d["bounds"] = {
            "left": round(l / EMU, 3),
            "top": round(t / EMU, 3),
            "width": round((r - l) / EMU, 3),
            "height": round((b - t) / EMU, 3),
        }
        tol = 0.02 * EMU
        d["outOfBounds"] = bool(l < -tol or t < -tol or r > sw + tol or b > sh + tol)
    if kind == "group":
        d["children"] = [c.shape_id for c in shape.shapes]
    try:
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame and kind in ("text", "placeholder", "shape", "freeform"):
            tf_ = shape.text_frame
            d["paragraphs"] = describe_paragraphs(tf_)
            d["autofit"] = autofit_mode(tf_)
    except Exception:
        pass
    if kind == "table":
        tbl = shape.table
        d["rows"] = len(tbl.rows)
        d["columns"] = len(tbl.columns)
        d["cells"] = [[c.text for c in row.cells] for row in tbl.rows]
    elif kind == "chart":
        d.update(describe_chart(shape.chart))
    elif kind == "picture" or (kind == "placeholder" and localname(shape._element) == "pic"):
        try:
            img = shape.image
            d["image"] = {"contentType": img.content_type, "pixels": list(img.size), "sha1": img.sha1}
        except Exception:
            d["image"] = None
        cnv = shape._element.find(".//{%s}cNvPr" % NS_P)
        if cnv is not None and cnv.get("descr"):
            d["alt"] = cnv.get("descr")
    return d


def describe_slide(prs: Any, slide: Any, index: int) -> Dict[str, Any]:
    sw, sh = prs.slide_width, prs.slide_height
    notes = ""
    if slide.has_notes_slide:
        try:
            notes = slide.notes_slide.notes_text_frame.text
        except Exception:
            notes = ""
    return {
        "index": index,
        "layout": slide.slide_layout.name,
        "hidden": slide._element.get("show") == "0",
        "notes": notes,
        "shapes": [describe_shape(s, p, d, tf, sw, sh) for s, p, d, tf in iter_shapes_abs(slide.shapes)],
    }


def parse_slide_list(spec: str, count: int) -> List[int]:
    """Parse "1,3-5" into sorted 1-based indices within 1..count."""
    out = set()
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            lo, hi = int(a), int(b)
        else:
            lo = hi = int(part)
        for i in range(lo, hi + 1):
            if not 1 <= i <= count:
                raise ValueError("slide %d is out of range 1..%d" % (i, count))
            out.add(i)
    return sorted(out)


def build_inventory(path: str, slides: Optional[str] = None) -> Dict[str, Any]:
    prs = Presentation(path)
    n = len(prs.slides)
    wanted = parse_slide_list(slides, n) if slides else list(range(1, n + 1))
    return {
        "file": str(path),
        "slideWidthIn": round(prs.slide_width / EMU, 3),
        "slideHeightIn": round(prs.slide_height / EMU, 3),
        "slideCount": n,
        "slides": [describe_slide(prs, prs.slides[i - 1], i) for i in wanted],
    }


def main(argv: Optional[List[str]] = None) -> int:
    ap = argparse.ArgumentParser(
        description="Dump slides, layouts, notes and every shape (groups flattened, with paths) of a .pptx as JSON.",
    )
    ap.add_argument("input", help="input .pptx")
    ap.add_argument("-o", "--output", help="write JSON here instead of stdout")
    ap.add_argument("--slides", help="1-based slide list, e.g. 1,3 or 2-5 (default: all)")
    args = ap.parse_args(argv)
    try:
        inv = build_inventory(args.input, args.slides)
    except (OSError, ValueError, KeyError) as exc:
        print("error: %s" % exc, file=sys.stderr)
        return 2
    text = json.dumps(inv, indent=2, ensure_ascii=False)
    if args.output:
        Path(args.output).write_text(text + "\n", encoding="utf-8")
        print("wrote %s (%d slide(s))" % (args.output, len(inv["slides"])), file=sys.stderr)
    else:
        print(text)
    return 0


if __name__ == "__main__":
    sys.exit(main())
