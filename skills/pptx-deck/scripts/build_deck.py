#!/usr/bin/env python3
"""Build a .pptx deck from a pptx-deck specification (JSON).

Pipeline
--------
1. Validate the spec with validate_spec.py (skip with --skip-validate).
2. Create a 16:9 (or 4:3) presentation, or open --template and reuse its slide
   masters and layouts (all of the template's existing slides are removed
   cleanly; no orphan slide parts remain).
3. Render every slide with the renderer for its role. Each of the 20 roles
   implements every layout variant described in references/layout-roles.md.
   Slides are made of native text boxes, shapes, pictures and native charts;
   no slide numbers or footers are added.
4. Add speaker notes (the slide's notes, or its coreMessage when notes are
   missing) and fixed document properties, then write the file.

Text policy: the full text is used; when it cannot fit at the minimum font
sizes (body 18 pt, caption 12 pt, title 28 pt, hero numbers 2.5 x body) the
short text is used; if even that does not fit, the slide is still written at the
minimum size and a "split this slide" warning is printed. Fonts are never
shrunk below the minimums.

Determinism: the same spec, theme and template produce byte-identical slide and
chart XML (no timestamps or random values; core properties are derived from
the spec; the .pptx zip uses fixed entry timestamps).

Template mode (--template existing.pptx): reuses the template's masters,
layouts, theme fonts and theme colours (unless --theme forces a skill theme).
It builds NEW slides on the template's blank layout. Filling the template's own
existing slides is a different job handled by fill_template.py (built later).
"""

from __future__ import annotations

import argparse
import copy
import datetime as _dt
import io
import math
import re
import sys
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

sys.path.insert(0, str(Path(__file__).resolve().parent))
import _common as C  # noqa: E402

C.require_pptx()

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_MARKER_STYLE, XL_TICK_MARK  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN  # noqa: E402
from pptx.opc.constants import RELATIONSHIP_TYPE as RT  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Emu, Pt  # noqa: E402

# --------------------------------------------------------------------------
# Constants
# --------------------------------------------------------------------------

M = 0.6  # outer margin, design inches
LINE_SPACING = 1.2
FIT_STEPS = (1.0, 0.75, 0.5, 0.25, 0.0)
INSET_X = 0.05  # text box inner margins (design inches)
INSET_Y = 0.03

_TYPO = C.typography_limits()
MIN_PT = {
    "body": float(_TYPO["minPt"]["body"]),
    "caption": float(_TYPO["minPt"]["caption"]),
    "title": float(_TYPO["minPt"]["title"]),
    "stat": float(_TYPO["minPt"]["title"]),
    "hero": float(_TYPO["minPt"]["body"]) * float(_TYPO["heroToBodyRatioMin"]),
}
PREF_PT = {"body": 20.0, "caption": 14.0, "title": 38.0, "stat": 36.0, "hero": 72.0}

# Average glyph advance in em for fonts wider or narrower than the 0.5 default;
# the builder is deliberately a little conservative.
FONT_EM = {
    "impact": 0.52,
    "archivo black": 0.62,
    "franklin gothic demi": 0.56,
    "space grotesk": 0.56,
    "georgia": 0.55,
    "playfair display": 0.55,
    "fraunces": 0.57,
    "source serif 4": 0.54,
    "consolas": 0.56,
    "courier new": 0.61,
    "bahnschrift": 0.52,
    "trebuchet ms": 0.54,
    "helvetica neue": 0.53,
    "segoe ui semibold": 0.54,
    "segoe ui": 0.52,
    "arial": 0.53,
    "calibri": 0.5,
}
HEAVY_FONTS = ("impact", "archivo black", "franklin gothic demi", "semibold", "black")
SERIF_FONTS = ("georgia", "playfair display", "fraunces", "source serif 4")

CHART_TYPES = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "stacked_bar": XL_CHART_TYPE.BAR_STACKED,
}

FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)


# --------------------------------------------------------------------------
# Determinism patches
# --------------------------------------------------------------------------


def _patch_chart_workbook_determinism() -> None:
    """Give the XlsxWriter workbooks embedded behind charts a fixed creation time."""
    try:
        from pptx.chart import xlsx as pxlsx
        from contextlib import contextmanager
        from xlsxwriter import Workbook

        if getattr(pxlsx._BaseWorkbookWriter, "_pptx_deck_patched", False):
            return

        @contextmanager
        def _open_worksheet(self, xlsx_file):  # noqa: ANN001
            workbook = Workbook(xlsx_file, {"in_memory": True})
            workbook.set_properties({"created": _dt.datetime(2000, 1, 1)})
            worksheet = workbook.add_worksheet()
            yield workbook, worksheet
            workbook.close()

        pxlsx._BaseWorkbookWriter._open_worksheet = _open_worksheet
        pxlsx._BaseWorkbookWriter._pptx_deck_patched = True
    except Exception:  # pragma: no cover - best effort only
        pass


_patch_chart_workbook_determinism()


# --------------------------------------------------------------------------
# Fit engine
# --------------------------------------------------------------------------


class P:
    """One paragraph of a text block.

    ``cands`` is the ordered list of text alternatives (full first, short second).
    ``kind`` is body / caption / title / stat / hero and fixes the minimum size.
    ``color`` is a symbolic name (text, muted, accent, primary, onfill) or a hex string.
    """

    def __init__(
        self,
        cands: Sequence[str],
        kind: str = "body",
        pref: Optional[float] = None,
        bold: bool = False,
        color: str = "text",
        font: str = "body",
        before: float = 0.0,
        italic: bool = False,
        bullet: bool = False,
        outline: bool = False,
        min_pt: Optional[float] = None,
        indent: float = 0.0,
        single_line: bool = False,
    ) -> None:
        cl = [c for c in cands if c]
        self.cands = cl if cl else [""]
        self.kind = kind
        self.min = max(MIN_PT[kind], min_pt or 0.0)
        self.pref = max(self.min, pref if pref is not None else PREF_PT[kind])
        self.bold = bold
        self.color = color
        self.font = font
        self.before = before
        self.italic = italic
        self.bullet = bullet
        self.outline = outline
        self.indent = indent
        self.single_line = single_line or kind == "hero"

    def text_at(self, level: int) -> str:
        return self.cands[min(level, len(self.cands) - 1)]

    def size_at(self, t: float) -> float:
        return round((self.min + t * (self.pref - self.min)) * 2.0) / 2.0


def fit_blocks(
    em_of,
    blocks: Sequence[Tuple[Sequence[P], float, float]],
) -> Tuple[int, float, List[bool]]:
    """Find the common (level, t) at which every block fits.

    ``blocks`` holds ``(paragraphs, usable_width_in, usable_height_in)``. Tries the
    full texts at every size from preferred down to the minimum, then the short
    texts. Returns ``(level, t, fits_per_block)``; when nothing fits entirely the
    last attempt (short text, minimum size) is returned with per-block flags.
    """
    max_level = 1 if any(len(p.cands) > 1 for b in blocks for p in b[0]) else 0
    last_flags: List[bool] = [False] * len(blocks)
    for level in range(max_level + 1):
        for t in FIT_STEPS:
            flags = []
            for paras, w, h in blocks:
                flags.append(_block_height(em_of, paras, w, level, t) <= h + 1e-6)
            last_flags = flags
            if all(flags):
                return level, t, flags
    return max_level, 0.0, last_flags


def _block_height(em_of, paras: Sequence[P], w: float, level: int, t: float) -> float:
    total = 0.0
    for p in paras:
        size = p.size_at(t)
        n, h = C.measure_text(p.text_at(level), w, size, LINE_SPACING, em_of(p))
        if p.single_line and n > 1:
            return 1e9  # numbers must stay on one line
        total += h + p.before / 72.0
    return total


# --------------------------------------------------------------------------
# XML helpers
# --------------------------------------------------------------------------


def _sub(parent, tag: str, **attrs):
    el = etree.SubElement(parent, qn(tag))
    for k, v in attrs.items():
        el.set(k, str(v))
    return el


def _strip_style(shape) -> None:
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def _add_shadow(shape) -> None:
    spPr = shape._element.spPr
    for old in spPr.findall(qn("a:effectLst")):
        spPr.remove(old)
    eff = _sub(spPr, "a:effectLst")
    sh = _sub(eff, "a:outerShdw", blurRad=152400, dist=50800, dir=5400000, algn="t", rotWithShape=0)
    clr = _sub(sh, "a:srgbClr", val="000000")
    _sub(clr, "a:alpha", val=24000)


def _set_alpha(shape, alpha: float) -> None:
    """Set the solid fill's opacity (0..1) on an autoshape."""
    sf = shape._element.spPr.find(qn("a:solidFill"))
    if sf is None:
        return
    clr = sf[0]
    for old in clr.findall(qn("a:alpha")):
        clr.remove(old)
    _sub(clr, "a:alpha", val=int(round(alpha * 100000)))


def _arrow_end(line_shape) -> None:
    ln = line_shape._element.spPr.find(qn("a:ln"))
    if ln is None:
        return
    for old in ln.findall(qn("a:tailEnd")):
        ln.remove(old)
    _sub(ln, "a:tailEnd", type="triangle", w="med", len="med")


def _para_props(paragraph, bullet: bool, indent_in: float, k: float, plain_indent: float = 0.0) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    if bullet:
        pPr.set("marL", str(int(round(indent_in * k * C.EMU_PER_INCH))))
        pPr.set("indent", str(-int(round(indent_in * k * C.EMU_PER_INCH))))
        for old in pPr.findall(qn("a:buNone")):
            pPr.remove(old)
        _sub(pPr, "a:buFont", typeface="Arial")
        _sub(pPr, "a:buChar", char="•")
    else:
        pPr.set("marL", str(int(round(plain_indent * k * C.EMU_PER_INCH))))
        pPr.set("indent", "0")
        _sub(pPr, "a:buNone")


# --------------------------------------------------------------------------
# Deck context (theme, geometry, warnings)
# --------------------------------------------------------------------------


class DeckContext:
    """Shared per-deck state: resolved theme, colours, fonts, geometry and the warning list."""

    def __init__(
        self,
        spec: Dict[str, Any],
        theme: Dict[str, Any],
        prs,
        spec_dir: Path,
        k: float,
        bg_mode: str,
        template_bg: Optional[str] = None,
    ) -> None:
        self.spec = spec
        self.theme = theme
        self.prs = prs
        self.spec_dir = spec_dir
        self.k = k
        self.W = prs.slide_width / C.EMU_PER_INCH / k
        self.H = prs.slide_height / C.EMU_PER_INCH / k
        self.bg_mode = bg_mode  # 'master' | 'slide' | 'inherit'
        self.warnings: List[str] = []
        pal = theme["palette"]
        self.deco = theme["decoration"]
        self.bg = C.normalize_hex(template_bg or pal["background"])
        self.dark = C.relative_luminance(self.bg) < 0.35
        self.text = C.normalize_hex(pal["text"])
        if C.contrast_ratio(self.text, self.bg) < 4.5:
            self.text = "FFFFFF" if self.dark else "000000"
        self.surface = C.normalize_hex(pal.get("surface", pal["background"]))
        self.primary = C.normalize_hex(pal["primary"])
        self.secondary = C.normalize_hex(pal.get("secondary", pal["primary"]))
        self.accent = C.normalize_hex(pal.get("accent", pal["primary"]))
        # Colour for big blocks and bands: a primary that is just the text colour (white on a
        # dark theme, ink on a bold light theme) would give an unstyled block, so use the accent.
        text_like = C.contrast_ratio(self.primary, self.text) < 1.6
        if text_like and (self.dark or theme.get("family") == "bold"):
            self.block, self.block2 = self.accent, self.accent
        else:
            self.block, self.block2 = self.primary, self.secondary
        muted = C.normalize_hex(pal.get("mutedText", pal["text"]))
        if C.contrast_ratio(muted, self.bg) < 4.5:
            muted = C.mix_hex(muted, self.text, 0.45)
        self.muted = muted
        self.rule = C.mix_hex(self.bg, self.text, 0.22)
        self.series = [C.normalize_hex(x) for x in pal.get("chartSeries", [pal["primary"], pal["secondary"], pal["accent"]])]
        self.accent_text_ok = C.accent_text_allowed(theme)
        self.fonts = {
            "heading": C.resolve_font(theme["fonts"]["heading"]),
            "body": C.resolve_font(theme["fonts"]["body"]),
            "mono": C.resolve_font(theme["fonts"].get("mono", theme["fonts"]["body"])),
        }
        hl = self.fonts["heading"].lower()
        self.heading_bold = not (any(h in hl for h in HEAVY_FONTS) or any(s in hl for s in SERIF_FONTS))
        lang = spec.get("language", "en")
        self.lang = lang if "-" in lang else {"en": "en-US", "nb": "nb-NO", "no": "nb-NO", "de": "de-DE", "fr": "fr-FR", "es": "es-ES", "sv": "sv-SE", "da": "da-DK"}.get(lang, lang)
        self.section_counter = 0
        self.media_cursor = 0

    # ---- colour logic --------------------------------------------------
    def fg(self, bg: Optional[str]) -> str:
        """Readable body-text colour on a given background (theme text if contrast allows)."""
        bg = bg or self.bg
        for cand in (self.text, self.bg, "FFFFFF", "000000"):
            if C.contrast_ratio(cand, bg) >= 4.5:
                return cand
        return "FFFFFF" if C.relative_luminance(bg) < 0.4 else "000000"

    def muted_on(self, bg: Optional[str]) -> str:
        bg = bg or self.bg
        if C.contrast_ratio(self.muted, bg) >= 4.5:
            return self.muted
        return C.mix_hex(self.fg(bg), bg, 0.25)

    def accent_text(self, bg: Optional[str]) -> str:
        """Accent colour for text if the theme allows it and it is legible, else a safe stand-in."""
        bg = bg or self.bg
        if self.accent_text_ok and C.contrast_ratio(self.accent, bg) >= 3.0:
            return self.accent
        return self.primary_text(bg)

    def primary_text(self, bg: Optional[str]) -> str:
        bg = bg or self.bg
        if C.contrast_ratio(self.primary, bg) >= 3.0:
            return self.primary
        return self.fg(bg)

    def resolve_color(self, name: str, bg: Optional[str]) -> str:
        if name == "text":
            return self.fg(bg)
        if name == "muted":
            return self.muted_on(bg)
        if name == "accent":
            return self.accent_text(bg)
        if name == "primary":
            return self.primary_text(bg)
        return C.normalize_hex(name)

    def on_fill(self, fill: str) -> str:
        """Text colour readable on a solid fill."""
        for cand in (self.text, self.bg, "FFFFFF", "000000"):
            if C.contrast_ratio(cand, fill) >= 4.5:
                return cand
        return "FFFFFF" if C.relative_luminance(fill) < 0.4 else "000000"

    def em_of(self, p: P) -> float:
        name = self.fonts.get(p.font, self.fonts["body"]).lower()
        em = FONT_EM.get(name, 0.52)
        if p.bold or (p.font == "heading" and self.heading_bold):
            em *= 1.06
        return em

    def warn(self, slide_label: str, msg: str) -> None:
        self.warnings.append("%s: %s" % (slide_label, msg))


# --------------------------------------------------------------------------
# Slide builder
# --------------------------------------------------------------------------


class SlideBuilder:
    """Drawing primitives for one slide. All coordinates are design inches."""

    def __init__(self, ctx: DeckContext, slide, spec_slide: Dict[str, Any], index: int, variant: int) -> None:
        self.ctx = ctx
        self.slide = slide
        self.s = spec_slide
        self.index = index
        self.variant = variant
        self.W = ctx.W
        self.H = ctx.H
        self.label = "slide %d (%s)" % (index + 1, spec_slide.get("id", "?"))
        self.bg = ctx.bg  # colour currently behind the content
        self._n_decor = 0
        self._n_pic = 0
        self.used_chart = False
        self.used_media = False

    # ---- spec accessors ------------------------------------------------
    @property
    def title(self) -> Dict[str, str]:
        return self.s["title"]

    @property
    def summary(self) -> Optional[Dict[str, str]]:
        return self.s.get("summary")

    @property
    def items(self) -> List[Dict[str, Any]]:
        return list(self.s.get("items") or [])

    @property
    def chart(self) -> Optional[Dict[str, Any]]:
        return self.s.get("chart")

    @property
    def media(self) -> List[Dict[str, Any]]:
        return list(self.s.get("media") or [])

    def warn(self, msg: str) -> None:
        self.ctx.warn(self.label, msg)

    # ---- conversions ---------------------------------------------------
    def E(self, v: float) -> int:
        return int(round(v * self.ctx.k * C.EMU_PER_INCH))

    def _name(self, name: Optional[str], prefix: str) -> str:
        if name:
            return name
        self._n_decor += 1
        return "%s %d" % (prefix, self._n_decor)

    # ---- background ----------------------------------------------------
    def set_background(self, fill: str, gradient_to: Optional[str] = None, angle: float = 0.0) -> None:
        """Set the slide background colour (optionally a two-stop gradient when the theme allows)."""
        fill = C.normalize_hex(fill)
        self.bg = fill
        f = self.slide.background.fill
        if gradient_to and self.ctx.deco.get("gradients"):
            f.gradient()
            f.gradient_angle = angle
            f.gradient_stops[0].color.rgb = C.hex_to_rgb(fill)
            f.gradient_stops[1].color.rgb = C.hex_to_rgb(gradient_to)
        else:
            f.solid()
            f.fore_color.rgb = C.hex_to_rgb(fill)

    def apply_base_background(self) -> None:
        """Paint the theme background where the deck's mode requires per-slide backgrounds."""
        if self.ctx.bg_mode == "slide":
            self.set_background(self.ctx.bg)

    # ---- shapes --------------------------------------------------------
    def rect(
        self,
        x: float,
        y: float,
        w: float,
        h: float,
        fill: Optional[str] = None,
        line: Optional[str] = None,
        line_w: float = 1.0,
        shape: str = "rect",
        rounded: bool = False,
        shadow: bool = False,
        alpha: Optional[float] = None,
        gradient_to: Optional[str] = None,
        name: Optional[str] = None,
    ):
        """Add a filled / outlined autoshape. ``rounded`` applies the theme corner radius (if > 0)."""
        ctx = self.ctx
        kinds = {
            "rect": MSO_SHAPE.RECTANGLE,
            "ellipse": MSO_SHAPE.OVAL,
            "chevron": MSO_SHAPE.CHEVRON,
            "pentagon": MSO_SHAPE.PENTAGON,
            "rarrow": MSO_SHAPE.RIGHT_ARROW,
            "trapezoid": MSO_SHAPE.TRAPEZOID,
        }
        mso = kinds[shape]
        radius_pt = float(ctx.deco.get("cornerRadius", 0) or 0)
        use_round = rounded and shape == "rect" and radius_pt > 0
        if use_round:
            mso = MSO_SHAPE.ROUNDED_RECTANGLE
        shp = self.slide.shapes.add_shape(mso, self.E(x), self.E(y), self.E(w), self.E(h))
        _strip_style(shp)
        shp.name = self._name(name, "Decor")
        if use_round:
            shp.adjustments[0] = max(0.0, min(0.5, (radius_pt / 72.0) / max(0.01, min(w, h))))
        if fill is None:
            shp.fill.background()
        elif gradient_to and ctx.deco.get("gradients"):
            shp.fill.gradient()
            shp.fill.gradient_angle = 90.0
            shp.fill.gradient_stops[0].color.rgb = C.hex_to_rgb(fill)
            shp.fill.gradient_stops[1].color.rgb = C.hex_to_rgb(gradient_to)
        else:
            shp.fill.solid()
            shp.fill.fore_color.rgb = C.hex_to_rgb(fill)
            if alpha is not None:
                _set_alpha(shp, alpha)
        if line is None:
            shp.line.fill.background()
        else:
            shp.line.fill.solid()
            shp.line.color.rgb = C.hex_to_rgb(line)
            shp.line.width = Pt(line_w * ctx.k)
        if shadow and ctx.deco.get("shadows"):
            _add_shadow(shp)
        return shp

    def line(self, x1: float, y1: float, x2: float, y2: float, color: str, width_pt: float = 1.0, arrow: bool = False, name: Optional[str] = None):
        """Add a straight connector line; ``arrow`` adds an arrowhead at the end point."""
        c = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, self.E(x1), self.E(y1), self.E(x2), self.E(y2))
        _strip_style(c)
        c.name = self._name(name, "Line")
        c.line.fill.solid()
        c.line.color.rgb = C.hex_to_rgb(color)
        c.line.width = Pt(width_pt * self.ctx.k)
        if arrow:
            _arrow_end(c)
        return c

    def card(self, x: float, y: float, w: float, h: float, fill: Optional[str] = None, name: Optional[str] = None):
        """A surface card: theme surface fill, subtle border on flat themes, shadow when the theme has shadows."""
        ctx = self.ctx
        fill = fill or ctx.surface
        border = None
        if not ctx.deco.get("shadows") and C.contrast_ratio(fill, self.bg) < 1.25:
            border = ctx.rule
        shp = self.rect(x, y, w, h, fill=fill, line=border, line_w=0.75, rounded=True, shadow=True, name=name or self._name(None, "Card"))
        return shp

    def rule_under_title(self, x: float, y: float, w: float) -> None:
        """Header rule honouring the theme ``rules`` style."""
        style = self.ctx.deco.get("rules", "none")
        if style == "thin":
            self.line(x, y, x + w, y, self.ctx.rule, 1.0, name="Header rule")
        elif style == "bold":
            self.rect(x, y - 0.04, 1.4, 0.09, fill=self.ctx.accent, name="Header rule")

    # ---- text ----------------------------------------------------------
    def blocks(self, specs: Sequence[Dict[str, Any]]) -> Tuple[int, float, List[bool]]:
        """Fit and emit several text blocks together so they share one text level and size scale.

        Each spec: ``rect=(x, y, w, h)``, ``paras=[P, ...]`` and optional ``anchor``
        ('t', 'm', 'b'), ``align`` ('l', 'c', 'r'), ``bg`` (colour behind the text), ``name``.
        Returns ``(level, t, fits)``.
        """
        ctx = self.ctx
        entries = []
        for sp in specs:
            x, y, w, h = sp["rect"]
            entries.append((sp["paras"], max(0.1, w - 2 * INSET_X), max(0.1, h - 2 * INSET_Y)))
        if not entries:
            return 0, 1.0, []
        level, t, flags = fit_blocks(ctx.em_of, entries)
        for sp, ok in zip(specs, flags):
            if not ok:
                label = sp.get("name") or "text"
                self.warn("text %r does not fit at the minimum font size: split this slide" % label)
            self._emit(sp, level, t)
        return level, t, flags

    def text(self, x: float, y: float, w: float, h: float, paras: Sequence[P], anchor: str = "t", align: str = "l", bg: Optional[str] = None, name: Optional[str] = None) -> bool:
        """Fit and emit one text block; returns whether it fit."""
        _, _, flags = self.blocks([{"rect": (x, y, w, h), "paras": list(paras), "anchor": anchor, "align": align, "bg": bg, "name": name}])
        return flags[0]

    def _emit(self, sp: Dict[str, Any], level: int, t: float) -> None:
        ctx = self.ctx
        x, y, w, h = sp["rect"]
        paras: Sequence[P] = sp["paras"]
        bg = sp.get("bg") or self.bg
        tb = self.slide.shapes.add_textbox(self.E(x), self.E(y), self.E(w), self.E(h))
        tb.name = self._name(sp.get("name"), "Text")
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = Emu(self.E(INSET_X))
        tf.margin_top = tf.margin_bottom = Emu(self.E(INSET_Y))
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[sp.get("anchor", "t")]
        align = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[sp.get("align", "l")]
        first = True
        for p in paras:
            para = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            para.alignment = align
            if p.before:
                para.space_before = Pt(p.before * ctx.k)
            _para_props(para, p.bullet, 0.3 if p.bullet else 0.0, ctx.k, p.indent)
            run = para.add_run()
            run.text = p.text_at(level)
            f = run.font
            f.size = Pt(p.size_at(t) * ctx.k)
            f.name = ctx.fonts.get(p.font, ctx.fonts["body"])
            bold = p.bold or (p.font == "heading" and ctx.heading_bold and p.kind in ("title", "stat", "hero"))
            f.bold = bool(bold)
            f.italic = bool(p.italic)
            color = ctx.resolve_color(p.color, bg)
            rPr = run._r.get_or_add_rPr()
            rPr.set("lang", ctx.lang)
            if p.outline:
                for old in list(rPr):
                    rPr.remove(old)
                ln = _sub(rPr, "a:ln", w=int(round(1.5 * ctx.k * 12700)))
                sf = _sub(ln, "a:solidFill")
                _sub(sf, "a:srgbClr", val=color)
                _sub(rPr, "a:noFill")
                _sub(rPr, "a:latin", typeface=ctx.fonts.get(p.font, ctx.fonts["body"]))
            else:
                f.color.rgb = C.hex_to_rgb(color)

    def plain_text(self, x, y, w, h, text: str, size_pt: float, color: str, font: str = "body", align: str = "l", anchor: str = "t", bold: bool = False, name: Optional[str] = None) -> None:
        """Emit non-fitted decorative text (glyphs, huge ornamental numbers). Not subject to minimum sizes."""
        sp = {"rect": (x, y, w, h), "paras": [P([text], "caption", pref=size_pt, bold=bold, color=color, font=font, min_pt=size_pt)], "anchor": anchor, "align": align, "name": name}
        self._emit(sp, 0, 1.0)

    # ---- header --------------------------------------------------------
    def header(self, x: float = M, w: Optional[float] = None, with_summary: bool = True, bg: Optional[str] = None, color: str = "text") -> float:
        """Title (+ rule, + summary). Returns the y where content may start."""
        w = w if w is not None else self.W - 2 * M
        bg = bg or self.bg
        t = self.title
        title_h = 1.15
        self.text(x, 0.45, w, title_h, [P([t["full"], t.get("short", "")], "title", pref=38, font="heading", color=color)], anchor="b", bg=bg, name="Title")
        self.rule_under_title(x, 1.72, w)
        y = 1.92
        sm = self.summary
        if with_summary and sm:
            self.text(x, y, w, 0.95, [P([sm["full"], sm.get("short", "")], "body", pref=22, color="muted")], bg=bg, name="Summary")
            y += 1.0
        return y + 0.1

    # ---- composite pieces ----------------------------------------------
    def badge(self, cx: float, cy: float, d: float, label: str, fill: Optional[str] = None, name: Optional[str] = None) -> None:
        """Numbered circle (or square for radius-0 themes keeps a circle: badges are always round)."""
        fill = fill or self.ctx.block
        self.rect(cx - d / 2, cy - d / 2, d, d, fill=fill, shape="ellipse", name=(name or "Badge") + " shape")
        inner = d * 0.78
        self.text(
            cx - inner / 2, cy - inner / 2, inner, inner,
            [P([label], "body", pref=22, bold=True, color=self.ctx.on_fill(fill), font="heading")],
            anchor="m", align="c", bg=fill, name=name or "Badge",
        )

    def chip(self, x: float, y: float, w: float, h: float, label: str, fill: str, name: Optional[str] = None) -> None:
        """Small label chip (caption size) on a solid fill."""
        self.rect(x, y, w, h, fill=fill, rounded=True, name=(name or "Chip") + " shape")
        self.text(x, y, w, h, [P([label], "caption", pref=14, bold=True, color=self.ctx.on_fill(fill))], anchor="m", align="c", bg=fill, name=name or "Chip")

    def initials(self, label: str) -> str:
        parts = [p for p in re.split(r"[\s\-]+", label.strip()) if p]
        if not parts:
            return "?"
        if len(parts) == 1:
            return parts[0][:2].upper()
        return (parts[0][0] + parts[-1][0]).upper()

    # ---- pictures ------------------------------------------------------
    def picture(self, media: Optional[Dict[str, Any]], x: float, y: float, w: float, h: float, fit: Optional[str] = None, ellipse: bool = False, name: Optional[str] = None) -> bool:
        """Place an image in the box with cover (crop) or contain fit. Draws a labelled placeholder if the file is unusable."""
        ctx = self.ctx
        alt = (media or {}).get("alt", "")
        fit = fit or (media or {}).get("fit", "cover")
        data = _load_image(ctx, media) if media else None
        if data is None:
            if media:
                self.warn("image %r could not be loaded; placeholder drawn" % media.get("path"))
            self.rect(x, y, w, h, fill=ctx.surface, line=ctx.rule, shape="ellipse" if ellipse else "rect", name=(name or "Picture") + " placeholder")
            if alt and not ellipse:
                self.text(x + 0.1, y + h / 2 - 0.3, w - 0.2, 0.6, [P([alt], "caption", pref=12, color="muted")], anchor="m", align="c", bg=ctx.surface, name=(name or "Picture") + " caption")
            return False
        blob, iw, ih = data
        box_ratio = w / h
        img_ratio = iw / ih
        px, py, pw, ph = x, y, w, h
        crop = (0.0, 0.0, 0.0, 0.0)
        if fit == "contain":
            if img_ratio > box_ratio:
                ph = w / img_ratio
                py = y + (h - ph) / 2
            else:
                pw = h * img_ratio
                px = x + (w - pw) / 2
        else:
            if img_ratio > box_ratio:
                keep = box_ratio / img_ratio
                crop = ((1 - keep) / 2, 0.0, (1 - keep) / 2, 0.0)
            else:
                keep = img_ratio / box_ratio
                crop = (0.0, (1 - keep) / 2, 0.0, (1 - keep) / 2)
        pic = self.slide.shapes.add_picture(io.BytesIO(blob), self.E(px), self.E(py), self.E(pw), self.E(ph))
        pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = crop
        self._n_pic += 1
        pic.name = name or "Picture %d" % self._n_pic
        pic._element.nvPicPr.cNvPr.set("descr", alt)
        if ellipse:
            pic.auto_shape_type = MSO_SHAPE.OVAL
        return True

    # ---- charts --------------------------------------------------------
    def add_chart(self, chart: Dict[str, Any], x: float, y: float, w: float, h: float, type_override: Optional[str] = None, legend: Optional[bool] = None, name: str = "Chart") -> None:
        """Add a native, theme-styled chart (bar, column, line, pie, doughnut, stacked_bar)."""
        _add_native_chart(self, chart, x, y, w, h, type_override, legend, name)

    def callout(self, x: float, y: float, w: float, h: float, texts: Sequence[str], pref: float = 22, name: str = "Insight") -> None:
        """Insight callout: surface block, accent bar on the left, text inside."""
        ctx = self.ctx
        self.card(x, y, w, h, name=name + " card")
        self.rect(x, y, 0.1, h, fill=ctx.accent, name=name + " bar")
        self.text(x + 0.25, y + 0.05, w - 0.4, h - 0.1, [P(list(texts), "body", pref=pref, color="text")], anchor="m", bg=ctx.surface, name=name)


# --------------------------------------------------------------------------
# Images
# --------------------------------------------------------------------------


def _load_image(ctx: DeckContext, media: Dict[str, Any]) -> Optional[Tuple[bytes, int, int]]:
    """Read an image, apply the theme image treatment, return ``(bytes, width_px, height_px)`` or None."""
    path = media.get("path", "")
    if re.match(r"^[a-z]+://", path):
        return None
    p = Path(path)
    if not p.is_absolute():
        p = ctx.spec_dir / p
    if not p.is_file():
        return None
    try:
        from PIL import Image, ImageOps

        with Image.open(p) as im:
            im.load()
            treat = ctx.deco.get("imageTreatment", "natural")
            fmt = (im.format or "PNG").upper()
            if treat == "natural" and fmt in ("PNG", "JPEG", "GIF", "BMP", "TIFF"):
                return p.read_bytes(), im.width, im.height
            rgb = im.convert("RGB")
            gray = ImageOps.grayscale(rgb)
            if treat == "duotone":
                dark_end = C.mix_hex(ctx.bg, "000000", 0.2) if ctx.dark else ctx.text
                light_end = ctx.primary if ctx.dark else ctx.bg
                out = ImageOps.colorize(gray, black="#" + dark_end, white="#" + light_end)
            elif treat == "grayscale":
                out = gray.convert("RGB")
            else:
                out = rgb
            buf = io.BytesIO()
            out.save(buf, format="PNG", optimize=False)
            return buf.getvalue(), out.width, out.height
    except Exception:
        return None


# --------------------------------------------------------------------------
# Charts
# --------------------------------------------------------------------------


def _chart_text_color(ctx: DeckContext) -> str:
    return ctx.muted_on(ctx.bg) if ctx.fg(ctx.bg) == ctx.text else ctx.fg(ctx.bg)


def _add_native_chart(sb: SlideBuilder, spec_chart: Dict[str, Any], x, y, w, h, type_override, legend, name) -> None:
    ctx = sb.ctx
    ctype = type_override or spec_chart.get("type", "column")
    xl = CHART_TYPES[ctype]
    cats = list(spec_chart["categories"])
    series = [dict(name=s["name"], values=list(s["values"])) for s in spec_chart["series"]]
    horizontal = ctype in ("bar", "stacked_bar")
    if ctype in ("pie", "doughnut"):
        series = series[:1]
    if horizontal:  # first category at the top
        cats = cats[::-1]
        for s in series:
            s["values"] = s["values"][::-1]
        if ctype == "bar":
            series = series[::-1]
    cd = CategoryChartData()
    cd.categories = cats
    for s in series:
        cd.add_series(s["name"], s["values"])
    gf = sb.slide.shapes.add_chart(xl, sb.E(x), sb.E(y), sb.E(w), sb.E(h), cd)
    gf.name = name
    chart = gf.chart
    k = ctx.k
    chart.has_title = False
    txt = _chart_text_color(ctx)
    chart.font.size = Pt(14 * k)
    chart.font.name = ctx.fonts["body"]
    chart.font.color.rgb = C.hex_to_rgb(txt)
    multi = len(series) > 1
    want_legend = legend if legend is not None else (multi or ctype in ("pie", "doughnut"))
    chart.has_legend = bool(want_legend)
    if want_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(14 * k)
        chart.legend.font.color.rgb = C.hex_to_rgb(txt)
    plot = chart.plots[0]
    # series colours: the chart's logical order maps to palette order
    order = list(range(len(series)))
    if ctype == "bar":
        order = order[::-1]
    pal = ctx.series
    if ctype in ("pie", "doughnut"):
        plot.vary_by_categories = True
        ser = plot.series[0]
        for i in range(len(cats)):
            pt = ser.points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = C.hex_to_rgb(pal[i % len(pal)])
            pt.format.line.fill.solid()
            pt.format.line.color.rgb = C.hex_to_rgb(ctx.bg)
            pt.format.line.width = Pt(1.5 * k)
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.show_percentage = True
        dl.show_value = False
        dl.show_category_name = False
        dl.number_format = "0%"
        dl.number_format_is_linked = False
        dl.font.size = Pt(14 * k)
        dl.font.bold = True
        cols = [pal[i % len(pal)] for i in range(len(cats))]
        white = sum(C.contrast_ratio("FFFFFF", c) for c in cols)
        black = sum(C.contrast_ratio("111111", c) for c in cols)
        dl.font.color.rgb = C.hex_to_rgb("FFFFFF" if white >= black else "111111")
        if ctype == "pie":
            dl.position = XL_LABEL_POSITION.INSIDE_END
    else:
        for logical, ser in zip(order, plot.series):
            col = pal[logical % len(pal)]
            if ctype == "line":
                ser.format.line.fill.solid()
                ser.format.line.color.rgb = C.hex_to_rgb(col)
                ser.format.line.width = Pt(3 * k)
                ser.smooth = False
                ser.marker.style = XL_MARKER_STYLE.CIRCLE
                ser.marker.size = 9
                ser.marker.format.fill.solid()
                ser.marker.format.fill.fore_color.rgb = C.hex_to_rgb(col)
                ser.marker.format.line.fill.solid()
                ser.marker.format.line.color.rgb = C.hex_to_rgb(ctx.bg)
            else:
                ser.format.fill.solid()
                ser.format.fill.fore_color.rgb = C.hex_to_rgb(col)
                ser.format.line.fill.background()
        if ctype in ("bar", "column", "stacked_bar"):
            plot.gap_width = 55 if ctype != "stacked_bar" else 40
            if ctype == "stacked_bar":
                plot.overlap = 100
            else:
                plot.overlap = -10 if multi else 0
        # axes
        ca, va = chart.category_axis, chart.value_axis
        ca.has_major_gridlines = False
        ca.major_tick_mark = XL_TICK_MARK.NONE
        ca.format.line.fill.solid()
        ca.format.line.color.rgb = C.hex_to_rgb(ctx.rule)
        ca.tick_labels.font.size = Pt(14 * k)
        ca.tick_labels.font.color.rgb = C.hex_to_rgb(txt)
        va.has_major_gridlines = True
        va.major_gridlines.format.line.fill.solid()
        va.major_gridlines.format.line.color.rgb = C.hex_to_rgb(C.mix_hex(ctx.bg, ctx.text, 0.12))
        va.major_gridlines.format.line.width = Pt(0.75 * k)
        va.format.line.fill.background()
        va.major_tick_mark = XL_TICK_MARK.NONE
        va.tick_labels.font.size = Pt(14 * k)
        va.tick_labels.font.color.rgb = C.hex_to_rgb(txt)
        if not multi and len(cats) <= 8 and ctype != "stacked_bar":
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.show_value = True
            dl.font.size = Pt(14 * k)
            dl.font.bold = True
            dl.font.color.rgb = C.hex_to_rgb(ctx.fg(ctx.bg))
            dl.position = XL_LABEL_POSITION.ABOVE if ctype == "line" else XL_LABEL_POSITION.OUTSIDE_END
    # transparent chart area
    cs = chart._chartSpace
    for old in cs.findall(qn("c:spPr")):
        cs.remove(old)
    sp = etree.Element(qn("c:spPr"))
    _sub(sp, "a:noFill")
    ln = _sub(sp, "a:ln")
    _sub(ln, "a:noFill")
    cs.find(qn("c:chart")).addnext(sp)
    # alt text
    gf._element.nvGraphicFramePr.cNvPr.set("descr", spec_chart.get("insight", ""))


def chart_from_items(items: Sequence[Dict[str, Any]], kind: str = "doughnut") -> Optional[Dict[str, Any]]:
    """Build a chart spec from items that carry numeric ``value`` fields (None if fewer than two do)."""
    vals = [(it["label"], it["value"]) for it in items if isinstance(it.get("value"), (int, float))]
    if len(vals) < 2:
        return None
    return {"type": kind, "categories": [v[0] for v in vals], "series": [{"name": "Value", "values": [v[1] for v in vals]}], "insight": ""}


# --------------------------------------------------------------------------
# Layout helpers shared by the role renderers
# --------------------------------------------------------------------------


def _tp(t: Dict[str, str], kind: str = "title", **kw) -> P:
    """Paragraph from a {full, short} text object."""
    return P([t["full"], t.get("short", "")], kind, **kw)


def _bottom(sb: SlideBuilder) -> float:
    return sb.H - 0.55


def _grid(x: float, y: float, w: float, h: float, cols: int, rows: int, gx: float = 0.3, gy: float = 0.3) -> List[Tuple[float, float, float, float]]:
    cw = (w - (cols - 1) * gx) / cols
    rh = (h - (rows - 1) * gy) / rows
    return [(x + c * (cw + gx), y + r * (rh + gy), cw, rh) for r in range(rows) for c in range(cols)]


def _item_paras(it: Dict[str, Any], label_pref: float = 22, detail_pref: float = 20, label_color: str = "text", detail_color: str = "muted", gap: float = 2.0, align_value: bool = False) -> List[P]:
    ps = [P([it["label"]], "body", pref=label_pref, bold=True, color=label_color)]
    d = it.get("detail")
    if d:
        ps.append(P([d["full"], d.get("short", "")], "body", pref=detail_pref, color=detail_color, before=gap))
    return ps


def _detail_para(it: Dict[str, Any], pref: float = 20, color: str = "muted") -> Optional[P]:
    d = it.get("detail")
    return P([d["full"], d.get("short", "")], "body", pref=pref, color=color) if d else None


def _value_text(it: Dict[str, Any]) -> str:
    return C.format_item_value(it)


def _priority_key(it: Dict[str, Any]) -> Tuple[int, int]:
    p = it.get("priority")
    return (p if isinstance(p, int) else 99, 0)


def _chart_with_insight(sb: SlideBuilder, chart: Dict[str, Any], x: float, y: float, w: float, h: float, type_override: Optional[str] = None) -> None:
    """Chart on top and the insight sentence as a callout below it."""
    ins = chart.get("insight") if h >= 3.4 else None
    ch = h - 1.05 if ins else h
    sb.add_chart(chart, x, y, w, ch, type_override)
    if ins:
        sb.callout(x, y + ch + 0.1, w, 0.95, [ins], pref=20)
    sb.used_chart = True


def _stack_rows(sb: SlideBuilder, items: Sequence[Dict[str, Any]], x: float, y: float, w: float, h: float, label_pref: float = 22, detail_pref: float = 20, max_row: float = 1.5, rules: bool = True, bg: Optional[str] = None, prefix: str = "Item") -> None:
    """Items as stacked label+detail rows separated by hairlines."""
    n = len(items)
    if n == 0:
        return
    rh = min(max_row, h / n)
    specs = []
    for i, it in enumerate(items):
        ry = y + i * rh
        if rules and i > 0:
            sb.line(x, ry, x + w, ry, sb.ctx.rule, 0.75, name="%s rule %d" % (prefix, i))
        specs.append({"rect": (x, ry + 0.04, w, rh - 0.08), "paras": _item_paras(it, label_pref, detail_pref), "anchor": "m", "bg": bg, "name": "%s %s" % (prefix, it["id"])})
    sb.blocks(specs)


def _dist_entries(sb: SlideBuilder) -> List[Dict[str, Any]]:
    """Entries (label, number, text, detail) for distribution slides from items or from the chart."""
    ctx = sb.ctx
    out: List[Dict[str, Any]] = []
    items = sb.items
    chart = sb.chart
    if items:
        for i, it in enumerate(items):
            out.append({"label": it["label"], "value": it.get("value"), "text": _value_text(it), "detail": it.get("detail"), "color": ctx.series[i % len(ctx.series)], "id": it["id"]})
        if chart and chart.get("series") and all(e["value"] is None for e in out) and len(chart["series"][0]["values"]) == len(out):
            for e, v in zip(out, chart["series"][0]["values"]):
                e["value"] = v
        return out
    if chart:
        vals = chart["series"][0]["values"]
        total = float(sum(vals)) or 1.0
        for i, (c, v) in enumerate(zip(chart["categories"], vals)):
            out.append({"label": c, "value": v, "text": "%d%%" % round(v / total * 100) if total else "", "detail": None, "color": ctx.series[i % len(ctx.series)], "id": "c%d" % i})
    return out


# --------------------------------------------------------------------------
# Role renderers
# --------------------------------------------------------------------------


def role_cover(sb: SlideBuilder) -> None:
    ctx, W, H, v = sb.ctx, sb.W, sb.H, sb.variant
    sm, med = sb.summary, sb.media
    owner = ctx.spec.get("owner")
    ttl = _tp(sb.title, "title", pref=54, font="heading")
    if v == 1:
        bw = 2.6 if W > 11 else 2.0
        bx = W - bw
        if med and sb.picture(med[0], bx, 0, bw, H, "cover", name="Cover picture"):
            sb.used_media = True
        else:
            sb.rect(bx, 0, bw, H, fill=ctx.accent, gradient_to=ctx.block, name="Accent band")
        tw = bx - M - 0.5
        sb.text(M, 1.7, tw, 2.8, [ttl], anchor="b", name="Title")
        sb.rect(M, 4.7, 1.2, 0.09, fill=ctx.accent, name="Accent rule")
        if sm:
            sb.text(M, 4.95, tw, 1.3, [_tp(sm, "body", pref=24, color="muted")], name="Summary")
        if owner:
            sb.text(M, H - 0.95, tw, 0.45, [P([owner], "caption", pref=14, color="muted")], anchor="m", name="Owner")
    elif v == 2:
        if med:
            if sb.picture(med[0], 0, 0, W, H, "cover", name="Cover backdrop"):
                sb.used_media = True
                sb.rect(0, 0, W, H, fill=ctx.bg, alpha=0.86, name="Backdrop veil")
        sb.text(M, 1.5, W - 2 * M, 2.6, [ttl], anchor="b", align="c", name="Title")
        sb.rect(W / 2 - 0.7, 4.3, 1.4, 0.08, fill=ctx.accent, name="Accent rule")
        if sm:
            sb.text(M + 1.0, 4.6, W - 2 * M - 2.0, 1.4, [_tp(sm, "body", pref=24, color="muted")], align="c", name="Summary")
        if owner:
            sb.text(M, H - 0.95, W - 2 * M, 0.45, [P([owner], "caption", pref=14, color="muted")], anchor="m", align="c", name="Owner")
    else:
        hx = W / 2
        if med and sb.picture(med[0], hx, 0, W - hx, H, "cover", name="Cover picture"):
            sb.used_media = True
        else:
            sb.rect(hx, 0, W - hx, H, fill=ctx.block, gradient_to=ctx.block2, name="Primary block")
        tw = hx - M - 0.5
        sb.text(M, 1.5, tw, 3.0, [_tp(sb.title, "title", pref=50, font="heading")], anchor="b", name="Title")
        sb.rect(M, 4.7, 1.2, 0.09, fill=ctx.accent, name="Accent rule")
        if sm:
            sb.text(M, 4.95, tw, 1.4, [_tp(sm, "body", pref=22, color="muted")], name="Summary")
        if owner:
            sb.text(M, H - 0.95, tw, 0.45, [P([owner], "caption", pref=14, color="muted")], anchor="m", name="Owner")


def role_statement(sb: SlideBuilder) -> None:
    ctx, W, H, v = sb.ctx, sb.W, sb.H, sb.variant
    sm = sb.summary
    if v == 1:
        sb.text(M, 0.9, W - 2 * M, 3.8, [_tp(sb.title, "title", pref=60, font="heading")], anchor="m", align="c", name="Title")
        if sm:
            sb.text(M + 1.0, 4.95, W - 2 * M - 2.0, 1.5, [_tp(sm, "body", pref=22, color="muted")], align="c", name="Summary")
    else:
        sb.plain_text(M, 0.3, 2.2, 2.7, "“", 130, ctx.accent_text(ctx.bg), font="heading", name="Quote mark")
        sb.rect(M, 3.1, 0.09, 3.5, fill=ctx.accent, name="Vertical rule")
        tx, tw = M + 0.5, W - 2 * M - 0.5
        sb.text(tx, 3.0, tw, 2.4, [_tp(sb.title, "title", pref=46, font="heading")], anchor="t", name="Title")
        if sm:
            sb.text(tx, 5.5, tw, 1.4, [_tp(sm, "body", pref=22, color="muted")], name="Summary")


def role_breakdown(sb: SlideBuilder) -> None:
    ctx, W, H, v = sb.ctx, sb.W, sb.H, sb.variant
    top, bot = sb.header(), _bottom(sb)
    items = sb.items
    n = len(items)
    if not n:
        return
    if v == 1:
        rh = min(1.0, (bot - top) / n)
        specs = []
        for i, it in enumerate(items):
            y = top + i * rh
            sb.line(M, y, W - M, y, ctx.rule, 0.75, name="Row rule %d" % (i + 1))
            specs.append({"rect": (M, y, 1.2, rh), "paras": [P(["%02d" % (i + 1)], "stat", pref=32, bold=True, color="accent", font="heading")], "anchor": "m", "name": "Index %d" % (i + 1)})
            specs.append({"rect": (M + 1.3, y, 4.4, rh), "paras": [P([it["label"]], "body", pref=22, bold=True)], "anchor": "m", "name": "Item %s label" % it["id"]})
            dp = _detail_para(it)
            if dp:
                specs.append({"rect": (M + 5.8, y, W - M - (M + 5.8), rh), "paras": [dp], "anchor": "m", "name": "Item %s detail" % it["id"]})
        sb.blocks(specs)
    elif v == 2:
        cols = {1: 1, 2: 2, 3: 3, 4: 2}.get(n, 3)
        rows = int(math.ceil(n / cols))
        rects = _grid(M, top, W - 2 * M, bot - top, cols, rows, 0.3, 0.3)
        specs = []
        for i, (it, (x, y, w, h)) in enumerate(zip(items, rects)):
            sb.card(x, y, w, h, name="Card %s" % it["id"])
            sb.badge(x + 0.5, y + 0.55, 0.5, str(i + 1), name="Badge %d" % (i + 1))
            specs.append({"rect": (x + 0.9, y + 0.1, w - 1.05, 0.9), "paras": [P([it["label"]], "body", pref=22, bold=True)], "anchor": "m", "bg": ctx.surface, "name": "Item %s label" % it["id"]})
            dp = _detail_para(it)
            if dp:
                specs.append({"rect": (x + 0.2, y + 1.05, w - 0.4, h - 1.15), "paras": [dp], "bg": ctx.surface, "name": "Item %s detail" % it["id"]})
        sb.blocks(specs)
    else:
        gap = 0.08
        cw = (W - 2 * M - (n - 1) * gap) / n
        hh = 1.25
        specs = []
        for i, it in enumerate(items):
            x = M + i * (cw + gap)
            sb.rect(x, top, cw, hh, fill=ctx.block, gradient_to=ctx.block2, name="Segment %s" % it["id"])
            specs.append({"rect": (x + 0.05, top, cw - 0.1, hh), "paras": [P([it["label"]], "body", pref=22, bold=True, color=ctx.on_fill(ctx.block))], "anchor": "m", "bg": ctx.block, "name": "Item %s label" % it["id"]})
            dp = _detail_para(it)
            if dp:
                specs.append({"rect": (x + 0.05, top + hh + 0.15, cw - 0.1, bot - top - hh - 0.15), "paras": [dp], "name": "Item %s detail" % it["id"]})
        sb.blocks(specs)


def role_transition(sb: SlideBuilder) -> None:
    ctx, W, H, v = sb.ctx, sb.W, sb.H, sb.variant
    sm = sb.summary
    ctx.section_counter += 1
    num = "%02d" % ctx.section_counter
    if v == 1:
        sb.set_background(ctx.block, gradient_to=ctx.block2, angle=0)
        on = ctx.on_fill(ctx.block)
        soft = C.mix_hex(on, ctx.block, 0.3)
        sb.plain_text(W - 4.6, 1.6, 4.0, 3.2, num, 150, soft, font="heading", align="r", anchor="m", bold=True, name="Section number")
        tw = W - 2 * M - 4.6
        sb.text(M, 2.0, tw, 2.6, [_tp(sb.title, "title", pref=54, font="heading", color=on)], anchor="b", bg=ctx.block, name="Title")
        if sm:
            sb.text(M, 4.8, tw, 1.5, [_tp(sm, "body", pref=22, color=soft)], bg=ctx.block, name="Summary")
    else:
        sb.text(M, 2.0, W - 2 * M, 2.2, [_tp(sb.title, "title", pref=50, font="heading")], anchor="b", align="c", name="Title")
        sb.rect(W / 2 - 0.8, 4.4, 1.6, 0.09, fill=ctx.accent, name="Accent underline")
        if sm:
            sb.text(M + 1.5, 4.7, W - 2 * M - 3.0, 1.4, [_tp(sm, "body", pref=22, color="muted")], align="c", name="Summary")


def role_context(sb: SlideBuilder) -> None:
    ctx, W, H, v = sb.ctx, sb.W, sb.H, sb.variant
    sm, items = sb.summary, sb.items
    n = len(items)
    top, bot = sb.header(with_summary=False), _bottom(sb)
    full = W - 2 * M
    if v == 1:
        lw = full * 0.55 - 0.25 if (sm and n) else full
        if sm:
            sb.text(M, top, lw, bot - top, [_tp(sm, "body", pref=26)], name="Summary")
        if n:
            rx = M + (full * 0.55 if sm else 0)
            rw = full - (full * 0.55 if sm else 0)
            if sm:
                sb.line(rx - 0.15, top, rx - 0.15, bot, ctx.rule, 0.75, name="Column rule")
            _stack_rows(sb, items, rx, top, rw, bot - top, max_row=1.7)
    elif v == 2:
        y = top
        if sm:
            sb.card(M, y, full, 1.45, name="Summary card")
            sb.text(M + 0.3, y + 0.05, full - 0.6, 1.35, [_tp(sm, "body", pref=24)], anchor="m", bg=ctx.surface, name="Summary")
            y += 1.7
        if n:
            rects = _grid(M, y, full, bot - y, n, 1, 0.3, 0.3)
            specs = []
            for it, (x, yy, w, h) in zip(items, rects):
                sb.card(x, yy, w, h, name="Card %s" % it["id"])
                sb.rect(x, yy, w, 0.09, fill=ctx.accent, name="Card bar %s" % it["id"])
                specs.append({"rect": (x + 0.2, yy + 0.25, w - 0.4, h - 0.4), "paras": _item_paras(it), "bg": ctx.surface, "name": "Item %s" % it["id"]})
            sb.blocks(specs)
    else:
        lw = full * 0.46
        if sm:
            sb.text(M, top, lw, bot - top, [_tp(sm, "body", pref=26)], name="Summary")
        rx = M + (lw + 0.4 if sm else 0)
        rw = full - (lw + 0.4 if sm else 0)
        if n:
            first = items[0]
            ph = min(2.5, (bot - top) * (0.55 if n > 1 else 1.0))
            sb.card(rx, top, rw, ph, name="Pull-fact card")
            sb.rect(rx, top, 0.1, ph, fill=ctx.accent, name="Pull-fact bar")
            paras = [P([first["label"]], "stat", pref=54, bold=True, font="heading", color="accent")]
            dp = _detail_para(first)
            if dp:
                dp.before = 4
                paras.append(dp)
            sb.text(rx + 0.3, top + 0.1, rw - 0.5, ph - 0.2, paras, anchor="m", bg=ctx.surface, name="Pull-fact")
            if n > 1:
                _stack_rows(sb, items[1:], rx, top + ph + 0.2, rw, bot - (top + ph + 0.2), max_row=1.3)


def _tile(sb: SlideBuilder, rect, it: Dict[str, Any], kind: str, pref: float, specs: List[Dict[str, Any]], centered: bool = False) -> None:
    ctx = sb.ctx
    x, y, w, h = rect
    sb.card(x, y, w, h, name="Tile %s" % it["id"])
    val = _value_text(it)
    paras: List[P] = []
    if val:
        paras.append(P([val], kind, pref=pref, bold=True, font="heading", color="primary", single_line=True))
        paras.append(P([it["label"]], "body", pref=22, bold=True, before=2))
    else:
        paras.append(P([it["label"]], "stat", pref=36, bold=True, font="heading", color="primary"))
    dp = _detail_para(it)
    if dp:
        dp.before = 2
        paras.append(dp)
    specs.append({"rect": (x + 0.2, y + 0.1, w - 0.4, h - 0.2), "paras": paras, "anchor": "m", "align": "c" if centered else "l", "bg": ctx.surface, "name": "Tile %s text" % it["id"]})


def _tiles_flow(sb: SlideBuilder, items, x, y, w, h, kind="hero", pref=66, cols: Optional[int] = None) -> None:
    n = len(items)
    if cols is None:
        cols = n if n <= 4 else 3
    rows = int(math.ceil(n / cols))
    rects = _grid(x, y, w, h, cols, rows, 0.25, 0.25)
    specs: List[Dict[str, Any]] = []
    for it, r in zip(items, rects):
        _tile(sb, r, it, kind, pref, specs)
    sb.blocks(specs)


def role_metrics(sb: SlideBuilder) -> None:
    ctx, W, H, v = sb.ctx, sb.W, sb.H, sb.variant
    top, bot = sb.header(), _bottom(sb)
    items, chart = sb.items, sb.chart
    n = len(items)
    full = W - 2 * M
    if v == 3 and not chart:
        v = 1
    if v == 1:
        if chart and n <= 4:
            th = 2.1
            _tiles_flow(sb, items, M, top, full, th)
            _chart_with_insight(sb, chart, M, top + th + 0.2, full, bot - (top + th + 0.2))
        else:
            if chart:
                sb.warn("chart not drawn: metrics variant 1 shows a chart only with 4 or fewer tiles; use layoutVariant 3")
            _tiles_flow(sb, items, M, top, full, min(bot - top, 3.4 if n <= 4 else 4.4))
    elif v == 2:
        order = sorted(range(n), key=lambda i: (items[i].get("priority") if isinstance(items[i].get("priority"), int) else 99, i))
        hero = items[order[0]]
        rest = [items[i] for i in range(n) if i != order[0]]
        specs: List[Dict[str, Any]] = []
        if not rest:
            _tile(sb, (M, top, full, bot - top), hero, "hero", 110, specs, centered=True)
        else:
            hw = 5.0 if W > 11 else 4.0
            _tile(sb, (M, top, hw, bot - top), hero, "hero", 96, specs)
            cols = 2 if len(rest) > 1 else 1
            rows = int(math.ceil(len(rest) / cols))
            rects = _grid(M + hw + 0.3, top, full - hw - 0.3, bot - top, cols, rows, 0.25, 0.25)
            for it, r in zip(rest, rects):
                _tile(sb, r, it, "stat", 40, specs)
        sb.blocks(specs)
    else:
        lw = full * 0.48
        specs = []
        if n <= 3:
            rows_r = _grid(M, top, lw, bot - top, 1, n, 0.25, 0.2)
            for it, (x, y, w, h) in zip(items, rows_r):
                sb.card(x, y, w, h, name="Tile %s" % it["id"])
                val = _value_text(it)
                vw = 2.2 if val else 0.0
                if val:
                    specs.append({"rect": (x + 0.15, y, vw, h), "paras": [P([val], "stat", pref=40, bold=True, font="heading", color="primary", single_line=True)], "anchor": "m", "bg": ctx.surface, "name": "Tile %s value" % it["id"]})
                specs.append({"rect": (x + 0.15 + vw, y + 0.05, w - vw - 0.3, h - 0.1), "paras": _item_paras(it, 22, 18, gap=1), "anchor": "m", "bg": ctx.surface, "name": "Tile %s text" % it["id"]})
        else:
            cols = 2
            rows = int(math.ceil(n / cols))
            rects = _grid(M, top, lw, bot - top, cols, rows, 0.25, 0.25)
            for it, r in zip(items, rects):
                _tile(sb, r, it, "stat", 40, specs)
        sb.blocks(specs)
        _chart_with_insight(sb, chart, M + lw + 0.3, top, full - lw - 0.3, bot - top)


def role_trend(sb: SlideBuilder) -> None:
    ctx, W, H, v = sb.ctx, sb.W, sb.H, sb.variant
    top, bot = sb.header(), _bottom(sb)
    items, chart = sb.items, sb.chart
    n = len(items)
    full = W - 2 * M
    if v in (2,) and not chart:
        v = 1
    if v in (1, 3) and not items and chart:
        v = 2
    if not items and not chart:
        return
    if v == 1:
        ymid = top + (bot - top) / 2
        sb.line(M, ymid, W - M, ymid, ctx.rule, 3, name="Timeline axis")
        step = full / n
        bw = min(3.6, step * 1.6) if n > 1 else 5.0
        specs = []
        for i, it in enumerate(items):
            cx = M + (i + 0.5) * step
            sb.rect(cx - 0.15, ymid - 0.15, 0.3, 0.3, fill=ctx.block, line=ctx.bg, line_w=2, shape="ellipse", name="Marker %s" % it["id"])
            bx = max(M, min(cx - bw / 2, W - M - bw))
            above = i % 2 == 0
            paras = []
            val = _value_text(it)
            if val:
                paras.append(P([val], "body", pref=22, bold=True, color="primary"))
            paras += _item_paras(it, 22 if not val else 20, 20)
            if above:
                rect = (bx, top, bw, ymid - 0.3 - top)
                sb.line(cx, ymid - 0.3, cx, ymid - 0.15, ctx.rule, 1.0, name="Stem %s" % it["id"])
            else:
                rect = (bx, ymid + 0.3, bw, bot - (ymid + 0.3))
                sb.line(cx, ymid + 0.15, cx, ymid + 0.3, ctx.rule, 1.0, name="Stem %s" % it["id"])
            specs.append({"rect": rect, "paras": paras, "anchor": "b" if above else "t", "align": "c", "name": "Item %s" % it["id"]})
        sb.blocks(specs)
        if chart:
            sb.warn("chart not drawn: trend variant 1 is a timeline; use layoutVariant 2 or 3")
    elif v == 2:
        _chart_with_insight(sb, chart, M, top, full, bot - top)
    else:
        lw = full * 0.46 if chart else full
        rh = min(1.4, (bot - top) / n)
        lx = M + 0.3
        sb.line(lx, top + rh / 2, lx, top + rh * (n - 0.5), ctx.rule, 3, name="Timeline axis")
        specs = []
        for i, it in enumerate(items):
            y = top + i * rh
            sb.rect(lx - 0.13, y + rh / 2 - 0.13, 0.26, 0.26, fill=ctx.block, line=ctx.bg, line_w=2, shape="ellipse", name="Marker %s" % it["id"])
            val = _value_text(it)
            ips = _item_paras(it, 22, 20)
            if val:
                ips[0] = P([val + "  \u2013  " + it["label"]], "body", pref=22, bold=True)
            specs.append({"rect": (lx + 0.35, y + 0.03, lw - 0.7, rh - 0.06), "paras": ips, "anchor": "m", "name": "Item %s" % it["id"]})
        sb.blocks(specs)
        if chart:
            _chart_with_insight(sb, chart, M + lw + 0.3, top, full - lw - 0.3, bot - top)


def role_comparison(sb: SlideBuilder) -> None:
    ctx, W, H, v = sb.ctx, sb.W, sb.H, sb.variant
    top, bot = sb.header(), _bottom(sb)
    items, chart = sb.items, sb.chart
    n = len(items)
    full = W - 2 * M
    if v == 3 and not chart:
        v = 2
    if not items and chart:
        v = 3
    if not items and not chart:
        return
    if v == 1:
        cw = (full - 0.4) / 2 if n > 1 else full
        specs = []
        heads = [ctx.block, ctx.block2 if ctx.block2 != ctx.block else ctx.secondary]
        for col in range(2 if n > 1 else 1):
            x = M + col * (cw + 0.4)
            idxs = [i for i in range(n) if i % 2 == col]
            hi = items[idxs[0]]
            sb.rect(x, top, cw, 0.75, fill=heads[col], rounded=True, name="Header chip %d" % (col + 1))
            specs.append({"rect": (x + 0.15, top, cw - 0.3, 0.75), "paras": [P([hi["label"]], "body", pref=26, bold=True, font="heading", color=ctx.on_fill(heads[col]))], "anchor": "m", "bg": heads[col], "name": "Header %d" % (col + 1)})
            by = top + 0.9
            sb.card(x, by, cw, bot - by, name="Column card %d" % (col + 1))
            paras: List[P] = []
            for j, i in enumerate(idxs):
                it = items[i]
                if j == 0:
                    dp = _detail_para(it, 20, "text")
                    if dp:
                        paras.append(dp)
                else:
                    ps = _item_paras(it, 22, 20, detail_color="text")
                    if paras:
                        ps[0].before = 8
                    paras += ps
            if paras:
                specs.append({"rect": (x + 0.25, by + 0.15, cw - 0.5, bot - by - 0.3), "paras": paras, "bg": ctx.surface, "name": "Column %d" % (col + 1)})
        sb.blocks(specs)
    elif v == 2:
        has_val = any(_value_text(it) for it in items)
        rh = min(1.1, (bot - top) / n)
        specs = []
        lw = 3.2
        vw = 2.4 if has_val else 0.0
        for i, it in enumerate(items):
            y = top + i * rh
            if i % 2 == 0:
                sb.rect(M, y, full, rh, fill=ctx.surface, name="Row band %d" % (i + 1))
            row_bg = ctx.surface if i % 2 == 0 else ctx.bg
            specs.append({"rect": (M + 0.15, y, lw, rh), "paras": [P([it["label"]], "body", pref=22, bold=True)], "anchor": "m", "bg": row_bg, "name": "Item %s label" % it["id"]})
            if has_val:
                specs.append({"rect": (M + 0.15 + lw, y, vw, rh), "paras": [P([_value_text(it) or " "], "stat", pref=32, bold=True, font="heading", color="primary", single_line=True)], "anchor": "m", "bg": row_bg, "name": "Item %s value" % it["id"]})
            dp = _detail_para(it)
            if dp:
                specs.append({"rect": (M + 0.15 + lw + vw, y, full - lw - vw - 0.3, rh), "paras": [dp], "anchor": "m", "bg": row_bg, "name": "Item %s detail" % it["id"]})
        sb.blocks(specs)
        if chart:
            sb.warn("chart not drawn: comparison variant 2 is a table; use layoutVariant 3")
    else:
        lw = full * 0.5
        _chart_with_insight(sb, chart, M, top, lw, bot - top)
        if n:
            _stack_rows(sb, items, M + lw + 0.4, top, full - lw - 0.4, bot - top, max_row=1.5)


def role_distribution(sb: SlideBuilder) -> None:
    ctx, W, H, v = sb.ctx, sb.W, sb.H, sb.variant
    top, bot = sb.header(), _bottom(sb)
    full = W - 2 * M
    ents = _dist_entries(sb)
    n = len(ents)
    if not n:
        return
    if sb.chart:
        sb.used_chart = True  # the chart's numbers drive every distribution variant
    vals = [e["value"] for e in ents]
    have_vals = all(isinstance(x, (int, float)) for x in vals) and sum(vals) > 0
    total = float(sum(vals)) if have_vals else 0.0
    if have_vals:
        for e in ents:
            if not e["text"]:
                e["text"] = "%d%%" % round(e["value"] / total * 100)
    if v == 1:
        cw = full * 0.46 if have_vals else 0
        legend_x = M + (cw + 0.4 if have_vals else 0)
        legend_w = full - (cw + 0.4 if have_vals else 0)
        if have_vals:
            ch = {"type": "doughnut", "categories": [e["label"] for e in ents], "series": [{"name": "Share", "values": vals}], "insight": (sb.chart or {}).get("insight", "")}
            ctype = (sb.chart or {}).get("type")
            sb.add_chart(ch, M, top, cw, bot - top, ctype if ctype in ("pie", "doughnut") else "doughnut", legend=False)
            sb.used_chart = True
        rh = min(1.1, (bot - top) / n)
        specs = []
        for i, e in enumerate(ents):
            y = top + i * rh
            sb.rect(legend_x, y + rh / 2 - 0.15, 0.3, 0.3, fill=e["color"], name="Legend swatch %d" % (i + 1))
            paras = [P([e["label"]], "body", pref=22, bold=True)]
            if e["detail"]:
                paras.append(P([e["detail"]["full"], e["detail"].get("short", "")], "body", pref=20, color="muted", before=2))
            vw = 1.7 if e["text"] else 0.0
            specs.append({"rect": (legend_x + 0.5, y, legend_w - 0.5 - vw, rh), "paras": paras, "anchor": "m", "name": "Legend %s" % e["id"]})
            if e["text"]:
                specs.append({"rect": (legend_x + legend_w - vw, y, vw, rh), "paras": [P([e["text"]], "stat", pref=30, bold=True, font="heading", color="primary")], "anchor": "m", "align": "r", "name": "Legend %s value" % e["id"]})
        sb.blocks(specs)
    elif v == 2:
        fw = full * 0.5
        rh = (bot - top) / n
        bh = min(0.85, rh - 0.1)
        specs = []
        mx = max(vals) if have_vals else 1.0
        for i, e in enumerate(ents):
            frac = (e["value"] / mx) if have_vals else 1.0 - 0.6 * i / max(1, n - 1)
            frac = max(0.42, min(1.0, frac))
            bw = fw * frac
            y = top + i * rh + (rh - bh) / 2
            sb.rect(M + (fw - bw) / 2, y, bw, bh, fill=e["color"], rounded=True, name="Funnel bar %d" % (i + 1))
            specs.append({"rect": (M + (fw - bw) / 2 + 0.1, y, bw - 0.2, bh), "paras": [P([e["label"]], "body", pref=22, bold=True, color=ctx.on_fill(e["color"]))], "anchor": "m", "align": "c", "bg": e["color"], "name": "Funnel label %s" % e["id"]})
            rx = M + fw + 0.4
            paras = []
            if e["text"]:
                paras.append(P([e["text"]], "stat", pref=30, bold=True, font="heading", color="primary"))
            if e["detail"]:
                paras.append(P([e["detail"]["full"], e["detail"].get("short", "")], "body", pref=20, color="muted"))
            if paras:
                specs.append({"rect": (rx, top + i * rh, full - fw - 0.4, rh), "paras": paras, "anchor": "m", "name": "Funnel note %s" % e["id"]})
        sb.blocks(specs)
    else:
        bh = 1.6
        by = top + 0.15
        x = M
        specs = []
        for i, e in enumerate(ents):
            share = (e["value"] / total) if have_vals else 1.0 / n
            w = full * share
            sb.rect(x, by, w, bh, fill=e["color"], line=ctx.bg, line_w=1.5, name="Share segment %d" % (i + 1))
            if w >= 1.0 and e["text"]:
                specs.append({"rect": (x, by, w, bh), "paras": [P([e["text"]], "caption", pref=18, bold=True, color=ctx.on_fill(e["color"]))], "anchor": "m", "align": "c", "bg": e["color"], "name": "Share value %s" % e["id"]})
            x += w
        cols = 3 if n > 3 else n
        rows = int(math.ceil(n / cols))
        ly = by + bh + 0.35
        rects = _grid(M, ly, full, bot - ly, cols, rows, 0.3, 0.1)
        for e, (rx, ry, rw, rh) in zip(ents, rects):
            sb.rect(rx, ry + 0.12, 0.28, 0.28, fill=e["color"], name="Legend swatch %s" % e["id"])
            paras = [P([e["label"] + ("  " + e["text"] if e["text"] else "")], "body", pref=22, bold=True)]
            if e["detail"]:
                paras.append(P([e["detail"]["full"], e["detail"].get("short", "")], "body", pref=20, color="muted", before=2))
            specs.append({"rect": (rx + 0.42, ry, rw - 0.42, rh), "paras": paras, "name": "Legend %s" % e["id"]})
        sb.blocks(specs)


def role_relationship(sb: SlideBuilder) -> None:
    ctx, W, H, v = sb.ctx, sb.W, sb.H, sb.variant
    top, bot = sb.header(), _bottom(sb)
    items = sb.items
    n = len(items)
    full = W - 2 * M
    if not n:
        return
    if v == 1:
        cx, cy = W / 2, top + (bot - top) / 2
        hub_d = min(2.3, (bot - top) * 0.5)
        rest = items[1:]
        m = len(rest)
        bw, bh = min(4.0, full * 0.3), 1.08
        rx = max(1.5, full / 2 - bw / 2 - 0.1)
        ry = (bot - top) / 2 - 0.58
        pos = []
        if m:
            start = -90.0 if m % 2 == 1 else -90.0 + 180.0 / m
            if m == 1:
                start = 0.0
            for j in range(m):
                a = math.radians(start + 360.0 * j / m)
                pos.append((cx + rx * math.cos(a), cy + ry * math.sin(a)))
        for it, (px, py) in zip(rest, pos):
            sb.line(cx, cy, px, py, ctx.rule, 2.0, name="Link %s" % it["id"])
        specs = []
        sb.rect(cx - hub_d / 2, cy - hub_d / 2, hub_d, hub_d, fill=ctx.block, line=ctx.bg, line_w=3, shape="ellipse", name="Hub")
        inner = hub_d * 0.74
        hub = items[0]
        specs.append({"rect": (cx - inner / 2, cy - inner / 2, inner, inner), "paras": [P([hub["label"]], "body", pref=22, bold=True, color=ctx.on_fill(ctx.block))], "anchor": "m", "align": "c", "bg": ctx.block, "name": "Hub label"})
        for it, (px, py) in zip(rest, pos):
            sb.card(px - bw / 2, py - bh / 2, bw, bh, name="Node %s" % it["id"])
            specs.append({"rect": (px - bw / 2 + 0.08, py - bh / 2 + 0.03, bw - 0.16, bh - 0.06), "paras": _item_paras(it, 20, 18, gap=0), "anchor": "m", "align": "c", "bg": ctx.surface, "name": "Node %s text" % it["id"]})
        sb.blocks(specs)
        hd = _detail_para(hub, 18)
        if hd:
            pass  # hub detail is shown through the narrative; the circle holds the label only
    else:
        gap = 0.5
        bw = (full - (n - 1) * gap) / n
        bh = min(4.2, bot - top - 0.2)
        by = top + 0.1
        specs = []
        for i, it in enumerate(items):
            x = M + i * (bw + gap)
            sb.card(x, by, bw, bh, name="Box %s" % it["id"])
            sb.rect(x, by, bw, 0.1, fill=ctx.block, name="Box bar %s" % it["id"])
            specs.append({"rect": (x + 0.1, by + 0.25, bw - 0.2, bh - 0.35), "paras": _item_paras(it, 22, 18), "bg": ctx.surface, "name": "Box %s text" % it["id"]})
            if i < n - 1:
                sb.line(x + bw + 0.06, by + bh / 2, x + bw + gap - 0.06, by + bh / 2, ctx.primary_text(ctx.bg), 3, arrow=True, name="Arrow %d" % (i + 1))
        sb.blocks(specs)


def role_case(sb: SlideBuilder) -> None:
    ctx, W, H, v = sb.ctx, sb.W, sb.H, sb.variant
    sm, items, med = sb.summary, sb.items, sb.media
    n = len(items)
    top, bot = sb.header(with_summary=False), _bottom(sb)
    full = W - 2 * M
    if v == 1:
        mw = full * 0.4 if med else 0
        if med and sb.picture(med[0], M, top, mw, bot - top, "cover", name="Case picture"):
            sb.used_media = True
        rx = M + (mw + 0.4 if med else 0)
        rw = full - (mw + 0.4 if med else 0)
        paras: List[P] = []
        if sm:
            paras.append(_tp(sm, "body", pref=22, color="text"))
        for it in items:
            paras.append(P([it["label"]], "body", pref=22, bold=True, color="primary", before=10))
            dp = _detail_para(it, 20, "text")
            if dp:
                dp.before = 1
                paras.append(dp)
        if paras:
            sb.text(rx, top, rw, bot - top, paras, name="Case text")
    elif v == 2:
        y = top
        if sm:
            sb.plain_text(M, y - 0.2, 1.2, 1.9, "“", 100, ctx.accent_text(ctx.bg), font="heading", name="Quote mark")
            sb.text(M + 1.2, y, full - 1.2, 1.9, [_tp(sm, "body", pref=28, color="text")], anchor="m", name="Summary")
            y += 2.15
        if n:
            rects = _grid(M, y, full, bot - y, n, 1, 0.3, 0.3)
            specs = []
            for it, (x, yy, w, h) in zip(items, rects):
                sb.card(x, yy, w, h, name="Card %s" % it["id"])
                sb.rect(x, yy, w, 0.09, fill=ctx.accent, name="Card bar %s" % it["id"])
                specs.append({"rect": (x + 0.2, yy + 0.25, w - 0.4, h - 0.35), "paras": [P([it["label"]], "body", pref=22, bold=True, color="primary")] + ([_detail_para(it, 20, "text")] if it.get("detail") else []), "bg": ctx.surface, "name": "Item %s" % it["id"]})
            sb.blocks(specs)
    else:
        tw = full - (2.9 if med else 0)
        y = top
        if sm:
            sb.text(M, y, tw, 1.7, [_tp(sm, "body", pref=24, color="text")], name="Summary")
        if med and sb.picture(med[0], W - M - 2.6, y, 2.6, 1.7, "cover", name="Case thumbnail"):
            sb.used_media = True
        y += 2.0
        if n:
            ymid = y + 0.3
            sb.line(M, ymid, W - M, ymid, ctx.rule, 3, name="Timeline axis")
            cw = full / n
            specs = []
            for i, it in enumerate(items):
                cx = M + (i + 0.5) * cw
                sb.badge(cx, ymid, 0.5, str(i + 1), name="Step %d" % (i + 1))
                specs.append({"rect": (M + i * cw + 0.05, ymid + 0.4, cw - 0.1, bot - ymid - 0.4), "paras": _item_paras(it, 22, 20), "align": "c", "name": "Item %s" % it["id"]})
            sb.blocks(specs)


def role_image(sb: SlideBuilder) -> None:
    ctx, W, H, v = sb.ctx, sb.W, sb.H, sb.variant
    sm, items, med = sb.summary, sb.items, sb.media
    m0 = med[0] if med else None
    full = W - 2 * M
    if v == 1:
        if sb.picture(m0, 0, 0, W, H, "cover", name="Image"):
            sb.used_media = True
        bar = 3.2 if items else 2.55
        sb.rect(0, H - bar, W, bar, fill=ctx.bg, alpha=0.9, name="Caption bar")
        lw = W * 0.5 if items else W - 2 * M
        sb.text(M, H - bar + 0.1, lw, 1.4, [_tp(sb.title, "title", pref=36, font="heading")], anchor="b", name="Title")
        if sm:
            sb.text(M, H - bar + 1.6, lw, bar - 1.7, [_tp(sm, "body", pref=20, color="muted")], name="Summary")
        if items:
            paras: List[P] = []
            for it in items:
                ps = _item_paras(it, 20, 18, gap=0)
                if paras:
                    ps[0].before = 4
                paras += ps
            sb.text(W * 0.58, H - bar + 0.1, W - M - W * 0.58, bar - 0.2, paras, anchor="m", name="Callouts")
    elif v == 2:
        px = W * 0.4
        if sb.picture(m0, px, 0, W - px, H, "cover", name="Image"):
            sb.used_media = True
        lw = px - M - 0.4
        sb.text(M, 0.55, lw, 2.3, [_tp(sb.title, "title", pref=34, font="heading")], anchor="t", name="Title")
        sb.rule_under_title(M, 2.95, lw)
        paras = []
        if sm:
            paras.append(_tp(sm, "body", pref=20, color="muted"))
        for it in items:
            ps = _item_paras(it, 20, 18, gap=0)
            ps[0].before = 8 if paras else 0
            paras += ps
        if paras:
            sb.text(M, 3.1, lw, H - 0.3 - 3.1, paras, name="Caption and callouts")
    else:
        top, bot = sb.header(with_summary=False), _bottom(sb)
        cap_h = 0.95 if sm else 0.0
        fh = bot - top - cap_h - 0.1
        fw = min(7.2, full * 0.58)
        fw = min(fw, fh * 1.7)
        fx = (W - fw) / 2
        sb.rect(fx, top, fw, fh, fill=ctx.surface, line=ctx.rule, line_w=1.0, name="Frame")
        if sb.picture(m0, fx + 0.1, top + 0.1, fw - 0.2, fh - 0.2, "contain", name="Image"):
            sb.used_media = True
        if sm:
            sb.text(fx, top + fh + 0.1, fw, cap_h, [_tp(sm, "body", pref=20, color="muted")], align="c", name="Caption")
        specs = []
        lw = fx - M - 0.3
        for i, it in enumerate(items):
            if i % 2 == 0:
                rect = (M, top, lw, (fh - 0.2) / 2 if len(items) > 2 else fh)
                if i == 2:
                    rect = (M, top + (fh + 0.2) / 2, lw, (fh - 0.2) / 2)
            else:
                rect = (fx + fw + 0.3, top, W - M - (fx + fw + 0.3), fh)
            specs.append({"rect": rect, "paras": _item_paras(it, 20, 18), "anchor": "m", "name": "Callout %s" % it["id"]})
        if specs:
            sb.blocks(specs)


def role_process(sb: SlideBuilder) -> None:
    ctx, W, H, v = sb.ctx, sb.W, sb.H, sb.variant
    top, bot = sb.header(), _bottom(sb)
    items = sb.items
    n = len(items)
    full = W - 2 * M
    if not n:
        return
    if v == 1:
        cw = full / n
        specs = []
        for i, it in enumerate(items):
            x = M + i * cw
            fill = ctx.block
            sb.rect(x, top + 0.05, cw - 0.04, 0.85, fill=fill, shape="chevron" if i else "pentagon", name="Chevron %d" % (i + 1))
            specs.append({"rect": (x + 0.45, top + 0.05, cw - 0.04 - 0.9, 0.85), "paras": [P([str(i + 1)], "body", pref=24, bold=True, font="heading", color=ctx.on_fill(fill))], "anchor": "m", "align": "c", "bg": fill, "name": "Step %d number" % (i + 1)})
            specs.append({"rect": (x + 0.05, top + 1.1, cw - 0.15, bot - top - 1.1), "paras": _item_paras(it, 22, 18), "name": "Item %s" % it["id"]})
        sb.blocks(specs)
    elif v == 2:
        rh = min(1.1, (bot - top) / n)
        bx = M + 0.4
        if n > 1:
            sb.line(bx, top + rh / 2, bx, top + rh * (n - 0.5), ctx.rule, 3, name="Step line")
        specs = []
        d = min(0.6, rh - 0.1)
        for i, it in enumerate(items):
            y = top + i * rh
            sb.badge(bx, y + rh / 2, d, str(i + 1), name="Badge %d" % (i + 1))
            specs.append({"rect": (M + 1.1, y, 4.0, rh), "paras": [P([it["label"]], "body", pref=22, bold=True)], "anchor": "m", "name": "Item %s label" % it["id"]})
            dp = _detail_para(it)
            if dp:
                specs.append({"rect": (M + 5.3, y, W - M - (M + 5.3), rh), "paras": [dp], "anchor": "m", "name": "Item %s detail" % it["id"]})
        sb.blocks(specs)
    else:
        per = 3
        rows = 1 if n <= per else 2
        gx = 0.7
        cw = (full - (per - 1) * gx) / per
        gy = 0.7
        ch = min(2.6, (bot - top - 0.2 - (rows - 1) * gy) / rows)
        y0 = top + 0.3
        specs = []
        placed = []
        for i in range(n):
            r, c = divmod(i, per)
            col = c if r % 2 == 0 else per - 1 - c
            placed.append((r, col))
        for i, (it, (r, col)) in enumerate(zip(items, placed)):
            x = M + col * (cw + gx)
            y = y0 + r * (ch + gy)
            sb.card(x, y, cw, ch, name="Card %s" % it["id"])
            sb.badge(x + 0.05, y + 0.02, 0.55, str(i + 1), name="Badge %d" % (i + 1))
            specs.append({"rect": (x + 0.15, y + 0.35, cw - 0.3, ch - 0.45), "paras": _item_paras(it, 22, 18), "bg": ctx.surface, "name": "Item %s" % it["id"]})
            if i < n - 1:
                nr, ncol = placed[i + 1]
                if nr == r:
                    nx = M + ncol * (cw + gx)
                    if nx > x:
                        sb.line(x + cw + 0.08, y + ch / 2, nx - 0.08, y + ch / 2, ctx.primary_text(ctx.bg), 3, arrow=True, name="Arrow %d" % (i + 1))
                    else:
                        sb.line(x - 0.08, y + ch / 2, nx + cw + 0.08, y + ch / 2, ctx.primary_text(ctx.bg), 3, arrow=True, name="Arrow %d" % (i + 1))
                else:
                    sb.line(x + cw / 2, y + ch + 0.06, x + cw / 2, y + ch + gy - 0.06, ctx.primary_text(ctx.bg), 3, arrow=True, name="Arrow %d" % (i + 1))
        sb.blocks(specs)


_PRIORITY_LABEL = {1: "High", 2: "Medium"}


def _priority_style(sb: SlideBuilder, p: Optional[int]) -> Tuple[str, str]:
    """(chip text, fill colour) for a risk priority (1 high .. 3 low)."""
    ctx = sb.ctx
    if p is None:
        return "", ctx.rule
    if p <= 1:
        return "High", ctx.accent
    if p == 2:
        return "Medium", ctx.secondary
    return "Low", C.mix_hex(ctx.bg, ctx.fg(ctx.bg), 0.22)


def role_risks(sb: SlideBuilder) -> None:
    ctx, W, H, v = sb.ctx, sb.W, sb.H, sb.variant
    top, bot = sb.header(), _bottom(sb)
    items = sb.items
    n = len(items)
    full = W - 2 * M
    if not n:
        return
    if v == 1:
        rh = min(1.2, (bot - top) / n)
        specs = []
        for i, it in enumerate(items):
            y = top + i * rh
            sb.line(M, y, W - M, y, ctx.rule, 0.75, name="Row rule %d" % (i + 1))
            specs.append({"rect": (M, y, 3.3, rh), "paras": [P([it["label"]], "body", pref=22, bold=True)], "anchor": "m", "name": "Item %s label" % it["id"]})
            dp = _detail_para(it)
            if dp:
                specs.append({"rect": (M + 3.5, y, full - 3.5 - 1.7, rh), "paras": [dp], "anchor": "m", "name": "Item %s detail" % it["id"]})
            txt, fill = _priority_style(sb, it.get("priority"))
            if txt:
                sb.chip(W - M - 1.5, y + rh / 2 - 0.23, 1.5, 0.46, txt, fill, name="Priority %s" % it["id"])
        sb.blocks(specs)
    else:
        order = sorted(range(n), key=lambda i: (_priority_key(items[i]), i))
        cols = 1 if n <= 3 else 2
        rows = int(math.ceil(n / cols))
        rects = _grid(M, top, full, bot - top, cols, rows, 0.3, 0.2)
        specs = []
        for idx, (i, (x, y, w, h)) in enumerate(zip(order, rects)):
            it = items[i]
            txt, fill = _priority_style(sb, it.get("priority"))
            sb.card(x, y, w, h, name="Card %s" % it["id"])
            sb.rect(x, y, 0.14, h, fill=fill if txt else ctx.rule, name="Priority bar %s" % it["id"])
            chip_w = 1.3 if txt else 0.0
            specs.append({"rect": (x + 0.3, y + 0.05, w - 0.45 - chip_w, h - 0.1), "paras": _item_paras(it, 22, 18), "anchor": "m", "bg": ctx.surface, "name": "Item %s" % it["id"]})
            if txt:
                sb.chip(x + w - 1.4, y + 0.15, 1.3, 0.42, txt, fill, name="Priority %s" % it["id"])
        sb.blocks(specs)


def role_observation(sb: SlideBuilder) -> None:
    ctx, W, H, v = sb.ctx, sb.W, sb.H, sb.variant
    sm, items, chart = sb.summary, sb.items, sb.chart
    n = len(items)
    top, bot = sb.header(with_summary=(sb.variant == 2)), _bottom(sb)
    full = W - 2 * M
    if v == 1:
        lw = full * 0.5 - 0.3 if n else full
        if sm:
            sb.text(M, top, lw, bot - top, [_tp(sm, "body", pref=30)], name="Summary")
        if n:
            rx = M + (full * 0.5 if sm else 0)
            rw = full - (full * 0.5 if sm else 0)
            if sm:
                sb.line(rx - 0.15, top, rx - 0.15, bot, ctx.rule, 0.75, name="Column rule")
            rh = min(1.5, (bot - top) / n)
            specs = []
            for i, it in enumerate(items):
                y = top + i * rh
                if i:
                    sb.line(rx, y, rx + rw, y, ctx.rule, 0.75, name="Row rule %d" % i)
                specs.append({"rect": (rx, y, 0.8, rh), "paras": [P([str(i + 1)], "stat", pref=34, bold=True, font="heading", color="accent")], "anchor": "m", "name": "Index %d" % (i + 1)})
                specs.append({"rect": (rx + 0.85, y + 0.03, rw - 0.85, rh - 0.06), "paras": _item_paras(it, 22, 20), "anchor": "m", "name": "Item %s" % it["id"]})
            sb.blocks(specs)
        if chart:
            sb.warn("chart not drawn: observation variant 1 has no chart area; use layoutVariant 2")
    elif v == 2:
        y = top
        if chart:
            ch = 1.45 if n else bot - y
            sb.add_chart(chart, M, y, full, ch)
            sb.used_chart = True
            y += ch + 0.2
        if n:
            cols = min(n, 3)
            rows = int(math.ceil(n / cols))
            rects = _grid(M, y, full, bot - y, cols, rows, 0.3, 0.25)
            specs = []
            for it, (x, yy, w, h) in zip(items, rects):
                sb.card(x, yy, w, h, name="Card %s" % it["id"])
                sb.rect(x, yy, w, 0.09, fill=ctx.accent, name="Card bar %s" % it["id"])
                specs.append({"rect": (x + 0.2, yy + 0.22, w - 0.4, h - 0.3), "paras": _item_paras(it, 22, 18), "bg": ctx.surface, "name": "Item %s" % it["id"]})
            sb.blocks(specs)
    else:
        text = [sm["full"], sm.get("short", "")] if sm else [sb.s["coreMessage"]]
        ch = 2.0 if n else bot - top
        sb.card(M, top, full, ch, name="Callout card")
        sb.rect(M, top, 0.12, ch, fill=ctx.accent, name="Callout bar")
        sb.text(M + 0.4, top + 0.1, full - 0.7, ch - 0.2, [P(text, "stat", pref=30, font="heading", color="text")], anchor="m", bg=ctx.surface, name="Callout")
        if n:
            paras: List[P] = []
            for it in items:
                paras.append(P([it["label"]], "body", pref=22, bold=True, bullet=True, before=8 if paras else 0))
                dp = _detail_para(it, 20)
                if dp:
                    dp.indent = 0.3
                    paras.append(dp)
            sb.text(M, top + ch + 0.25, full, bot - top - ch - 0.25, paras, name="Observations")
        if chart:
            sb.warn("chart not drawn: observation variant 3 has no chart area; use layoutVariant 2")


def role_ambient(sb: SlideBuilder) -> None:
    ctx, W, H, v = sb.ctx, sb.W, sb.H, sb.variant
    sm, med = sb.summary, sb.media
    if v == 1:
        light = "FFFFFF"
        if med and sb.picture(med[0], 0, 0, W, H, "cover", name="Image"):
            sb.used_media = True
            sb.rect(0, H - 3.4, W, 3.4, fill="000000", alpha=0.55, name="Scrim")
            under = "000000"
            soft = "E6E6E6"
            sb.set_background(ctx.bg)
        else:
            sb.set_background(ctx.block, gradient_to=ctx.block2, angle=45)
            under = ctx.block
            light = ctx.on_fill(ctx.block)
            soft = C.mix_hex(light, ctx.block, 0.25)
        sb.text(M, H - 2.8, W * 0.72, 1.6, [_tp(sb.title, "title", pref=52, font="heading", color=light)], anchor="b", bg=under, name="Title")
        if sm:
            sb.text(M, H - 1.15, W * 0.72, 0.8, [_tp(sm, "body", pref=20, color=soft)], bg=under, name="Summary")
    else:
        sb.text(M, 1.0, W - 2 * M, 3.7, [P([sb.title["full"], sb.title.get("short", "")], "title", pref=110, bold=True, font="heading", color="primary", outline=True)], anchor="m", name="Title")
        sb.rect(M, 5.0, 1.2, 0.09, fill=ctx.accent, name="Accent rule")
        if sm:
            sb.text(M, 5.25, W - 2 * M - 2.0, 1.2, [_tp(sm, "body", pref=22, color="muted")], name="Summary")


def role_actions(sb: SlideBuilder) -> None:
    ctx, W, H, v = sb.ctx, sb.W, sb.H, sb.variant
    top, bot = sb.header(), _bottom(sb)
    items = sb.items
    n = len(items)
    full = W - 2 * M
    if not n:
        return
    if v == 1:
        cw = full / n
        ylin = top + 1.05
        sb.line(M, ylin, W - M, ylin, ctx.rule, 3, name="Roadmap axis")
        specs = []
        for i, it in enumerate(items):
            cx = M + (i + 0.5) * cw
            sb.rect(cx - 0.17, ylin - 0.17, 0.34, 0.34, fill=ctx.block, line=ctx.bg, line_w=2, shape="ellipse", name="Phase marker %d" % (i + 1))
            val = _value_text(it)
            if val:
                specs.append({"rect": (cx - cw / 2 + 0.05, top, cw - 0.1, 0.7), "paras": [P([val], "body", pref=22, bold=True, color="primary")], "anchor": "b", "align": "c", "name": "Phase %s date" % it["id"]})
            specs.append({"rect": (cx - cw / 2 + 0.05, ylin + 0.35, cw - 0.1, bot - ylin - 0.35), "paras": _item_paras(it, 22, 18), "align": "c", "name": "Item %s" % it["id"]})
        sb.blocks(specs)
    elif v == 2:
        rh = min(1.1, (bot - top) / n)
        specs = []
        for i, it in enumerate(items):
            y = top + i * rh
            if i:
                sb.line(M, y, W - M, y, ctx.rule, 0.75, name="Row rule %d" % i)
            sb.rect(M + 0.05, y + rh / 2 - 0.17, 0.34, 0.34, line=ctx.primary_text(ctx.bg), line_w=2, rounded=True, name="Checkbox %d" % (i + 1))
            specs.append({"rect": (M + 0.6, y, 4.2, rh), "paras": [P([it["label"]], "body", pref=22, bold=True)], "anchor": "m", "name": "Item %s label" % it["id"]})
            p = it.get("priority")
            cw_ = 1.0 if isinstance(p, int) else 0.0
            dp = _detail_para(it)
            if dp:
                specs.append({"rect": (M + 5.0, y, full - 5.0 - cw_ - 0.2, rh), "paras": [dp], "anchor": "m", "name": "Item %s detail" % it["id"]})
            if cw_:
                fill = ctx.accent if p <= 1 else (ctx.secondary if p == 2 else C.mix_hex(ctx.bg, ctx.fg(ctx.bg), 0.22))
                sb.chip(W - M - 1.0, y + rh / 2 - 0.21, 1.0, 0.42, "P%d" % p, fill, name="Priority %s" % it["id"])
        sb.blocks(specs)
    else:
        gap = 0.3
        cw = (full - 2 * gap) / 3
        lanes: List[List[Dict[str, Any]]] = [[], [], []]
        for it in items:
            p = it.get("priority")
            lanes[0 if (isinstance(p, int) and p <= 1) else 1 if p == 2 else 2].append(it)
        names = ["Now", "Next", "Later"]
        fills = [ctx.block, ctx.block2 if ctx.block2 != ctx.block else ctx.secondary, C.mix_hex(ctx.bg, ctx.fg(ctx.bg), 0.22)]
        specs = []
        for li in range(3):
            x = M + li * (cw + gap)
            sb.rect(x, top, cw, 0.55, fill=fills[li], rounded=True, name="Lane header %s" % names[li])
            specs.append({"rect": (x, top, cw, 0.55), "paras": [P([names[li]], "body", pref=22, bold=True, color=ctx.on_fill(fills[li]), font="heading")], "anchor": "m", "align": "c", "bg": fills[li], "name": "Lane %s" % names[li]})
            lane = lanes[li]
            if not lane:
                continue
            ay = top + 0.7
            ah = bot - ay
            ch = min(1.9, (ah - (len(lane) - 1) * 0.15) / len(lane))
            for j, it in enumerate(lane):
                y = ay + j * (ch + 0.15)
                sb.card(x, y, cw, ch, name="Card %s" % it["id"])
                specs.append({"rect": (x + 0.15, y + 0.04, cw - 0.3, ch - 0.08), "paras": _item_paras(it, 20, 18, gap=1), "anchor": "m", "bg": ctx.surface, "name": "Item %s" % it["id"]})
        sb.blocks(specs)


def role_result(sb: SlideBuilder) -> None:
    ctx, W, H, v = sb.ctx, sb.W, sb.H, sb.variant
    sm, items, chart = sb.summary, sb.items, sb.chart
    top, bot = sb.header(with_summary=False), _bottom(sb)
    full = W - 2 * M
    if v == 3 and not chart:
        v = 1
    if not items:
        if sm:
            sb.text(M, top, full, bot - top, [_tp(sm, "body", pref=30)], anchor="m", name="Summary")
        return
    first, rest = items[0], items[1:]
    val = _value_text(first)
    hero_text = val or first["label"]
    hero_kind = "hero" if val else "stat"

    def hero_paras(color: str, align_bg: Optional[str] = None, with_detail: bool = True) -> List[P]:
        ps = [P([hero_text], hero_kind, pref=110 if val else 54, bold=True, font="heading", color=color)]
        if val:
            ps.append(P([first["label"]], "body", pref=26, bold=True, before=4, color="text" if color != "onfill" else color))
        if with_detail:
            dp = _detail_para(first)
            if dp:
                dp.before = 4
                ps.append(dp)
        return ps

    if v == 1:
        lw = min(6.2, full * 0.48)
        sb.text(M, top, lw, bot - top, hero_paras("accent"), anchor="m", name="Result")
        rx = M + lw + 0.5
        rw = full - lw - 0.5
        sb.line(rx - 0.25, top + 0.2, rx - 0.25, bot - 0.2, ctx.rule, 0.75, name="Column rule")
        paras: List[P] = []
        if sm:
            paras.append(_tp(sm, "body", pref=22, color="muted"))
        for it in rest:
            ps = _item_paras(it, 22, 20, detail_color="text")
            ps[0].before = 12 if paras else 0
            paras += ps
        if paras:
            sb.text(rx, top, rw, bot - top, paras, anchor="m", name="Takeaways")
    elif v == 2:
        bh = 2.45
        sb.rect(0, top, W, bh, fill=ctx.surface, name="Result band")
        ps = [P([hero_text], hero_kind, pref=110 if val else 54, bold=True, font="heading", color="accent")]
        if val:
            ps.append(P([first["label"]], "body", pref=26, bold=True, before=2, color="text"))
        sb.text(M, top + 0.05, full, bh - 0.1, ps, anchor="m", align="c", bg=ctx.surface, name="Result")
        y = top + bh + 0.25
        if rest or sm:
            if sm and not rest:
                sb.text(M, y, full, bot - y, [_tp(sm, "body", pref=24, color="muted")], align="c", name="Summary")
            else:
                rects = _grid(M, y, full, bot - y, len(rest), 1, 0.3, 0.3)
                specs = []
                for it, (x, yy, w, h) in zip(rest, rects):
                    sb.card(x, yy, w, h, name="Card %s" % it["id"])
                    specs.append({"rect": (x + 0.2, yy + 0.1, w - 0.4, h - 0.2), "paras": _item_paras(it, 22, 18), "anchor": "m", "bg": ctx.surface, "name": "Item %s" % it["id"]})
                sb.blocks(specs)
    else:
        lw = min(5.4, full * 0.42)
        ch = min(2.8, (bot - top) * 0.58)
        sb.text(M, top, lw, ch, hero_paras("accent", with_detail=False), anchor="m", name="Result")
        _chart_with_insight(sb, chart, M + lw + 0.4, top, full - lw - 0.4, ch)
        y = top + ch + 0.25
        if rest:
            rects = _grid(M, y, full, bot - y, len(rest), 1, 0.3, 0.3)
            specs = []
            for it, (x, yy, w, h) in zip(rest, rects):
                sb.card(x, yy, w, h, name="Card %s" % it["id"])
                specs.append({"rect": (x + 0.2, yy + 0.05, w - 0.4, h - 0.1), "paras": _item_paras(it, 22, 18, gap=1), "anchor": "m", "bg": ctx.surface, "name": "Item %s" % it["id"]})
            sb.blocks(specs)
        elif sm:
            sb.text(M, y, full, bot - y, [_tp(sm, "body", pref=22, color="muted")], name="Summary")


def role_team(sb: SlideBuilder) -> None:
    ctx, W, H, v = sb.ctx, sb.W, sb.H, sb.variant
    top, bot = sb.header(), _bottom(sb)
    items, med = sb.items, sb.media
    n = len(items)
    full = W - 2 * M
    if not n:
        return
    if v == 1:
        cols = min(n, 4)
        rows = int(math.ceil(n / cols))
        rects = _grid(M, top, full, bot - top, cols, rows, 0.3, 0.15)
        d = max(0.7, min(1.5, (rects[0][3]) * 0.45))
        specs = []
        for i, (it, (x, y, w, h)) in enumerate(zip(items, rects)):
            cx = x + w / 2
            shown = False
            if i < len(med):
                shown = sb.picture(med[i], cx - d / 2, y, d, d, "cover", ellipse=True, name="Avatar %s" % it["id"])
                sb.used_media = sb.used_media or shown
            if not shown:
                sb.rect(cx - d / 2, y, d, d, fill=ctx.block, shape="ellipse", name="Avatar %s" % it["id"])
                inner = d * 0.7
                specs.append({"rect": (cx - inner / 2, y + (d - inner) / 2, inner, inner), "paras": [P([sb.initials(it["label"])], "body", pref=26, bold=True, font="heading", color=ctx.on_fill(ctx.block))], "anchor": "m", "align": "c", "bg": ctx.block, "name": "Initials %s" % it["id"]})
            specs.append({"rect": (x, y + d + 0.08, w, h - d - 0.08), "paras": _item_paras(it, 22, 18, gap=1), "align": "c", "name": "Item %s" % it["id"]})
        sb.blocks(specs)
    else:
        cols = 2 if n > 4 else 1
        rows = int(math.ceil(n / cols))
        rects = _grid(M, top, full, bot - top, cols, rows, 0.5, 0.0)
        specs = []
        for i, (it, (x, y, w, h)) in enumerate(zip(items, rects)):
            sb.line(x, y, x + w, y, ctx.rule, 0.75, name="Row rule %d" % (i + 1))
            d = min(0.7, h - 0.2)
            shown = False
            if i < len(med):
                shown = sb.picture(med[i], x + 0.05, y + (h - d) / 2, d, d, "cover", ellipse=True, name="Avatar %s" % it["id"])
                sb.used_media = sb.used_media or shown
            if not shown:
                sb.rect(x + 0.05, y + (h - d) / 2, d, d, fill=ctx.block, shape="ellipse", name="Avatar %s" % it["id"])
                inner = d * 0.76
                specs.append({"rect": (x + 0.05 + (d - inner) / 2, y + (h - inner) / 2, inner, inner), "paras": [P([sb.initials(it["label"])], "body", pref=20, bold=True, font="heading", color=ctx.on_fill(ctx.block))], "anchor": "m", "align": "c", "bg": ctx.block, "name": "Initials %s" % it["id"]})
            specs.append({"rect": (x + d + 0.25, y + 0.03, w - d - 0.3, h - 0.06), "paras": _item_paras(it, 22, 18, gap=1), "anchor": "m", "name": "Item %s" % it["id"]})
        sb.blocks(specs)


def role_closing(sb: SlideBuilder) -> None:
    ctx, W, H, v = sb.ctx, sb.W, sb.H, sb.variant
    sm, items, med = sb.summary, sb.items, sb.media
    n = len(items)
    full = W - 2 * M
    if v == 1:
        sb.set_background(ctx.block, gradient_to=ctx.block2, angle=0)
        on = ctx.on_fill(ctx.block)
        soft = C.mix_hex(on, ctx.block, 0.25)
        if med and sb.picture(med[0], 0, 0, W, H, "cover", name="Backdrop"):
            sb.used_media = True
            sb.rect(0, 0, W, H, fill=ctx.block, alpha=0.9, name="Backdrop veil")
        sb.text(M, 1.3, full, 2.3, [_tp(sb.title, "title", pref=54, font="heading", color=on)], anchor="b", align="c", bg=ctx.block, name="Title")
        if sm:
            sb.text(M + 1.0, 3.75, full - 2.0, 1.1, [_tp(sm, "body", pref=22, color=soft)], align="c", bg=ctx.block, name="Summary")
        if n:
            rects = _grid(M, 5.0, full, H - 0.55 - 5.0, n, 1, 0.4, 0.3)
            specs = []
            for it, (x, y, w, h) in zip(items, rects):
                specs.append({"rect": (x, y, w, h), "paras": _item_paras(it, 22, 20, label_color=on, detail_color=soft, gap=1), "align": "c", "bg": ctx.block, "name": "Item %s" % it["id"]})
            sb.blocks(specs)
    else:
        lw = full * 0.55
        sb.text(M, 1.5, lw, 2.8, [_tp(sb.title, "title", pref=50, font="heading")], anchor="b", name="Title")
        sb.rect(M, 4.5, 1.2, 0.09, fill=ctx.accent, name="Accent rule")
        if sm:
            sb.text(M, 4.75, lw, 1.7, [_tp(sm, "body", pref=22, color="muted")], name="Summary")
        rx = M + lw + 0.6
        rw = full - lw - 0.6
        sb.line(rx - 0.3, 1.5, rx - 0.3, H - 1.0, ctx.rule, 0.75, name="Column rule")
        y = 1.5
        if med and sb.picture(med[0], rx, y, rw, 1.4, "contain", name="Closing image"):
            sb.used_media = True
            y += 1.55
        if n:
            _stack_rows(sb, items, rx, y, rw, H - 0.8 - y, max_row=1.5, rules=True)


ROLE_RENDERERS = {
    "cover": role_cover,
    "statement": role_statement,
    "breakdown": role_breakdown,
    "transition": role_transition,
    "context": role_context,
    "metrics": role_metrics,
    "trend": role_trend,
    "comparison": role_comparison,
    "distribution": role_distribution,
    "relationship": role_relationship,
    "case": role_case,
    "image": role_image,
    "process": role_process,
    "risks": role_risks,
    "observation": role_observation,
    "ambient": role_ambient,
    "actions": role_actions,
    "result": role_result,
    "team": role_team,
    "closing": role_closing,
}


# --------------------------------------------------------------------------
# Presentation assembly
# --------------------------------------------------------------------------

_SCHEME_MAP = {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}


def _theme_part(prs):
    return prs.slide_master.part.part_related_by(RT.THEME)


def _theme_colors_from_xml(blob: bytes) -> Tuple[Dict[str, str], Dict[str, str]]:
    """Return (scheme colours, fonts) found in a theme part; colours keyed dk1, lt1, dk2, lt2, accent1..6."""
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    root = etree.fromstring(blob)
    cols: Dict[str, str] = {}
    cs = root.find(".//a:clrScheme", ns)
    if cs is not None:
        for child in cs:
            name = etree.QName(child).localname
            c = child[0] if len(child) else None
            if c is None:
                continue
            val = c.get("val") if etree.QName(c).localname == "srgbClr" else c.get("lastClr")
            if val and re.fullmatch(r"[0-9A-Fa-f]{6}", val):
                cols[name] = val.upper()
    fonts: Dict[str, str] = {}
    for key, tag in (("heading", "majorFont"), ("body", "minorFont")):
        lat = root.find(".//a:fontScheme/a:%s/a:latin" % tag, ns)
        if lat is not None and lat.get("typeface"):
            fonts[key] = lat.get("typeface")
    return cols, fonts


def _template_background(prs, layout, scheme: Dict[str, str]) -> str:
    """Best-effort background colour behind slides built on *layout* (white when unknown)."""
    for part_el in (layout._element, prs.slide_master._element):
        bg = part_el.find(".//" + qn("p:bg"))
        if bg is None:
            continue
        srgb = bg.find(".//" + qn("a:srgbClr"))
        if srgb is not None and srgb.get("val"):
            return srgb.get("val").upper()
        sch = bg.find(".//" + qn("a:schemeClr"))
        if sch is not None:
            key = _SCHEME_MAP.get(sch.get("val"), sch.get("val"))
            if key in scheme:
                return scheme[key]
    return scheme.get("lt1", "FFFFFF")


def theme_from_template(prs, layout, base_theme: Dict[str, Any]) -> Tuple[Dict[str, Any], str]:
    """Derive a theme dict from a template's own theme colours and fonts. Returns (theme, background hex)."""
    cols, fonts = _theme_colors_from_xml(_theme_part(prs).blob)
    bg = _template_background(prs, layout, cols)
    dark = C.relative_luminance(bg) < 0.35
    text = cols.get("lt1", "FFFFFF") if dark else cols.get("dk1", "000000")
    acc = [cols.get("accent%d" % i) for i in range(1, 7) if cols.get("accent%d" % i)]
    primary = cols.get("dk2") if not dark else (acc[0] if acc else text)
    if primary and C.contrast_ratio(primary, bg) < 3.0:
        primary = acc[0] if acc else text
    theme = {
        "id": "template",
        "name": "Template theme",
        "family": "template",
        "dark": dark,
        "autoPick": False,
        "palette": {
            "background": bg,
            "surface": C.mix_hex(bg, text, 0.06),
            "primary": primary or text,
            "secondary": acc[0] if acc else (primary or text),
            "accent": acc[1] if len(acc) > 1 else (primary or text),
            "text": text,
            "mutedText": C.mix_hex(text, bg, 0.35),
            "chartSeries": acc or [primary or text],
        },
        "fonts": {
            "heading": {"name": fonts.get("heading", "Calibri"), "fallback": fonts.get("heading", "Calibri")},
            "body": {"name": fonts.get("body", "Calibri"), "fallback": fonts.get("body", "Calibri")},
            "mono": {"name": "Consolas", "fallback": "Courier New"},
        },
        "decoration": {"shadows": False, "cornerRadius": 0, "gradients": False, "rules": "thin", "imageTreatment": "natural"},
    }
    return theme, bg


def _delete_all_slides(prs) -> None:
    """Remove every slide: drop the sldId entries and their relationships so no slide part is saved."""
    lst = prs.slides._sldIdLst
    for sld in list(lst):
        rid = sld.rId
        lst.remove(sld)
        prs.part.drop_rel(rid)
    root = prs.part._element
    for tag in ("p:custShowLst",):
        for el in root.findall(qn(tag)):
            root.remove(el)
    # section lists reference slide ids and would dangle
    for ext in root.findall(".//" + qn("p:extLst") + "/" + qn("p:ext")):
        if any(etree.QName(ch).localname == "sectionLst" for ch in ext):
            ext.getparent().remove(ext)


def _pick_blank_layout(prs):
    layouts = list(prs.slide_layouts)
    for lay in layouts:
        if lay.name.strip().lower() == "blank":
            return lay

    def count(lay) -> int:
        n = 0
        for ph in lay.placeholders:
            if ph.placeholder_format.type not in (13, 15, 16):  # slide number, footer, date
                n += 1
        return n

    return sorted(layouts, key=lambda l: (count(l), layouts.index(l)))[0]


def _style_fresh_masters(prs, ctx: DeckContext, old_width_emu: int) -> None:
    """Apply the theme to the default master: theme part colours/fonts, background, 16:9 placeholder geometry."""
    # theme colours and fonts
    part = _theme_part(prs)
    root = etree.fromstring(part.blob)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    cs = root.find(".//a:clrScheme", ns)
    pal = ctx.theme["palette"]
    series = ctx.series + [ctx.primary] * 6
    mapping = {
        "dk1": ctx.text if not ctx.dark else ctx.bg,
        "lt1": ctx.bg if not ctx.dark else ctx.text,
        "dk2": ctx.primary,
        "lt2": ctx.surface,
        "accent1": series[0],
        "accent2": series[1],
        "accent3": series[2],
        "accent4": series[3],
        "accent5": series[4],
        "accent6": series[5],
        "hlink": ctx.primary,
        "folHlink": ctx.secondary,
    }
    if cs is not None:
        cs.set("name", "pptx-deck")
        for child in list(cs):
            nm = etree.QName(child).localname
            if nm in mapping:
                for sub in list(child):
                    child.remove(sub)
                etree.SubElement(child, "{%s}srgbClr" % ns["a"]).set("val", mapping[nm])
    fs = root.find(".//a:fontScheme", ns)
    if fs is not None:
        fs.set("name", "pptx-deck")
        for tag, fkey in (("majorFont", "heading"), ("minorFont", "body")):
            lat = fs.find("a:%s/a:latin" % tag, ns)
            if lat is not None:
                lat.set("typeface", ctx.fonts[fkey])
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    # master background
    try:
        f = prs.slide_master.background.fill
        f.solid()
        f.fore_color.rgb = C.hex_to_rgb(ctx.bg)
    except Exception:
        pass
    # widen placeholders from the 4:3 default to the new width
    if prs.slide_width != old_width_emu:
        ratio = prs.slide_width / old_width_emu
        containers = [prs.slide_master] + list(prs.slide_layouts)
        for cont in containers:
            for shp in cont.shapes:
                if shp.left is not None and shp.width is not None:
                    shp.left = Emu(int(round(shp.left * ratio)))
                    shp.width = Emu(int(round(shp.width * ratio)))


def _set_core_properties(prs, spec: Dict[str, Any]) -> None:
    cp = prs.core_properties
    y, mo, d = C.spec_date(spec)
    when = _dt.datetime(y, mo, d, 0, 0, 0)
    cp.title = spec.get("title", "")
    cp.subject = spec.get("goal", "")
    cp.author = spec.get("owner", "") or ""
    cp.last_modified_by = spec.get("owner", "") or ""
    cp.keywords = ""
    cp.comments = ""
    cp.category = ""
    cp.language = spec.get("language", "en")
    cp.created = when
    cp.modified = when
    cp.last_printed = when
    cp.revision = 1
    cp.version = ""
    cp.content_status = ""
    cp.identifier = ""


def _set_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def _rewrite_zip_fixed(path: Path) -> None:
    """Rewrite the saved package with fixed entry timestamps so identical content gives identical bytes."""
    with zipfile.ZipFile(path, "r") as zin:
        entries = [(info.filename, zin.read(info.filename)) for info in zin.infolist()]
    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in entries:
            zi = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            zout.writestr(zi, data)
    tmp.replace(path)


class BuildResult:
    """Outcome of :func:`build_deck`."""

    def __init__(self, path: Path, slide_count: int, warnings: List[str], theme_id: str, variants: List[int]) -> None:
        self.path = path
        self.slide_count = slide_count
        self.warnings = warnings
        self.theme_id = theme_id
        self.variants = variants


def build_deck(
    spec: Dict[str, Any],
    out_path: Path | str,
    spec_dir: Path | str = ".",
    theme_id: Optional[str] = None,
    template: Optional[Path | str] = None,
    validate: bool = True,
) -> BuildResult:
    """Build a deck from a parsed spec and save it to *out_path*.

    ``theme_id`` forces a skill theme (also in template mode). ``template`` is the path of an
    existing .pptx whose masters and layouts are reused. With ``validate=True`` a spec with
    errors raises ``ValueError`` listing them. Returns a :class:`BuildResult`.
    """
    spec_dir = Path(spec_dir)
    out_path = Path(out_path)
    if validate:
        import validate_spec as V

        errors, _w = V.validate_spec(spec, spec_dir)
        if errors:
            raise ValueError("spec has %d error(s):\n  " % len(errors) + "\n  ".join(str(e) for e in errors))
    if template is None and isinstance(spec.get("template"), str):
        tp = Path(spec["template"])
        template = tp if tp.is_absolute() else spec_dir / tp

    base_theme = C.get_theme(theme_id or spec["theme"])
    notes_warn: List[str] = []
    template_bg: Optional[str] = None
    if template:
        prs = Presentation(str(template))
        _delete_all_slides(prs)
        layout = _pick_blank_layout(prs)
        if theme_id:
            theme, bg_mode = base_theme, "slide"
        else:
            theme, template_bg = theme_from_template(prs, layout, base_theme)
            bg_mode = "inherit"
        k = prs.slide_height / C.EMU_PER_INCH / 7.5
        if spec.get("aspect") and abs(prs.slide_width / prs.slide_height - C.slide_size(spec["aspect"])[0] / 7.5) > 0.02:
            notes_warn.append("deck: the template's slide size differs from aspect %s; the template's size is used" % spec["aspect"])
    else:
        prs = Presentation()
        old_w = prs.slide_width
        w_emu, h_emu = C.slide_size_emu(spec.get("aspect", "16:9"))
        prs.slide_width, prs.slide_height = Emu(w_emu), Emu(h_emu)
        layout = prs.slide_layouts[6]
        theme, bg_mode, k = base_theme, "master", 1.0
    ctx = DeckContext(spec, theme, prs, spec_dir, k, bg_mode, template_bg)
    ctx.warnings.extend(notes_warn)
    if not template:
        _style_fresh_masters(prs, ctx, old_w)

    variants: List[int] = []
    for i, s in enumerate(spec["slides"]):
        slide = prs.slides.add_slide(layout)
        slide._element.cSld.set("name", s["id"])
        v = C.resolve_variant(spec, i)
        variants.append(v)
        sb = SlideBuilder(ctx, slide, s, i, v)
        sb.apply_base_background()
        ROLE_RENDERERS[s["role"]](sb)
        if s.get("chart") and not sb.used_chart and not any(sb.label in w and "chart not drawn" in w for w in ctx.warnings):
            sb.warn("chart not shown by layout variant %d of role %s" % (sb.variant, s["role"]))
        if s.get("media") and not sb.used_media and s["role"] in ("cover", "case", "ambient", "closing", "image"):
            sb.warn("media not shown by layout variant %d of role %s" % (sb.variant, s["role"]))
        notes = (s.get("notes") or "").strip()
        if not notes:
            notes = s["coreMessage"]
            sb.warn("no speaker notes; coreMessage used as notes")
        _set_notes(slide, notes)

    _set_core_properties(prs, spec)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    _rewrite_zip_fixed(out_path)
    return BuildResult(out_path, len(spec["slides"]), ctx.warnings, theme["id"], variants)


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------


def main(argv: Optional[List[str]] = None) -> int:
    ap = argparse.ArgumentParser(
        description="Build a .pptx deck from a pptx-deck spec (JSON). The spec is validated first "
        "(see validate_spec.py), then every slide is rendered natively (text boxes, shapes, pictures, "
        "native charts) with the layout variant and theme of the spec. Output is deterministic.",
    )
    ap.add_argument("spec", help="path to the deck spec JSON file")
    ap.add_argument("-o", "--output", required=True, help="output .pptx path")
    ap.add_argument("--theme", help="force a theme id from references/themes.json (overrides the spec theme; with --template it also overrides the template's colours and fonts)")
    ap.add_argument(
        "--template",
        help="existing .pptx to reuse: its slide masters, layouts, theme fonts and colours are kept, "
        "its slides are removed, and new slides are built on its blank layout. (Filling a template's "
        "existing slides is the job of fill_template.py, built separately.)",
    )
    ap.add_argument("--skip-validate", action="store_true", help="do not run the spec validator before building")
    ap.add_argument("--quiet", "-q", action="store_true", help="print only errors and the one-line summary")
    args = ap.parse_args(argv)

    spec_path = Path(args.spec).resolve()
    try:
        spec = C.load_json(spec_path)
    except (OSError, ValueError) as exc:
        print("ERROR: cannot read spec: %s" % exc, file=sys.stderr)
        return 2
    if not args.skip_validate:
        import validate_spec as V

        errors, warnings = V.validate_spec(spec, spec_path.parent)
        for w in warnings:
            if not args.quiet:
                print(str(w))
        if errors:
            for e in errors:
                print(str(e))
            print("FAIL: spec has %d error(s); nothing was built (use --skip-validate to force)" % len(errors))
            return 1
    if args.theme:
        try:
            C.get_theme(args.theme)
        except KeyError as exc:
            print("ERROR: %s" % exc.args[0], file=sys.stderr)
            return 2
    template = Path(args.template).resolve() if args.template else None
    if template and not template.is_file():
        print("ERROR: template not found: %s" % template, file=sys.stderr)
        return 2
    try:
        res = build_deck(spec, args.output, spec_dir=spec_path.parent, theme_id=args.theme, template=template, validate=False)
    except Exception as exc:  # report as a clean error
        print("ERROR: build failed: %s: %s" % (type(exc).__name__, exc), file=sys.stderr)
        raise
    if not args.quiet:
        for w in res.warnings:
            print("WARNING: " + w)
    print("Built %s: %d slides, theme %s, %d warning(s)" % (res.path, res.slide_count, res.theme_id, len(res.warnings)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
