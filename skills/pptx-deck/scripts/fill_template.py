#!/usr/bin/env python3
"""Build a deck from a user template by selecting, ordering and filling its slides.

Usage:
    fill_template.py TEMPLATE.pptx EDITS.json -o OUT.pptx [--profile profile.json]
                     [--dry-run] [--strict] [--no-renumber]

EDITS.json
----------
{
  "slides": [
    {"source": 1, "slots": {"s1:2": "Quarterly review", "s1:3": ["Line one", "Line two"]}},
    {"source": 5, "slots": {"s5:6": "Heading", "s5:7/2": "text in a group member",
                            "s5:9@2,3": "text in table cell row 2, column 3"},
                   "images": {"s5:4": "photos/team.jpg"},
                   "notes": "Speaker notes for this slide"},
    {"source": 5, "slots": {"s5:6": "A second copy of template slide 5"}}
  ]
}

* ``source`` is the 1-based index of the template slide. Entries give the order
  of the output deck; a source used twice is duplicated (charts, notes slides
  and part names are cloned by rearrange.py); template slides that are not
  listed are dropped.
* Slot ids come from template_profile.py (``s<slide>:<shape id path>`` with an
  optional ``@row,col`` for table cells). A bare path such as ``6`` or ``7/2``
  is accepted too. Only slots flagged ``editable`` can be filled.
* A slot value is a string (``\\n`` separates paragraphs) or a list of
  paragraph strings. Text is written with replace_text.py's engine: position,
  size, colour, font and paragraph formatting stay; the first run's formatting
  is used for every paragraph. Text is never shortened and no "..." is added.
* ``images`` replace a picture and keep its frame; the picture is cropped (never
  stretched) to fill the frame. Image parts shared with other slides are not
  modified.
* Editable slots that are not filled have their sample text cleared (an
  unfilled placeholder shape is removed) with a warning. Running header/footer
  labels that repeat on most slides (``persistent`` in the profile) are kept.
  Static page numbers that match their template slide number are renumbered to
  the new position unless --no-renumber is given.
* A slot whose new text exceeds its soft capacity (profile ``capacity.softChars``,
  about 20 percent over the box estimate) gets a warning; the text is kept in full.

The output is checked with ooxml.py's validator. --dry-run does all of this
without writing OUT.pptx. --strict turns every warning into exit status 1.

Exit status: 0 ok, 1 invalid edits / output fails validation / warnings with --strict,
2 unreadable input.
"""

from __future__ import annotations

import argparse
import json
import sys
import tempfile
from pathlib import Path
from typing import Any, Dict, List, Optional

sys.path.insert(0, str(Path(__file__).resolve().parent))
import _common as C  # noqa: E402

C.require_pptx()

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402

import ooxml  # noqa: E402
import rearrange  # noqa: E402
import replace_text as RT  # noqa: E402
import template_profile as TP  # noqa: E402
from inventory import iter_shapes_abs, localname  # noqa: E402

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A = "{%s}" % NS_A
P = "{%s}" % NS_P
EMU = C.EMU_PER_INCH


class Plan:
    """Validated edits plus the problems found while validating."""

    def __init__(self) -> None:
        self.errors: List[str] = []
        self.warnings: List[str] = []
        self.entries: List[Dict[str, Any]] = []


def _paragraphs(value: Any) -> List[str]:
    if isinstance(value, str):
        return value.split("\n")
    if isinstance(value, list) and all(isinstance(x, str) for x in value):
        return list(value)
    raise ValueError('slot values are a string or a list of strings')


def load_edits(path: str) -> Dict[str, Any]:
    with open(path, "r", encoding="utf-8") as fh:
        data = json.load(fh)
    if not isinstance(data, dict) or not isinstance(data.get("slides"), list) or not data["slides"]:
        raise ValueError('edits file must be {"slides": [{"source": N, "slots": {...}}, ...]} with at least one slide')
    return data


def check_profile(profile: Dict[str, Any], prs: Any) -> None:
    if profile.get("slideCount") != len(prs.slides):
        raise ValueError(
            "profile describes %s slide(s) but the template has %d; regenerate it with template_profile.py"
            % (profile.get("slideCount"), len(prs.slides))
        )


def plan_edits(edits: Dict[str, Any], profile: Dict[str, Any], base_dir: Path) -> Plan:
    plan = Plan()
    n = profile["slideCount"]
    slot_index: Dict[int, Dict[str, Dict[str, Any]]] = {}
    pic_index: Dict[int, Dict[str, Dict[str, Any]]] = {}
    for s in profile["slides"]:
        slot_index[s["index"]] = {x["id"]: x for x in s["slots"]}
        pic_index[s["index"]] = {x["id"]: x for x in s["pictures"]}
    for pos, entry in enumerate(edits["slides"], start=1):
        label = "output slide %d" % pos
        if not isinstance(entry, dict) or not isinstance(entry.get("source"), int):
            plan.errors.append("%s: needs an integer \"source\" (template slide index)" % label)
            continue
        src = entry["source"]
        if not 1 <= src <= n:
            plan.errors.append("%s: source %d is out of range 1..%d" % (label, src, n))
            continue
        unknown = sorted(set(entry) - {"source", "slots", "images", "notes"})
        if unknown:
            plan.errors.append("%s: unknown key(s) %s" % (label, ", ".join(unknown)))
        slots: Dict[str, List[str]] = {}
        for raw_id, value in (entry.get("slots") or {}).items():
            try:
                idx, path, cell = TP.parse_slot_id(raw_id)
            except ValueError as exc:
                plan.errors.append("%s: %s" % (label, exc))
                continue
            if idx is not None and idx != src:
                plan.errors.append("%s: slot %s belongs to template slide %d but source is %d" % (label, raw_id, idx, src))
                continue
            sid = TP.slot_id(src, path, cell)
            slot = slot_index[src].get(sid)
            if slot is None:
                plan.errors.append("%s: no slot %s on template slide %d (see profile)" % (label, sid, src))
                continue
            if not slot["editable"]:
                plan.errors.append("%s: slot %s is not editable (%s)" % (label, sid, slot["editableReason"]))
                continue
            try:
                slots[sid] = _paragraphs(value)
            except ValueError as exc:
                plan.errors.append("%s slot %s: %s" % (label, sid, exc))
        images: Dict[str, Path] = {}
        for raw_id, img in (entry.get("images") or {}).items():
            try:
                idx, path, cell = TP.parse_slot_id(raw_id)
            except ValueError as exc:
                plan.errors.append("%s: %s" % (label, exc))
                continue
            if idx is not None and idx != src:
                plan.errors.append("%s: picture %s belongs to template slide %d but source is %d" % (label, raw_id, idx, src))
                continue
            sid = TP.slot_id(src, path)
            pic = pic_index[src].get(sid)
            if pic is None:
                plan.errors.append("%s: no picture slot %s on template slide %d (see profile)" % (label, sid, src))
                continue
            if not pic["replaceable"]:
                plan.errors.append("%s: picture %s is not replaceable (%s)" % (label, sid, pic.get("replaceableReason", "decorative")))
                continue
            ip = Path(img)
            if not ip.is_absolute():
                ip = base_dir / ip
            if not ip.is_file():
                plan.errors.append("%s: image file not found: %s" % (label, img))
                continue
            images[sid] = ip
        notes = entry.get("notes")
        if notes is not None and not isinstance(notes, (str, list)):
            plan.errors.append("%s: notes must be a string or list of strings" % label)
            notes = None
        plan.entries.append({"pos": pos, "source": src, "slots": slots, "images": images, "notes": notes})
    return plan


# --------------------------------------------------------------------------
# Applying edits to the rearranged deck
# --------------------------------------------------------------------------


def _chars(paras: List[str]) -> int:
    return sum(len(p) for p in paras)


def _remove_shape(shape: Any) -> None:
    el = shape._element
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def fill_picture(slide: Any, shape: Any, image_path: Path, warnings: List[str], label: str) -> None:
    """Replace the picture of *shape* with *image_path*, keeping the frame and cropping to fill it."""
    from PIL import Image

    with Image.open(image_path) as im:
        iw, ih = im.size
    if iw <= 0 or ih <= 0:
        raise ValueError("image has no pixels: %s" % image_path)
    el = shape._element
    # an empty picture placeholder: let python-pptx insert (and crop) the picture
    if hasattr(shape, "insert_picture") and localname(el) == "sp":
        shape.insert_picture(str(image_path))
        return
    blip = el.find(".//" + A + "blip")
    if blip is None:
        raise ValueError("shape has no picture to replace")
    _img_part, rid = slide.part.get_or_add_image_part(str(image_path))
    old = blip.get("{%s}embed" % NS_R)
    blip.set("{%s}embed" % NS_R, rid)
    if old and old != rid:
        try:
            slide.part.drop_rel(old)
        except Exception:
            pass
    fw, fh = shape.width, shape.height
    if fw and fh:
        frame_ar = fw / float(fh)
        img_ar = iw / float(ih)
        l = r = t = b = 0
        if img_ar > frame_ar:
            l = r = int(round((1 - frame_ar / img_ar) / 2 * 100000))
        elif img_ar < frame_ar:
            t = b = int(round((1 - img_ar / frame_ar) / 2 * 100000))
        blip_fill = blip.getparent()
        src = blip_fill.find(A + "srcRect")
        if src is None:
            src = etree.Element(A + "srcRect")
            blip.addnext(src)
        for k in ("l", "t", "r", "b"):
            src.attrib.pop(k, None)
        for k, v in (("l", l), ("t", t), ("r", r), ("b", b)):
            if v:
                src.set(k, str(v))
        if blip_fill.find(A + "stretch") is None and blip_fill.find(A + "tile") is None:
            stretch = etree.SubElement(blip_fill, A + "stretch")
            etree.SubElement(stretch, A + "fillRect")
        kept_w = iw * (1 - 2 * l / 100000.0)
        dpi = kept_w / (fw / float(EMU))
        if dpi < 96:
            warnings.append("%s: image %s is only about %d dpi at this frame size; it may look soft" % (label, image_path.name, dpi))


def apply_plan(prs: Any, plan: Plan, profile: Dict[str, Any], renumber: bool) -> Dict[str, int]:
    """Edit the rearranged *prs* in place (output slide k is plan.entries[k-1])."""
    stats = {"filled": 0, "cleared": 0, "pictures": 0, "renumbered": 0, "kept": 0}
    profile_slides = {s["index"]: s for s in profile["slides"]}
    for entry, slide in zip(plan.entries, prs.slides):
        pos, src = entry["pos"], entry["source"]
        label = "output slide %d (template %d)" % (pos, src)
        pslide = profile_slides[src]
        slot_by_id = {s["id"]: s for s in pslide["slots"]}
        to_remove: List[Any] = []

        # text fills
        for sid, paras in entry["slots"].items():
            shape, txb = TP.find_slot_txbody(slide, sid)
            if txb is None:
                plan.errors.append("%s: slot %s did not resolve to a text body on the copied slide" % (label, sid))
                continue
            RT.set_body_paragraphs(txb, paras)
            stats["filled"] += 1
            cap = (slot_by_id[sid] or {}).get("capacity")
            if cap and cap["softChars"] and _chars(paras) > cap["softChars"]:
                plan.warnings.append(
                    "%s: slot %s over capacity (%d characters, soft limit about %d at %.4g pt); text kept in full"
                    % (label, sid, _chars(paras), cap["softChars"], slot_by_id[sid]["fontPt"])
                )

        # unfilled editable slots
        for slot in pslide["slots"]:
            sid = slot["id"]
            if sid in entry["slots"] or not slot["editable"]:
                continue
            if slot.get("persistent"):
                stats["kept"] += 1
                continue
            shape, txb = TP.find_slot_txbody(slide, sid)
            if txb is None:
                continue
            if not TP.txbody_text(txb).strip():
                continue  # already empty
            sample = slot["sample"].strip().replace("\n", " / ")
            plan.warnings.append(
                "%s: slot %s (%s) was not filled; sample text %r cleared"
                % (label, sid, slot["role"], sample if len(sample) <= 40 else sample[:37] + "...")
            )
            if slot["kind"] == "text" and slot["placeholder"] and slot["placeholder"] not in ("SLIDE_NUMBER", "DATE", "FOOTER"):
                to_remove.append(shape)
            else:
                RT.set_body_paragraphs(txb, [""])
            stats["cleared"] += 1

        # page numbers follow the new position
        if renumber:
            for slot in pslide["slots"]:
                if slot["role"] == "page-number" and slot["sample"].strip() == str(src):
                    shape, txb = TP.find_slot_txbody(slide, slot["id"])
                    if txb is not None and not TP.has_field(txb, ("slidenum",)):
                        RT.set_body_paragraphs(txb, [str(pos)])
                        stats["renumbered"] += 1

        # pictures
        for sid, img in entry["images"].items():
            path = TP.parse_slot_id(sid)[1]
            shape = TP.find_shape_by_path(slide, path)
            if shape is None:
                plan.errors.append("%s: picture %s did not resolve on the copied slide" % (label, sid))
                continue
            try:
                fill_picture(slide, shape, img, plan.warnings, "%s picture %s" % (label, sid))
                stats["pictures"] += 1
            except Exception as exc:
                plan.errors.append("%s: picture %s: %s" % (label, sid, exc))

        for shape in to_remove:
            _remove_shape(shape)

        # notes
        if entry.get("notes") is not None:
            texts = entry["notes"].split("\n") if isinstance(entry["notes"], str) else list(entry["notes"])
            RT.set_body_paragraphs(slide.notes_slide.notes_text_frame._txBody, texts)
    return stats


def main(argv: Optional[List[str]] = None) -> int:
    ap = argparse.ArgumentParser(
        description="Select, order, duplicate and fill the slides of a user template (text and pictures only). "
        "Slot ids come from template_profile.py. Never shortens text; warns on over-capacity and unfilled slots.",
    )
    ap.add_argument("template", help="template .pptx")
    ap.add_argument("edits", help="edits JSON (see the module docstring)")
    ap.add_argument("-o", "--output", required=True, help="output .pptx")
    ap.add_argument("--profile", help="profile JSON from template_profile.py (default: computed on the fly)")
    ap.add_argument("--dry-run", action="store_true", help="do everything except writing the output file")
    ap.add_argument("--strict", action="store_true", help="exit 1 when there are any warnings")
    ap.add_argument("--no-renumber", action="store_true", help="keep static page numbers as in the template")
    args = ap.parse_args(argv)

    try:
        before = ooxml.validate(args.template)
        prs = Presentation(args.template)
        if args.profile:
            with open(args.profile, "r", encoding="utf-8") as fh:
                profile = json.load(fh)
            check_profile(profile, prs)
        else:
            profile = TP.build_profile(prs, args.template)
        edits = load_edits(args.edits)
    except (OSError, ValueError, KeyError) as exc:
        print("error: %s" % exc, file=sys.stderr)
        return 2
    except Exception as exc:
        print("error: cannot read inputs: %s" % exc, file=sys.stderr)
        return 2

    plan = plan_edits(edits, profile, Path(args.edits).resolve().parent)
    if plan.errors:
        for e in plan.errors:
            print("ERROR: " + e, file=sys.stderr)
        print("%d error(s); nothing written" % len(plan.errors), file=sys.stderr)
        return 1

    order = [e["source"] for e in plan.entries]
    summary = rearrange.rearrange(prs, order)
    stats = apply_plan(prs, plan, profile, renumber=not args.no_renumber)
    if plan.errors:
        for e in plan.errors:
            print("ERROR: " + e, file=sys.stderr)
        print("%d error(s); nothing written" % len(plan.errors), file=sys.stderr)
        return 1

    if args.dry_run:
        out_path = Path(tempfile.mkdtemp(prefix="fill-template-")) / "dry-run.pptx"
    else:
        out_path = Path(args.output)
        out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))

    after = ooxml.validate(str(out_path))
    known = {(p["code"], p["part"], p["message"]) for p in before.problems}
    fresh = [p for p in after.problems if (p["code"], p["part"], p["message"]) not in known]

    for w in plan.warnings:
        print("WARNING: " + w)
    for p in after.problems:
        print("PROBLEM [%s] %s: %s" % (p["code"], p["part"], p["message"]), file=sys.stderr)
    print(
        "%s %s: %d slide(s) (%d duplicated), %d slot(s) filled, %d sample(s) cleared, %d kept, %d page number(s) renumbered, %d picture(s) replaced, %d warning(s)"
        % (
            "dry run for" if args.dry_run else "wrote",
            args.output,
            summary["total"],
            summary["copied"],
            stats["filled"],
            stats["cleared"],
            stats["kept"],
            stats["renumbered"],
            stats["pictures"],
            len(plan.warnings),
        )
    )
    print("validation: %s" % ("OK" if not after.problems else "%d problem(s)" % len(after.problems)))
    if fresh:
        print("output failed validation", file=sys.stderr)
        return 1
    if args.strict and plan.warnings:
        print("--strict: %d warning(s) treated as failure" % len(plan.warnings), file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
