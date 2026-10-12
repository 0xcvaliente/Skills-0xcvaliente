#!/usr/bin/env python3
"""Styled workbook from a design preset (needs openpyxl).

  xlsx_starter.py OUT.xlsx [--preset business-modern]

Copy style_header / style_table / add_bar_chart into your own build script. Formulas are
written as formulas; run `python3 scripts/xlsx_check.py OUT.xlsx --recalc` afterwards.
"""
import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TOKENS = Path(__file__).resolve().parents[1] / "design-tokens.json"


def argb(hexstr):
    return hexstr.lstrip("#").upper()


def style_header(ws, p, row, ncols):
    c = p["colors"]
    tb = p["table"]
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(name=p["fonts"]["body"], size=p["type"]["sheet"]["header"], bold=True, color=argb(c[tb["header_text"]]))
        cell.fill = PatternFill("solid", fgColor=argb(c[tb["header_fill"]]))
        cell.alignment = Alignment(horizontal="left" if col == 1 else "right", vertical="center", wrap_text=True)


def style_table(ws, p, first_row, last_row, ncols, number_formats):
    c = p["colors"]
    thin = Side(style="thin", color=argb(c["rule"]))
    for r in range(first_row, last_row + 1):
        for col in range(1, ncols + 1):
            cell = ws.cell(row=r, column=col)
            cell.font = Font(name=p["fonts"]["body"], size=p["type"]["sheet"]["body"], color=argb(c["text"]))
            cell.border = Border(bottom=thin)
            if col > 1:
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = number_formats.get(col, "General")


def add_bar_chart(ws, p, title, data_col, cat_col, first_row, last_row, anchor):
    ch = BarChart()
    ch.type = "col"
    ch.title = title
    ch.legend = None
    ch.y_axis.title = "k EUR"
    ch.x_axis.title = "Region"
    ch.add_data(Reference(ws, min_col=data_col, min_row=first_row - 1, max_row=last_row), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=cat_col, min_row=first_row, max_row=last_row))
    ch.series[0].graphicalProperties.solidFill = argb(p["chart"][0])
    ch.y_axis.scaling.min = 0
    ch.height, ch.width = 8, 14
    ws.add_chart(ch, anchor)


def build(out, preset):
    p = json.loads(TOKENS.read_text(encoding="utf-8"))["presets"][preset]
    if "sheet" not in p.get("media", ["sheet"]):
        raise SystemExit(f"error: preset {preset!r} is not for sheets (allowed media: {', '.join(p['media'])}); choose another preset")
    c = p["colors"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Revenue"
    ws["A1"] = "Revenue by region"
    ws["A1"].font = Font(name=p["fonts"]["heading"], size=p["type"]["sheet"]["title"], bold=True, color=argb(c["primary"]))
    ws["A2"] = "Source: finance ledger, September. Figures in thousands of EUR."
    ws["A2"].font = Font(name=p["fonts"]["body"], size=10, italic=True, color=argb(c["muted"]))
    header = ["Region", "Q2", "Q3", "Change"]
    for i, h in enumerate(header, 1):
        ws.cell(row=4, column=i, value=h)
    data = [("Nordics", 1200, 1380), ("DACH", 980, 1040), ("UK", 820, 930)]
    for r, (name, q2, q3) in enumerate(data, 5):
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=q2)
        ws.cell(row=r, column=3, value=q3)
        ws.cell(row=r, column=4, value=f"=C{r}/B{r}-1")
    last = 4 + len(data)
    tot = last + 1
    ws.cell(row=tot, column=1, value="Total")
    ws.cell(row=tot, column=2, value=f"=SUM(B5:B{last})")
    ws.cell(row=tot, column=3, value=f"=SUM(C5:C{last})")
    ws.cell(row=tot, column=4, value=f"=C{tot}/B{tot}-1")
    style_header(ws, p, 4, 4)
    style_table(ws, p, 5, tot, 4, {2: "#,##0", 3: "#,##0", 4: "0.0%"})
    for col in range(1, 5):
        ws.cell(row=tot, column=col).font = Font(name=p["fonts"]["body"], size=p["type"]["sheet"]["body"], bold=True, color=argb(c["text"]))
        ws.cell(row=tot, column=col).border = Border(top=Side(style="medium", color=argb(c["text"])))
    ws.column_dimensions["A"].width = 18
    for col in "BCD":
        ws.column_dimensions[col].width = 14
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "4:4"
    add_bar_chart(ws, p, "Q3 revenue by region", 3, 1, 5, last, "F4")
    wb.properties.title = "Revenue by region"
    wb.save(out)
    print(f"wrote {out} ({preset})")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("out")
    ap.add_argument("--preset", default="business-modern")
    a = ap.parse_args()
    build(a.out, a.preset)
