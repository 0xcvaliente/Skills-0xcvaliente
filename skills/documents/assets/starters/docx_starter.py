#!/usr/bin/env python3
"""Styled Word document from a design preset (needs python-docx).

  docx_starter.py OUT.docx [--preset report-classic]

Copy the functions you need (apply_tokens, add_table, add_page_number) into your own build
script. Formatting lives in styles, not in individual paragraphs.
"""
import argparse
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

TOKENS = Path(__file__).resolve().parents[1] / "design-tokens.json"


def rgb(hexstr):
    h = hexstr.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def set_font(style, name, size, color=None, bold=None, italic=None):
    f = style.font
    f.name = name
    f.size = Pt(size)
    if color:
        f.color.rgb = rgb(color)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def apply_tokens(doc, p, language="en-US"):
    c, t, space = p["colors"], p["type"]["document"], p["space"]
    sec = doc.sections[0]
    a4 = p["page"]["size"].upper() == "A4"
    sec.page_width, sec.page_height = (Mm(210), Mm(297)) if a4 else (Mm(215.9), Mm(279.4))
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Mm(p["page"]["margin_mm"]))
    heading, body = p["fonts"]["heading"], p["fonts"]["body"]
    st = doc.styles
    set_font(st["Normal"], body, t["body"], c["text"])
    pf = st["Normal"].paragraph_format
    pf.space_after = Pt(space["paragraph_after_pt"])
    pf.line_spacing = t["line_height"]
    pf.widow_control = True
    set_font(st["Title"], heading, t["title"], c["primary"], bold=True)
    st["Title"].paragraph_format.space_after = Pt(space["scale"][3])
    for bottom in st["Title"].element.iter(qn("w:bottom")):      # the default template draws a blue rule under titles
        bottom.set(qn("w:color"), c["rule"].lstrip("#"))
    for level, key in ((1, "h1"), (2, "h2"), (3, "h3")):
        s = st[f"Heading {level}"]
        set_font(s, heading, t[key], c["primary"] if level < 3 else c["text"], bold=True)
        s.paragraph_format.space_before = Pt(space["heading_before_pt"] - (level - 1) * 4)
        s.paragraph_format.space_after = Pt(space["scale"][1])
        s.paragraph_format.keep_with_next = True
        s.paragraph_format.line_spacing = 1.15
    set_font(st["Subtitle"], body, t["h2"], c["muted"], bold=False, italic=False)
    set_font(st["Caption"], body, t["caption"], c["muted"], bold=False, italic=True)
    st["Caption"].paragraph_format.space_before = Pt(space["scale"][1])
    for name in ("List Bullet", "List Number"):
        st[name].paragraph_format.space_after = Pt(space["scale"][0])
    rpr = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), language)


def shade(cell, fill_hex):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex.lstrip("#"))
    tcpr.append(shd)


def borders(cell, rule_hex, top=False, bottom=True):
    tcpr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for edge, on in (("top", top), ("bottom", bottom)):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single" if on else "nil")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), rule_hex.lstrip("#"))
        b.append(e)
    tcpr.append(b)


def add_table(doc, p, header, rows, number_cols=(), widths_mm=None):
    c, tb = p["colors"], p["table"]
    t = doc.add_table(rows=1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    hdr_fill = c[tb["header_fill"]]
    for i, text in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.color.rgb = rgb(c[tb["header_text"]])
        shade(cell, hdr_fill)
        borders(cell, c[tb["rule"]])
        if i in number_cols:
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    trpr = t.rows[0]._tr.get_or_add_trPr()          # repeat header row on every page
    trpr.append(OxmlElement("w:tblHeader"))
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = str(text)
            borders(cells[i], c[tb["rule"]])
            if i in number_cols:
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if tb["zebra"] and ri % 2 == 1:
                shade(cells[i], c["surface"])
    for row in t.rows:
        for cell in row.cells:
            for par in cell.paragraphs:
                par.paragraph_format.space_after = Pt(2)
                par.paragraph_format.space_before = Pt(2)
    if widths_mm:
        for row in t.rows:
            for i, w in enumerate(widths_mm):
                row.cells[i].width = Mm(w)
    return t


def add_page_number(doc, p):
    par = doc.sections[0].footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = par.add_run()
    for kind, text in (("begin", None), (None, " PAGE "), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run.font.size = Pt(p["type"]["document"]["footer"])
    run.font.color.rgb = rgb(p["colors"]["muted"])


def build(out, preset):
    p = json.loads(TOKENS.read_text(encoding="utf-8"))["presets"][preset]
    if "document" not in p.get("media", ["document"]):
        raise SystemExit(f"error: preset {preset!r} is not for documents (allowed media: {', '.join(p['media'])}); choose another preset")
    doc = Document()
    apply_tokens(doc, p)
    doc.core_properties.title = "Quarterly review"
    doc.core_properties.author = "Team"
    doc.add_paragraph("Quarterly review", style="Title")
    doc.add_paragraph("Prepared for the leadership team", style="Subtitle")
    doc.add_heading("Summary", level=1)
    doc.add_paragraph("Revenue grew 12% quarter on quarter, driven by the new onboarding flow. "
                      "Churn held flat. The one risk is support load, which rose 18%.")
    doc.add_heading("Key results", level=1)
    for text in ("Revenue: 4.2 M EUR, up 12%", "Churn: 2.1%, unchanged", "Support tickets: up 18%"):
        doc.add_paragraph(text, style="List Bullet")
    doc.add_heading("By region", level=2)
    add_table(doc, p, ["Region", "Q2 (k EUR)", "Q3 (k EUR)", "Change"],
              [["Nordics", "1 200", "1 380", "+15%"], ["DACH", "980", "1 040", "+6%"], ["UK", "820", "930", "+13%"]],
              number_cols=(1, 2, 3), widths_mm=[55, 35, 35, 30])
    doc.add_paragraph("Table 1. Revenue by region. Source: finance ledger, September.", style="Caption")
    add_page_number(doc, p)
    doc.save(out)
    print(f"wrote {out} ({preset})")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("out")
    ap.add_argument("--preset", default="report-classic")
    a = ap.parse_args()
    try:
        build(a.out, a.preset)
    except KeyError as e:
        sys.exit(f"error: unknown preset or style {e}")
