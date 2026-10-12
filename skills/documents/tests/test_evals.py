import importlib.util
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
CHECK = os.path.join(ROOT, "evals", "check_outputs.py")
FILES = os.path.join(ROOT, "evals", "files")
SCRIPTS = os.path.join(ROOT, "scripts")
sys.path.insert(0, SCRIPTS)
from _common import find_soffice  # noqa: E402


def has(m):
    return importlib.util.find_spec(m) is not None


def check(eval_id, out):
    r = subprocess.run([sys.executable, CHECK, str(eval_id), out], capture_output=True, text=True)
    return r.returncode, r.stdout + r.stderr


class EvalFiles(unittest.TestCase):
    def test_fixtures_match_checksums(self):
        import hashlib
        import json
        with open(os.path.join(FILES, "checksums.json")) as fh:
            sums = json.load(fh)
        for name, digest in sums.items():
            with open(os.path.join(FILES, name), "rb") as fh:
                self.assertEqual(hashlib.sha256(fh.read()).hexdigest(), digest, name)

    def test_evals_json_is_well_formed(self):
        import json
        with open(os.path.join(ROOT, "evals", "evals.json")) as fh:
            data = json.load(fh)
        self.assertEqual(data["skill_name"], "documents")
        for e in data["evals"]:
            self.assertTrue(e["prompt"] and e["expected_output"] and e["assertions"])
            for f in e["files"]:
                self.assertTrue(os.path.exists(os.path.join(ROOT, f)), f)
        with open(os.path.join(ROOT, "evals", "trigger-queries.json")) as fh:
            q = json.load(fh)
        self.assertEqual(len(q), 20)
        self.assertEqual(sum(1 for x in q if x["should_trigger"]), 10)
        self.assertTrue(any(x["split"] == "validation" and x["should_trigger"] for x in q))
        self.assertTrue(any(x["split"] == "validation" and not x["should_trigger"] for x in q))


@unittest.skipUnless(has("pypdf") and has("reportlab"), "pypdf and reportlab needed")
class PdfEvals(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = self.tmp.name

    def tearDown(self):
        self.tmp.cleanup()

    def test_eval5_pass_and_fail(self):
        import json
        out = os.path.join(self.d, "out")
        os.makedirs(out)
        vals = os.path.join(self.d, "v.json")
        with open(vals, "w") as fh:
            json.dump({"full_name": "Ingrid Solheim", "start_date": "1 December 2025"}, fh)
        subprocess.run([sys.executable, os.path.join(SCRIPTS, "pdf_tool.py"), "fill", os.path.join(FILES, "form.pdf"),
                        os.path.join(out, "filled.pdf"), "--values", vals], check=True, capture_output=True)
        code, text = check(5, out)
        self.assertEqual(code, 0, text)
        with open(vals, "w") as fh:
            json.dump({"full_name": "Ingrid Solheim", "start_date": "1 December 2025", "email": "x@example.test"}, fh)
        subprocess.run([sys.executable, os.path.join(SCRIPTS, "pdf_tool.py"), "fill", os.path.join(FILES, "form.pdf"),
                        os.path.join(out, "filled2.pdf"), "--values", vals], check=True, capture_output=True)
        os.remove(os.path.join(out, "filled.pdf"))
        code, text = check(5, out)
        self.assertEqual(code, 1, text)
        self.assertIn("email left empty", text)

    def test_eval6_detects_surviving_text(self):
        out = os.path.join(self.d, "out")
        os.makedirs(out)
        shutil.copy(os.path.join(FILES, "statement.pdf"), os.path.join(out, "sent.pdf"))   # not redacted at all
        code, text = check(6, out)
        self.assertEqual(code, 1, text)
        from reportlab.pdfgen import canvas
        c = canvas.Canvas(os.path.join(out, "sent.pdf"))
        c.drawString(60, 700, "Account statement, September")
        c.drawString(60, 680, "Account number: [REDACTED]")
        c.save()
        code, text = check(6, out)
        self.assertEqual(code, 0, text)


@unittest.skipUnless(has("openpyxl"), "openpyxl needed")
class SheetEvals(unittest.TestCase):
    def test_eval2_catches_formulas_and_lost_zeros(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as d:
            wb = Workbook()
            ws = wb.active
            ws.append(["customer_id", "name"])
            ws.append([123, "=HYPERLINK(\"x\")"])
            wb.save(os.path.join(d, "clean.xlsx"))
            code, text = check(2, d)
            self.assertEqual(code, 1, text)
            self.assertIn("FAIL", text)

    @unittest.skipUnless(find_soffice(), "LibreOffice needed")
    def test_eval1_pass(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as d:
            wb = load_workbook(os.path.join(FILES, "budget.xlsx"))
            ws = wb.active
            for r in range(2, 6):
                ws[f"D{r}"] = f"=IFERROR((C{r}-B{r})/B{r},\"n/a\")"
            ws["D8"] = "=IFERROR(AVERAGE(D2:D5),\"n/a\")"
            ws["D9"] = "=IFERROR(VLOOKUP(\"Software\",A2:D5,2,FALSE),\"not found\")"
            wb.save(os.path.join(d, "budget-fixed.xlsx"))
            code, text = check(1, d)
            self.assertEqual(code, 0, text)
            self.assertIn("NEEDS-GRADER", text)


if __name__ == "__main__":
    unittest.main()
